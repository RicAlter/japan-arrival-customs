import sys,tempfile,unittest
from pathlib import Path
from openpyxl import Workbook
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT/'workbuddy/japan-arrival-customs/scripts'))
from arrival_core.workbuddy_router import route,inspect_workbooks,historical_references,sync_d1
from arrival_core.core import Rejected
class RouterTests(unittest.TestCase):
 def test_phrases_have_one_controlled_command(self): self.assertEqual(route('生成这批到货的审核表')['command'],'return-export')
 def test_ambiguous_message_is_rejected(self):
  with self.assertRaises(Rejected):route('处理这批日本到货并生成这批到货的审核表')
 def test_unknown_workbook_never_guessed(self):
  with tempfile.TemporaryDirectory() as d:
   p=Path(d)/'x.xlsx';Workbook().save(p);r=inspect_workbooks([p]);self.assertFalse(r['ready']);self.assertEqual(r['unknown'],['x.xlsx'])
 def test_controlled_template_signatures_and_conflict(self):
  with tempfile.TemporaryDirectory() as d:
   d=Path(d)
   def book(name,values):
    wb=Workbook();ws=wb.active
    for cell,value in values.items():ws[cell]=value
    path=d/name;wb.save(path);wb.close();return path
   stage=book('stage.xlsx',{'C1':'ProductCode(JANCode)','J1':'Qty','K1':'UnitPrice(日本JPY)','L1':'Subtotal'})
   invoice=book('invoice.xlsx',{'D17':'JAN','M17':'Qty'})
   weight=book('weight.xlsx',{'A1':'箱番','D1':'Total KGS'})
   result=inspect_workbooks([stage,invoice,weight]);self.assertTrue(result['ready'])
   duplicate=inspect_workbooks([stage,book('stage2.xlsx',{'C1':'JAN','J1':'Qty','K1':'UnitPrice(日本JPY)','L1':'Subtotal'})])
   self.assertIn('stage3_authority',duplicate['conflicts']);self.assertFalse(duplicate['ready'])
 def test_d1_reference_and_sync_failures_are_nonblocking(self):
  with tempfile.TemporaryDirectory() as d:
   root=Path(d)
   reply=sync_d1(root,root/'missing-library','missing-batch');self.assertFalse(reply['batch_saved']);self.assertTrue(reply['retry_required'])
   reference=historical_references(root/'missing-library',{'product_type':'挂件'},'2026-10-21')
   self.assertEqual(reference['status'],'unavailable');self.assertFalse(reference['automatic_adoption'])
if __name__=='__main__':unittest.main()
