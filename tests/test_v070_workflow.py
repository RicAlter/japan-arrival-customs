"""Portable full Excel roundtrip and B2->D1 integration, explicitly synthetic."""
import copy, hashlib, json, os, sys, tempfile, unittest
from datetime import datetime,timedelta,timezone
from pathlib import Path
from unittest.mock import patch
from openpyxl import Workbook,load_workbook
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT/'workbuddy/japan-arrival-customs/scripts'))
from arrival_core.workflow import execute,describe,invoke
from arrival_core.storage import Store,atomic_json
from arrival_core.core import Rejected,digest
from arrival_core.decision_library import DecisionLibrary
from arrival_core.source_adapters import official_payload

class WorkflowTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.root=Path(self.tmp.name)
        self.delivery=self.root/'delivery';self.delivery.mkdir()
        self.inputs=self.root/'inputs';self.inputs.mkdir();self.output=self.root/'store'
        self.stage=self.inputs/'any-name.xlsx'
        w=Workbook();s=w.active
        for cell,value in {'C1':'ProductCode(JANCode)','J1':'Qty','K1':'UnitPrice(日本JPY)','L1':'Subtotal',
                           'A2':'1','B2':'JP','C2':'TEST0001','E2':'TEST内部名','F2':'v1','I2':'TEST商品',
                           'J2':'2','K2':'10','L2':'20','M2':'装饰品','N2':'PVC','P2':'TEST-YD'}.items():s[cell]=value
        w.save(self.stage);w.close()
        self.invoice=self.inputs/'japanese.xlsx';w=Workbook();s=w.active;s['D17']='JAN';s['M17']='Qty'
        for cell,value in {'A18':'1','D18':'TEST0001','F18':'TEST内部名','G18':'v1','H18':'TEST商品','K18':'JP','L18':'10','M18':'2','O18':'20'}.items():s[cell]=value
        weights=w.create_sheet('weights');weights['A1']='箱番';weights['D1']='Total KGS';weights['A2']='1';weights['C2']='1';weights['D2']='2'
        w.save(self.invoice);w.close()
        self.history=self.inputs/'material-history.xlsx';w=Workbook();s=w.active
        s['C2']='TEST0001';s['F2']='TEST商品';s['E2']='v1';s['J2']='PVC、铁、纸等';w.save(self.history);w.close()
        self.imported=execute(self.output,{'message':'处理这批日本到货','files':[str(self.stage),str(self.invoice)],'currency':'JPY'})
        self.batch_id=self.imported['batch']['batch_id']
    def tearDown(self):self.tmp.cleanup()
    def export(self):
        return execute(self.output,{'message':'生成这批到货的审核表','batch':self.batch_id,'material_history':str(self.history),'file':str(self.delivery/'review.xlsx')})
    def configure_synthetic(self):
        store=Store(self.output);batch=store.load(self.batch_id)
        # New technical batch from an unchanged source, never upgrade a stored revision.
        synthetic=copy.deepcopy(batch);synthetic['id']=digest(['explicit-synthetic',batch['id']])[:24]
        from arrival_core.contracts import sign_contract,ROLES
        ctr=synthetic['input_contract'];ctr['responsibilities']={role:['TECH'] for role in ROLES}
        ctr.update(authority_quantity_unit='件',invoice_quantity_unit='件')
        synthetic['input_contract']=sign_contract(ctr,actor='TECH',evidence='TEST ONLY D08',permissions={'TECH':['contract']})
        synthetic['execution_mode']='synthetic_technical_test';store.save(synthetic);self.batch_id=synthetic['id']
        config={'purpose':'synthetic_technical_test','sessions':[{'principal':'TECH','roles':['material','customs'],
                'mode':'synthetic_technical_test','credential_ref':'TECH-NOT-REAL-D08',
                'token_sha256':hashlib.sha256(b'TECH-ONLY').hexdigest(),
                'expires_at':(datetime.now(timezone.utc)+timedelta(days=1)).isoformat()}]}
        (self.output/'security').mkdir();atomic_json(self.output/'security/trusted_sessions.json',config)
    def modified_preview(self):
        self.export();w=load_workbook(self.delivery/'review.xlsx');s=w['待处理问题']
        for row in s.iter_rows(min_row=7):
            if '材质' in str(row[2].value) or '材质' in str(row[3].value):
                row[5].value='补充资料';row[6].value='范围=本组全部；申报材质=TEST人工草稿塑料；原因=TEST核实，待审';break
        else:raise AssertionError('material question missing')
        w.save(self.delivery/'modified.xlsx');w.close()
        return execute(self.output,{'message':'这是我修改后的审核表，继续处理','file':str(self.delivery/'modified.xlsx'),'batch':self.batch_id})
    def test_same_import_resumes_without_revision_or_source_changes(self):
        before=Store(self.output).load(self.batch_id)
        rerun=execute(self.output,{'message':'处理这批日本到货','files':[str(self.stage),str(self.invoice)],'currency':'JPY'})
        self.assertTrue(rerun['result']['resumed']);self.assertEqual(digest(before),digest(Store(self.output).load(self.batch_id)))
    def test_new_source_cannot_resume_old_batch(self):
        unknown=self.inputs/'different.txt';unknown.touch()
        with self.assertRaises(Rejected):execute(self.output,{'message':'处理这批日本到货','batch':self.batch_id,'files':[str(unknown)]})
    def test_compact_export_idempotent_and_three_sheets(self):
        r=self.export();again=self.export();self.assertTrue(again['result']['replayed'])
        w=load_workbook(self.delivery/'review.xlsx');self.assertEqual(w.sheetnames,['处理总览','待处理问题','完整商品明细']);w.close()
        self.assertEqual(r['result']['export_id'],again['result']['export_id'])
    def test_auto_stage_recovers_persistent_preview(self):
        preview=self.modified_preview();self.assertFalse(preview['result']['commit_executed'])
        self.assertEqual(describe(self.output,self.batch_id)['stage'],'modified_pending_confirmation')
    def test_missing_trusted_identity_rejects_commit_and_preserves_batch(self):
        preview=self.modified_preview()['result'];before=digest(Store(self.output).load(self.batch_id))
        req={'message':'确认提交这批审核','batch':self.batch_id,'confirm':True,'preview_id':preview['preview_id'],
             'confirmed_hash':preview['confirmation_hash'],'reference':'TEST explicit','session_token_env':'MISSING_TEST_SESSION'}
        with self.assertRaises(Rejected):execute(self.output,req)
        self.assertEqual(before,digest(Store(self.output).load(self.batch_id)))
    def test_full_roundtrip_sync_drafts_and_cli_consistency(self):
        self.configure_synthetic();preview=self.modified_preview()['result']
        lib=DecisionLibrary.create(self.root/'knowledge',source_root=self.output,source_kind='synthetic')
        req={'message':'确认提交这批审核','batch':self.batch_id,'confirm':True,'preview_id':preview['preview_id'],
             'confirmed_hash':preview['confirmation_hash'],'reference':'TECH explicit preview confirmation',
             'session_token_env':'TECH_SESSION','library':str(lib.root)}
        with patch.dict(os.environ,{'TECH_SESSION':'TECH-ONLY'}):
            saved=execute(self.output,req);again=execute(self.output,req)
        self.assertTrue(again['result']['replayed']);self.assertTrue(saved['result']['knowledge_sync']['batch_saved'])
        self.assertEqual(lib.statistics(as_of='2026-10-21')['formal_learning_samples'],0)
        direct=invoke(self.output,'render-drafts','--batch',self.batch_id)
        pair=execute(self.output,{'message':'生成海关和内部到货表','batch':self.batch_id})
        self.assertEqual(pair['result'],direct);self.assertFalse(pair['formal_output'])
        with self.assertRaises(Rejected):invoke(self.output,'formal-export','--batch',self.batch_id)
    def test_d1_failure_keeps_successful_B2_commit_and_can_retry(self):
        self.configure_synthetic();preview=self.modified_preview()['result']
        request={'message':'确认提交这批审核','batch':self.batch_id,'confirm':True,'preview_id':preview['preview_id'],
                 'confirmed_hash':preview['confirmation_hash'],'reference':'TEST save','session_token_env':'TECH_SESSION','library':str(self.root/'missing-library')}
        with patch.dict(os.environ,{'TECH_SESSION':'TECH-ONLY'}):value=execute(self.output,request)
        self.assertTrue(value['result']['knowledge_sync']['batch_saved']);self.assertTrue(value['result']['knowledge_sync']['retry_required'])
        lib=DecisionLibrary.create(self.root/'missing-library',source_root=self.output,source_kind='synthetic')
        from arrival_core.workbuddy_router import sync_d1
        self.assertFalse(sync_d1(self.output,lib.root,self.batch_id)['retry_required'])
    def test_official_identity_reference_never_overwrites_human_final(self):
        batch=Store(self.output).load(self.batch_id);before=digest(batch)
        evidence={'kind':'official_web','complete_composition':{'original':'PVC、鉄、紙等','translation':'PVC、铁、纸等'},
                  'identity_evidence':{'code':'TEST0001','version':'v1','specification':''},'url':'https://official.example/p',
                  'content_sha256':'TESTHASH','collected_at':'2026-10-21'}
        r=official_payload(batch,[evidence]);self.assertEqual(len(r['official_materials']),1);self.assertEqual(before,digest(batch))
    def test_multiple_workbooks_detect_invoice_and_embedded_weight(self):
        invoice=self.inputs/'source.xlsx';w=Workbook();s=w.active;s['D17']='JAN';s['M17']='Qty'
        s['A18']='1';s['D18']='TEST0001';s['G18']='TEST商品';s['M18']='2';s['O18']='20'
        weights=w.create_sheet('weights');weights['A1']='箱番';weights['D1']='Total KGS';weights['A2']='1';weights['C2']='1';weights['D2']='2'
        w.save(invoice);w.close()
        from arrival_core.workbuddy_router import inspect_workbooks
        result=inspect_workbooks([self.stage,invoice]);self.assertTrue(result['ready'])
        self.assertIn('weight',result['files'][1]['roles']);self.assertIn('invoice',result['files'][1]['roles'])

if __name__=='__main__':unittest.main()
