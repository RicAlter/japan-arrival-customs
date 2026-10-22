"""Raw Japanese workbook intake: no fabricated stage-3 authority or approval."""
import hashlib
import os
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook,load_workbook

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'workbuddy/japan-arrival-customs/scripts'))
from arrival_core.core import Rejected,digest
from arrival_core.importer import import_raw_invoice
from arrival_core.storage import Store
from arrival_core.workflow import execute
from arrival_core.workflow import describe
from arrival_core.workbuddy_router import inspect_workbooks


def make_raw(path):
    wb=Workbook();ws=wb.active;ws.title='invoice (計算)'
    for cell,value in {'A17':'CartonNo.','C17':'ProductCode(JANCode)','D17':'DescriptionofGoods',
        'G17':'Origin','H17':'UnitPrice(日本JPY)','I17':'Qty','K17':'Subtotal',
        'A18':'HA001','C18':'4900000000001','D18':'商品A','G18':'JP','H18':100,'I18':2,'K18':200,
        'A19':'HA002','C19':'4900000000002','D19':'商品B','G19':'CN','H19':300,'I19':3,'K19':900,
        'G21':'TOTAL','I21':5,'K21':1100}.items():ws[cell]=value
    weights=wb.create_sheet('重さとサイズ')
    for cell,value in {'A1':'箱番','B1':'重さ','C1':'カートン個数','D1':'合計重量（KGS)',
                       'A2':'HA001','B2':2,'C2':1,'D2':2,'A3':'HA002','B3':3,'C3':1,'D3':3}.items():weights[cell]=value
    wb.save(path);wb.close()


class RawIntakeTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.out_tmp=tempfile.TemporaryDirectory();self.root=Path(self.tmp.name)
        self.raw=self.root/'raw.xlsx';make_raw(self.raw);self.before=hashlib.sha256(self.raw.read_bytes()).hexdigest()
        self.store=Path(self.out_tmp.name)/'output'
    def tearDown(self):self.out_tmp.cleanup();self.tmp.cleanup()

    def test_raw_layout_and_source_cells_are_exact(self):
        result=inspect_workbooks([self.raw]);self.assertTrue(result['ready'])
        self.assertEqual(result['input_mode'],'raw_invoice_preliminary')
        self.assertEqual(result['files'][0]['weight_sheets'],['重さとサイズ'])
        batch=import_raw_invoice(self.raw)
        self.assertEqual(len(batch['rows']),2);self.assertEqual(batch['arrival_summary']['subtotal_jpy'],'1100')
        self.assertEqual(batch['arrival_summary']['qty_numeric_sum'],'5');self.assertIsNone(batch['arrival_summary']['qty_common_unit'])
        row=batch['rows'][0]
        self.assertEqual(row['source']['sheet'],'invoice (計算)');self.assertEqual(row['source']['row'],18)
        self.assertEqual(row['source']['cells']['code']['cell'],'C18');self.assertEqual(row['source']['cells']['qty']['cell'],'I18')
        self.assertEqual(row['source']['cells']['amount']['cell'],'K18')
        self.assertEqual(row['authority']['stage'],1);self.assertEqual(row['authority']['status'],'preliminary_from_raw_invoice')
        self.assertEqual(row['arrival']['quantity_raw'],'2');self.assertEqual(row['arrival']['unit_status'],'pending_product_or_purchase_lookup')
        self.assertEqual(hashlib.sha256(self.raw.read_bytes()).hexdigest(),self.before)

    def test_raw_workbook_alone_creates_preliminary_and_registered_review(self):
        response=execute(self.store,{'message':'处理这批日本到货','files':[str(self.raw)]})
        batch_id=response['batch']['batch_id'];batch=Store(self.store).load(batch_id)
        self.assertEqual(batch['execution_mode'],'preliminary_untrusted');self.assertIsNone(batch['input_contract']['approval'])
        artifacts=response['result']['artifacts'];pre=Path(artifacts['preliminary']['file']);review=Path(artifacts['review']['file'])
        self.assertTrue(pre.is_file());self.assertTrue(review.is_file())
        brief=describe(self.store,batch_id)
        self.assertNotIn('rows',brief['checks']['quantity'])
        self.assertNotIn('rows',brief['blockers'])
        self.assertNotIn('original_identity_pending',brief['checks']['quantity']['pending_reason_counts'])
        self.assertNotIn('frontend_status_pending',brief['blockers']['pending_reason_counts'])
        self.assertLess(len(str(brief)),20000)
        wb=load_workbook(pre,read_only=True);self.assertEqual(wb.sheetnames,['初步整理']);self.assertEqual(wb['初步整理'].max_row-6,2);wb.close()
        wb=load_workbook(review);self.assertEqual(wb.sheetnames,['处理总览','待处理问题','完整商品明细'])
        detail=wb['完整商品明细'];self.assertEqual(detail.max_row-6,2);self.assertIsNone(detail['E7'].value)
        issues=wb['待处理问题']
        for row in issues.iter_rows(min_row=7):
            if '基础资料候选缺失' in str(row[0].value):
                row[5].value='补充资料';row[6].value='中文名=业务补录A；范围=1；依据=TEST only';break
        else:raise AssertionError('missing product issue')
        modified=self.root/'modified.xlsx';wb.save(modified);wb.close()
        preview=execute(self.store,{'message':'这是我修改后的审核表，继续处理','file':str(modified),'batch':batch_id})
        self.assertFalse(preview['result']['commit_executed'])
        before=digest(Store(self.store).load(batch_id))
        with self.assertRaises(Rejected):
            execute(self.store,{'message':'确认提交这批审核','batch':batch_id,'confirm':True,
                'preview_id':preview['result']['preview_id'],'confirmed_hash':preview['result']['confirmation_hash'],
                'reference':'TEST','session_token_env':'ABSENT'})
        self.assertEqual(before,digest(Store(self.store).load(batch_id)))
        drafts=execute(self.store,{'message':'生成海关和内部到货表','batch':batch_id})['result']
        self.assertFalse(drafts['comparison']['formal_output']);self.assertEqual(drafts['comparison']['privacy_scan'],'passed')
        self.assertEqual(hashlib.sha256(self.raw.read_bytes()).hexdigest(),self.before)

    @unittest.skipUnless(os.environ.get('KMT_REAL_RAW_WORKBOOK'),'real KMT workbook path not supplied')
    def test_real_kmt_2687_rows_and_cached_totals(self):
        path=Path(os.environ['KMT_REAL_RAW_WORKBOOK']);before=hashlib.sha256(path.read_bytes()).hexdigest()
        batch=import_raw_invoice(path)
        self.assertEqual(len(batch['rows']),2687)
        self.assertEqual(batch['arrival_summary']['subtotal_jpy'],'87473076')
        self.assertEqual(batch['arrival_summary']['qty_numeric_sum'],'79924')
        self.assertEqual(hashlib.sha256(path.read_bytes()).hexdigest(),before)


if __name__=='__main__':unittest.main()
