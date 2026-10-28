"""Operations intake tests use synthetic files; no price/rate inference."""
import copy
import hashlib
import json
import sys
import tempfile
import unittest
from pathlib import Path
from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'workbuddy/japan-arrival-customs/scripts'))
from arrival_core.operations_adapter import run, prepare
from arrival_core.core import Rejected
from arrival_core.report_templates import cache_formulas
from arrival_core.workflow import execute


def fixtures(root, formula=False):
    business = root / 'arbitrary-name-a.xlsx'; declaration = root / 'arbitrary-name-b.xlsx'
    wb = Workbook(); ws = wb.active; ws.title = 'Sheet1'
    ws.append(['CartonNo.', 'Origin', 'ProductCode(JANCode)', None, None, None, None, None,
               'DescriptionofGoods', 'Qty', 'UnitPrice(日本JPY)', 'Subtotal', '品类', '材质', '渠道', 'yd'])
    ws.append(['HA001', 'JP', '0010001', None, 'PRIVATE-INTERNAL', None, None, None, 'Synthetic Japanese',
               2, 50, '=J2*K2' if formula else 100, 'toy', 'PVC', 'PRIVATE-CHANNEL', None])
    weights = wb.create_sheet('重さとサイズ (2)')
    weights.append(['箱番', '重さ', 'カートン個数', '合計重量（KGS)'])
    weights.append(['HA001', 2, 1, 2]); wb.save(business); wb.close()
    if formula: cache_formulas(business, {1: {'L2': 100}, 2: {}})
    wb = Workbook(); ws = wb.active; ws.title = '图片处理'
    ws.append(['箱号', '品番', 'Items & Descriptions', '原產地', '数量（盒装', 'Qty（数量按单个）',
               '單價', '品牌', '品名', '品类', '是否\n盲盒', '是否需要打标签', '材质'])
    ws.append(['HA001', '0010001', 'Synthetic toy', 'Japan', 1, 3, None, None, None, 'toy', None, None, 'PVC'])
    wb.save(declaration); wb.close()
    return [str(business), str(declaration)]


def confirmation(preview):
    return {'confirm': True, 'business_unit': '件', 'internal_unit': '件', 'declaration_unit': '个',
            'reference': 'TEST:explicit reviewed batch units and conversion list',
            'source_bundle_hash': preview['source_bundle_hash'], 'special_conversion_hash': preview['special_conversion_hash']}


class AdapterTests(unittest.TestCase):
    def setUp(self):
        self.source = tempfile.TemporaryDirectory(); self.output = tempfile.TemporaryDirectory()
        self.root = Path(self.source.name); self.dest = Path(self.output.name) / 'state'
        self.files = fixtures(self.root, formula=True)

    def tearDown(self): self.output.cleanup(); self.source.cleanup()

    def change(self, file, sheet, cell, value):
        wb = load_workbook(file); wb[sheet][cell] = value; wb.save(file); wb.close()

    def test_preview_confirmation_provenance_pair_and_replay(self):
        before = [hashlib.sha256(Path(p).read_bytes()).hexdigest() for p in self.files]
        preview = execute(self.dest, {'message': '处理已确认运营表', 'files': self.files})
        self.assertEqual(preview['status'], 'needs_confirmation'); self.assertEqual(len(preview['special_conversions']), 1)
        request = {'message': '处理已确认运营表', 'files': self.files, 'units_confirmation': confirmation(preview)}
        result = execute(self.dest, request)
        self.assertEqual(result['status'], 'ok', result)
        plan = json.loads(Path(result['provenance']).read_text(encoding='utf-8'))
        proof = plan['details'][0]['fields']['业务确认金额JPY']
        self.assertEqual(proof['formula'], '=J2*K2'); self.assertEqual(proof['cell'], 'L2')
        self.assertEqual(proof['validation'], 'exact_product_matched')
        self.assertEqual(proof['path'], self.files[0]); self.assertEqual(len(proof['precedents']), 2)
        self.assertEqual(plan['details'][0]['values']['商品编码'], '0010001')
        self.assertEqual(result['gates']['draft']['price_pending_rows'], 1)
        self.assertFalse(result['formal_output']); self.assertTrue(execute(self.dest, request)['result']['replayed'])
        self.assertEqual(before, [hashlib.sha256(Path(p).read_bytes()).hexdigest() for p in self.files])

    def test_units_missing_placeholder_or_stale_cannot_generate(self):
        preview = run(self.dest, self.files); self.assertFalse(list(self.dest.rglob('*.xlsx')))
        for field, value in [('business_unit', '待确认'), ('source_bundle_hash', 'stale'), ('confirm', False)]:
            units = confirmation(preview); units[field] = value
            self.assertEqual(run(self.dest, self.files, units_confirmation=units)['status'], 'needs_input')

    def test_missing_source_duplicate_identity_and_no_code_are_reported(self):
        self.assertEqual(run(self.dest, self.files[:1])['status'], 'needs_input')
        self.assertEqual(run(self.dest, self.files + self.files[1:])['status'], 'needs_input')
        self.change(self.files[1], '图片处理', 'B2', None)
        self.assertEqual(run(self.dest, self.files)['status'], 'needs_input')

    def test_unknown_formula_missing_cache_and_conflicting_material_are_reported(self):
        self.change(self.files[0], 'Sheet1', 'L2', '=SUM(J2:K2)')
        self.assertIn('白名单', run(self.dest, self.files)['reason'])
        fixtures(self.root, formula=True)
        self.change(self.files[0], 'Sheet1', 'L2', '=J2*K2')
        self.assertIn('缓存', run(self.dest, self.files)['reason'])
        fixtures(self.root, formula=True)
        self.change(self.files[1], '图片处理', 'M2', 'METAL')
        self.assertIn('冲突', run(self.dest, self.files)['reason'])

    def test_provenance_tamper_and_source_change_cannot_replay(self):
        p = run(self.dest, self.files); units = confirmation(p)
        r = run(self.dest, self.files, units_confirmation=units)
        with Path(r['provenance']).open('a', encoding='utf-8') as stream: stream.write(' ')
        self.assertEqual(run(self.dest, self.files, units_confirmation=units)['status'], 'needs_input')

    def test_primary_priced_customs_file_is_not_step4_input(self):
        self.change(self.files[1], '图片处理', 'G2', 50)
        self.assertIn('历史对照', run(self.dest, self.files)['reason'])

    def test_wrong_trailing_amount_total_cannot_hide_a_bad_row(self):
        fixtures(self.root, formula=False)
        self.change(self.files[0], 'Sheet1', 'L3', 999)
        self.assertIn('合计', run(self.dest, self.files)['reason'])

    def test_saved_batch_binds_original_sources(self):
        preview = run(self.dest, self.files)
        r = run(self.dest, self.files, units_confirmation=confirmation(preview))
        from arrival_core.storage import check_sources
        batch_file = Path(r['standard_input']).parent.parent / 'batch.json'
        batch = json.loads(batch_file.read_text(encoding='utf-8'))
        self.change(self.files[1], '图片处理', 'F2', 4)
        with self.assertRaisesRegex(Rejected, '源文件已变更'): check_sources(batch)

    def test_portable_workbuddy_cli_preview(self):
        import subprocess
        request = self.root / 'request.json'
        request.write_text(json.dumps({'message': '处理已确认运营表', 'files': self.files}), encoding='utf-8')
        process = subprocess.run([sys.executable, '-X', 'utf8', '-B', str(ROOT / 'workbuddy/japan-arrival-customs/scripts/run.py'),
            '--output', str(self.dest), 'workbuddy', '--request', str(request)], capture_output=True, text=True, encoding='utf-8')
        self.assertEqual(process.returncode, 0, process.stderr + process.stdout)
        self.assertEqual(json.loads(process.stdout)['status'], 'needs_confirmation')


if __name__ == '__main__': unittest.main()
