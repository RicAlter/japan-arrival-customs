"""rc.7 P0 regressions: stable Excel recovery and scoped per-row amount approval."""
import hashlib
import json
import os
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "workbuddy" / "japan-arrival-customs" / "scripts"))

from arrival_core.business_amount import effective_basis
from arrival_core.compact_review import export_compact
from arrival_core.core import Rejected, digest
from arrival_core.excel_return import ReturnService
from arrival_core.storage import Store, atomic_json
try:  # direct unittest discovery places tests/ on sys.path
    from test_report_templates import TemplateTests
    from test_v070_workflow import WorkflowTests
except ModuleNotFoundError:  # module-qualified local invocation
    from tests.test_report_templates import TemplateTests
    from tests.test_v070_workflow import WorkflowTests


class ReturnRecoveryTests(TemplateTests):
    """No filename, row number or code-only association is accepted on recovery."""

    def _registered(self, name="review.xlsx"):
        state = self.root / "state"
        store = Store(state)
        if not (store.folder(self.batch["id"]) / "head.json").exists():
            store.save(self.batch)
        seed = self.root / ("seed-" + name)
        export_compact(self.batch, None, seed)
        if not hasattr(self, "_delivery_tmp"):
            self._delivery_tmp = tempfile.TemporaryDirectory()
            self.addCleanup(self._delivery_tmp.cleanup)
        service = ReturnService(state)
        destination = Path(self._delivery_tmp.name) / name
        result = service.register_b1(self.batch["id"], seed, destination)
        return service, result, destination

    @staticmethod
    def _drop_identifier_and_reorder(path, target):
        wb = load_workbook(path)
        wb.properties.identifier = None
        # Leave a physical blank row, then reverse the physical order of both
        # stable-key tables.  Values stay with their problem ID / ordinal.
        for title in ("待处理问题", "完整商品明细"):
            ws = wb[title]
            rows = [[cell.value for cell in row] for row in ws.iter_rows(min_row=7, max_row=ws.max_row)]
            rows = [row for row in rows if any(value is not None for value in row)]
            for row in ws.iter_rows(min_row=7, max_row=ws.max_row):
                for cell in row:
                    cell.value = None
            for index, values in enumerate(reversed(rows), 8):
                for column, value in enumerate(values, 1):
                    ws.cell(index, column).value = value
        wb.save(target)
        wb.close()

    def test_missing_identifier_recovers_after_sort_and_blank_row(self):
        service, registered, source = self._registered()
        uploaded = self.root / "save-as-sorted.xlsx"
        self._drop_identifier_and_reorder(source, uploaded)
        preview = service.preview(uploaded, batch_id=self.batch["id"])
        self.assertEqual(preview["export_id"], registered["export_id"])
        self.assertEqual(preview["binding"]["method"], "trusted_readonly_content_match")
        self.assertEqual(preview["binding"]["batch_id"], self.batch["id"])
        self.assertEqual(preview["edits"], [])

    def test_missing_identifier_is_rejected_when_two_exports_match(self):
        service, _, source = self._registered("first.xlsx")
        self._registered("second.xlsx")
        uploaded = self.root / "save-as.xlsx"
        self._drop_identifier_and_reorder(source, uploaded)
        with self.assertRaisesRegex(Rejected, "多个已登记导出"):
            service.preview(uploaded, batch_id=self.batch["id"])

    def test_recovery_rejects_readonly_change_even_with_correct_batch(self):
        service, _, source = self._registered()
        uploaded = self.root / "tampered.xlsx"
        wb = load_workbook(source)
        wb.properties.identifier = None
        wb["完整商品明细"]["D7"] = "altered-readonly-code"
        wb.save(uploaded)
        wb.close()
        with self.assertRaisesRegex(Rejected, "没有匹配的可信导出"):
            service.preview(uploaded, batch_id=self.batch["id"])


class AmountApprovalTests(WorkflowTests):
    def _items(self, *, amount="21", quantity="2", unit="件"):
        row = Store(self.output).load(self.batch_id)["rows"][0]
        return {row["id"]: {"amount": amount, "currency": "JPY", "quantity": quantity,
                            "quantity_unit": unit, "reason": "TEST 本票逐行金额依据"}}

    def _preview(self, **kwargs):
        return __import__("arrival_core.workflow", fromlist=["execute"]).execute(
            self.output, {"message": "预览这批最终金额", "batch": self.batch_id,
                          "amounts": self._items(**kwargs), "reference": "TEST 业务金额来源"})["result"]

    def test_amount_is_previewed_then_confirmed_per_row_and_idempotent(self):
        self.configure_synthetic()
        before = Store(self.output).load(self.batch_id)
        preview = self._preview()
        self.assertEqual(digest(before), digest(Store(self.output).load(self.batch_id)))
        self.assertEqual(preview["changes"][0]["original"]["amount"], "20")
        request = {"message": "确认这批最终金额", "batch": self.batch_id,
                   "preview_id": preview["preview_id"], "confirmed_hash": preview["confirmation_hash"],
                   "confirm": True, "reference": "TEST 明确确认", "session_token_env": "TECH_SESSION"}
        with patch.dict(os.environ, {"TECH_SESSION": "TECH-ONLY"}):
            saved = __import__("arrival_core.workflow", fromlist=["execute"]).execute(self.output, request)["result"]
            replay = __import__("arrival_core.workflow", fromlist=["execute"]).execute(self.output, request)["result"]
        self.assertFalse(saved["formal_output"])
        self.assertFalse(saved["d1_learning_eligible"])
        self.assertTrue(replay["replayed"])
        row = Store(self.output).load(self.batch_id)["rows"][0]
        self.assertEqual(row["authority"]["amount"], "20")
        self.assertEqual(effective_basis(row)["amount"], "21")
        self.assertEqual(row["business_amount"]["current_version"], 1)

    def test_amount_needs_amount_role_and_never_allocates_a_total(self):
        self.configure_synthetic()
        preview = self._preview()
        security = self.output / "security" / "trusted_sessions.json"
        config = json.loads(security.read_text(encoding="utf-8"))
        config["sessions"].append({"principal": "NO_AMOUNT", "roles": ["business"],
            "mode": "synthetic_technical_test", "credential_ref": "TEST", 
            "token_sha256": hashlib.sha256(b"NO-AMOUNT").hexdigest(),
            "expires_at": config["sessions"][0]["expires_at"]})
        atomic_json(security, config)
        before = digest(Store(self.output).load(self.batch_id))
        request = {"message": "确认这批最终金额", "batch": self.batch_id,
                   "preview_id": preview["preview_id"], "confirmed_hash": preview["confirmation_hash"],
                   "confirm": True, "reference": "TEST", "session_token_env": "NO_AMOUNT_SESSION"}
        with patch.dict(os.environ, {"NO_AMOUNT_SESSION": "NO-AMOUNT"}):
            with self.assertRaisesRegex(Rejected, "amount"):
                __import__("arrival_core.workflow", fromlist=["execute"]).execute(self.output, request)
        self.assertEqual(before, digest(Store(self.output).load(self.batch_id)))
        with self.assertRaisesRegex(Rejected, "逐行"):
            __import__("arrival_core.workflow", fromlist=["execute"]).execute(
                self.output, {"message": "预览这批最终金额", "batch": self.batch_id,
                              "amounts": {}, "reference": "TEST total only"})


def load_tests(loader, tests, pattern):
    names = ["test_missing_identifier_recovers_after_sort_and_blank_row",
             "test_missing_identifier_is_rejected_when_two_exports_match",
             "test_recovery_rejects_readonly_change_even_with_correct_batch"]
    suite = unittest.TestSuite(ReturnRecoveryTests(name) for name in names)
    suite.addTests(AmountApprovalTests(name) for name in (
        "test_amount_is_previewed_then_confirmed_per_row_and_idempotent",
        "test_amount_needs_amount_role_and_never_allocates_a_total"))
    return suite


if __name__ == "__main__":
    unittest.main()
