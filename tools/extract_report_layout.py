"""Development only: extract approved presentation, never workbook data."""
import argparse
import hashlib
import json
from pathlib import Path
from xml.etree.ElementTree import tostring

from openpyxl import load_workbook


def style(cell):
    return {key: tostring(getattr(cell, key).to_tree(), encoding="unicode")
            for key in ("font", "fill", "border", "alignment", "protection")} | {"number_format": cell.number_format}


def extract(path, sheet, columns):
    wb = load_workbook(path)
    ws = wb[sheet]
    result = {"sheet": sheet, "columns": columns, "source_sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
              "headers": [ws.cell(1, c).value for c in range(1, columns + 1)],
              "header_styles": [style(ws.cell(1, c)) for c in range(1, columns + 1)],
              "data_styles": [style(ws.cell(2, c)) for c in range(1, columns + 1)],
              "column_dimensions": {key: {field: getattr(dim, field) for field in ("width", "hidden", "min", "max")}
                                    for key, dim in ws.column_dimensions.items()},
              "header_height": ws.row_dimensions[1].height, "data_height": ws.row_dimensions[2].height,
              "freeze_panes": ws.freeze_panes, "show_grid_lines": ws.sheet_view.showGridLines}
    wb.close()
    return result


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("handoff", type=Path)
    parser.add_argument("destination", type=Path)
    args = parser.parse_args()
    files = sorted((args.handoff / "01_已验收版式").glob("*.xlsx"))
    five = next(p for p in files if p.name.startswith("第5步"))
    six = next(p for p in files if p.name.startswith("第6步"))
    data = {"schema": "japan-arrival/accepted-layout-1", "customs": extract(five, "图片处理", 33),
            "internal": extract(six, "Sheet1", 20), "weights": extract(six, "重さとサイズ (2)", 8)}
    args.destination.parent.mkdir(parents=True, exist_ok=True)
    args.destination.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"output": str(args.destination), "sections": list(data), "business_rows_copied": 0}, ensure_ascii=False))


if __name__ == "__main__":
    main()
