"""Read-only KMT rc.7 draft check; creates only a new isolated test result folder."""
import argparse
import hashlib
import json
import sys
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'workbuddy/japan-arrival-customs/scripts'))
from arrival_core.core import digest
from arrival_core.storage import Store,atomic_json
from arrival_core.workflow import execute

def sha(path): return hashlib.sha256(path.read_bytes()).hexdigest()

def main():
    parser=argparse.ArgumentParser();parser.add_argument('raw',type=Path);parser.add_argument('output',type=Path);args=parser.parse_args()
    raw=args.raw.resolve(strict=True);out=args.output.resolve()
    if out.exists(): raise RuntimeError('拒绝覆盖已有KMT验证目录')
    before=sha(raw);out.mkdir(parents=True)
    imported=execute(out/'state',{'message':'处理这批日本到货','files':[str(raw)]})
    batch_id=imported['batch']['batch_id'];store=Store(out/'state');batch=store.load(batch_id);state_hash=digest(batch)
    rendered=execute(out/'state',{'message':'生成海关和内部到货表','batch':batch_id})['result']
    folder=Path(rendered['directory']);customs=load_workbook(folder/'海关报关单_受控测试草稿.xlsx',data_only=True)
    internal=load_workbook(folder/'内部到货单_受控测试草稿.xlsx',data_only=True)
    c=customs['图片处理'];i=internal['Sheet1']
    assert len(batch['rows'])==2687
    assert Decimal(batch['arrival_summary']['subtotal_jpy'])==Decimal('87473076')
    assert Decimal(batch['arrival_summary']['qty_numeric_sum'])==Decimal('79924')
    # The internal template reserves its final notes/footer row; validate the
    # 2,687 data rows rather than treating that layout row as a product.
    assert c.max_column==33 and i.max_column==20 and c.max_row==2688 and i.max_row>=2688
    assert [str(c.cell(r,2).value) for r in range(2,2689)]==[row['identity']['code'] for row in batch['rows']]
    assert sum(Decimal(str(i.cell(r,9).value)) for r in range(2,2689))==Decimal('79924')
    assert sum(c.cell(r,17).value or 0 for r in range(2,2689))==397
    assert sum(Decimal(str(c.cell(r,19).value or 0)) for r in range(2,2689))==Decimal('2858.813')
    assert all(c.cell(r,6).value is None and c.cell(r,7).value is None and c.cell(r,16).value is None for r in range(2,2689))
    assert all(i.cell(r,15).value is None and i.cell(r,18).value is None for r in range(2,2689))
    assert digest(store.load(batch_id))==state_hash and sha(raw)==before
    customs.close();internal.close()
    receipt={'batch_id':batch_id,'source_rows':2687,'raw_qty_sum':'79924','subtotal_jpy':'87473076','cartons':397,
             'weight_kg':'2858.813','customs_columns':33,'internal_columns':20,'source_sha256':before,
             'source_unchanged':True,'batch_unchanged_by_render':True,'formal_output':False,
             'final_business_amount':'not_confirmed','workbuddy_live_test':'not_run'}
    atomic_json(out/'verification.json',receipt);print(json.dumps(receipt,ensure_ascii=False))

if __name__=='__main__': main()
