"""Developer-only read-only-source KMT check; calls the single Skill entrypoint."""
import argparse
import hashlib
import json
import sys
from collections import Counter
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT/'workbuddy/japan-arrival-customs/scripts'))
from arrival_core.workflow import execute
from arrival_core.storage import Store, atomic_json
from arrival_core.core import digest
from arrival_core.excel_return import ReturnService


def sha(path): return hashlib.sha256(path.read_bytes()).hexdigest()


def main():
    p=argparse.ArgumentParser();p.add_argument('handoff',type=Path);p.add_argument('output',type=Path)
    args=p.parse_args(); handoff=args.handoff; out=args.output.resolve(); out.mkdir(parents=True,exist_ok=False)
    manifest={}
    for line in (handoff/'MANIFEST.sha256').read_text(encoding='utf-8-sig').splitlines():
        h, name=line.split('  ',1);manifest[name]=h
    mismatches=[name for name,h in manifest.items() if sha(handoff/name)!=h]
    if mismatches: raise RuntimeError('交接包哈希不一致:'+repr(mismatches))
    raw=next((handoff/'05_原始输入副本').glob('日方原表*.xlsx'))
    state=out/'state'
    imported=execute(state,{'message':'处理这批日本到货','files':[str(raw)],'currency':'JPY'})
    atomic_json(out/'01-import.json',imported)
    batch_id=imported['batch']['batch_id']; store=Store(state); batch=store.load(batch_id)
    original_hash=digest(batch)
    source=next((handoff/'04_依赖数据').glob('候选资料来源*.json'))
    text=source.read_text(encoding='utf-8-sig'); changes=json.loads(text[text.find('{'):])['result']['changes']
    product={}; material={}
    for change in changes:
        for rid in change['row_ids']:
            if change['action']=='edit_product': product.setdefault(rid,{}).update(change['values'])
            elif change['action']=='edit_decision': material.setdefault(rid,{}).update(change['values'])
    regfile=next(path for path in (state/'exports').glob('*.json')
                 if json.loads(path.read_text(encoding='utf-8'))['batch_id']==batch_id)
    reg=json.loads(regfile.read_text(encoding='utf-8'))
    wb=load_workbook(reg['base_file'])
    issue_cells={str(row[0].value).split('｜',1)[0]:row for row in wb['待处理问题'].iter_rows(min_row=7)}
    aliases={'internal_name':'中文名','customs_name':'申报品名','category':'品类','channel':'渠道','yd':'YD',
             'declaration_material':'申报材质'}
    populated=Counter(); skipped=[]
    for issue in reg['issues']:
        values_by_id=product if issue['domain']=='product' else material if issue['domain']=='material' else {}
        records=[values_by_id.get(rid,{}) for rid in issue['row_ids']]
        if not records or not all(records):continue
        fields={k:v for k,v in records[0].items() if k in aliases}
        if not fields or any({k:v for k,v in item.items() if k in aliases}!=fields for item in records):
            skipped.append(issue['id']);continue
        row=issue_cells[issue['id']]
        row[5].value='补充资料'
        row[6].value='范围=本组全部；依据=WorkBuddy交接候选，仅复核版式；'+'；'.join(aliases[k]+'='+v for k,v in fields.items() if v)
        populated[issue['domain']]+=len(issue['row_ids'])
    returned=out/'candidate-review.xlsx';wb.save(returned);wb.close()
    preview=execute(state,{'message':'这是我修改后的审核表，继续处理','file':str(returned)})
    atomic_json(out/'02-preview.json',preview)
    result=execute(state,{'message':'生成海关和内部到货表','batch':batch_id,'preview_id':preview['result']['preview_id']})
    atomic_json(out/'03-render.json',result)
    repeat=execute(state,{'message':'生成海关和内部到货表','batch':batch_id,'preview_id':preview['result']['preview_id']})
    assert result['result']['directory']==repeat['result']['directory']
    folder=Path(result['result']['directory'])
    customs=load_workbook(folder/'海关报关单_受控测试草稿.xlsx',data_only=True)
    internal=load_workbook(folder/'内部到货单_受控测试草稿.xlsx',data_only=True)
    c=customs['图片处理']; i=internal['Sheet1']
    assert len(batch['rows'])==2687
    assert Decimal(batch['arrival_summary']['subtotal_jpy'])==Decimal('87473076')
    assert sum(Decimal(str(i.cell(er,9).value)) for er in range(2,2689))==Decimal('79924')
    assert sum(c.cell(er,17).value or 0 for er in range(2,2689))==397
    assert sum(Decimal(str(c.cell(er,19).value or 0)) for er in range(2,2689))==Decimal('2858.813')
    assert [str(c.cell(er,2).value) for er in range(2,2689)]==[row['identity']['code'] for row in batch['rows']]
    assert all(c.cell(er,6).value is None and c.cell(er,7).value is None and c.cell(er,16).value is None for er in range(2,2689))
    assert all(i.cell(er,15).value is None and i.cell(er,18).value is None for er in range(2,2689))
    assert sum(bool(i.cell(er,19).value) for er in range(2,2689))==2665
    assert digest(store.load(batch_id))==original_hash
    assert not store.load(batch_id).get('decision_events')
    final_mismatches=[name for name,h in manifest.items() if sha(handoff/name)!=h]
    assert not final_mismatches
    report={'batch_id':batch_id,'source_rows':2687,'subtotal_jpy':'87473076','raw_qty_sum':'79924',
            'cartons':397,'weight_kg':'2858.813','candidate_prefill':dict(populated),'skipped_nonuniform_groups':skipped,
            'yd_nonempty':2665,'yd_empty':22,'batch_unchanged':True,'human_final_written':0,'preview_committed':False,
            'handoff_files_verified':len(manifest),'source_files_unchanged':True,'idempotent_render':True,
            'formal_output':False,'workbuddy_live_test':'not_run','output':str(folder),
            'comparison':result['result']['comparison']}
    customs.close();internal.close()
    atomic_json(out/'verification.json',report)
    print(json.dumps(report,ensure_ascii=False,indent=2))


if __name__=='__main__': main()
