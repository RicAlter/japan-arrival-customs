"""KMT operations mapping, source binding and controlled preservation.

Reuse the frozen arithmetic/import/formal gate. Presentation preservation is
implemented only in this independent successor; frozen releases are untouched.
"""
import copy
import hashlib
import json
import re
from decimal import Decimal, localcontext
from fractions import Fraction
from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter, column_index_from_string
from . import _p0_frozen as p0
from .core import Rejected, digest, decimal
from .storage import atomic_json, check_sources
from .excel_writer import literal
from .confirmed_excel import DRAFT_DETAIL_HEADERS, import_confirmed_arrival, gate_status
from .drafts import render_pair, draft_payload, verify_pair

MAPPINGS = Path(__file__).resolve().parents[2] / 'mappings'
UNITS = {'件', '个', '盒', '套', '包', '张', '枚', '支', '只', '本'}


def sha(path): return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def finite(value):
    value = Fraction(value)
    with localcontext() as ctx:
        ctx.prec = 100
        result = Decimal(value.numerator) / Decimal(value.denominator)
    if Fraction(result) != value: raise Rejected('派生数量无法精确表示，需补充明确换算')
    return str(result)


def configs():
    result = []
    for name in ('kmt-stage3-v1.json', 'kmt-stage4-v1.json', 'kmt-weights-v1.json', 'kmt-operations-hij-v1.json', 'kmt-operations-ijk-v1.json'):
        path = MAPPINGS / name
        config = json.loads(path.read_text(encoding='utf-8'))
        config['source'] = {'path': str(path.resolve()), 'sha256': sha(path), 'kind': 'mapping'}
        result.append(config)
    return result


def locate(sheet, config):
    observed = {}
    for col, record in sheet.rows.get(1, {}).items():
        text = p0.text(record.get('v')).strip()
        if text: observed.setdefault(text, []).append(col)
    columns = {}
    for field, header in config['headers'].items():
        found = observed.get(header, [])
        if len(found) != 1: return None
        columns[field] = found[0]
    for field in config['anchored_blank_headers']:
        # An unnamed column can only be mapped when all named anchors stay fixed.
        if any(columns[f] != config['columns'][f] for f in columns): return None
        col = config['columns'][field]
        if p0.good(sheet.v(1, col)): return None
        columns[field] = col
    return columns


def evidence(book, sheet, row, col):
    cell = sheet.cells.get(f'{col}{row}', {})
    return {'path': book.label, 'sha256': book.sha256, 'sheet': sheet.name, 'row': row,
            'cell': f'{col}{row}', 'raw_value': p0.text(cell.get('v')), 'formula': cell.get('f'),
            'cached_value': p0.text(cell.get('v')), 'cell_type': cell.get('type'),
            'literal_value': original_literal(cell.get('v'))}


def original_literal(value):
    if isinstance(value, Decimal):
        return int(value) if value == value.to_integral_value() else float(value)
    return value


def value(book, sheet, row, field, columns, config):
    col = columns[field]; proof = evidence(book, sheet, row, col)
    raw = sheet.v(row, col); formula = sheet.f(row, col)
    if formula:
        factors = config['formula_products'].get(field)
        if not factors: raise Rejected(f'{sheet.name}!{col}{row}：不支持的公式')
        refs = [columns[f] + str(row) for f in factors]
        normalized = formula.replace('$', '').replace(' ', '').upper()
        if normalized not in ('=' + '*'.join(refs), '=' + '*'.join(reversed(refs))):
            raise Rejected(f'{sheet.name}!{col}{row}：公式超出逐行乘法白名单')
        operands = []
        for ref in refs:
            c = sheet.cells.get(ref, {})
            if c.get('f') or c.get('type') == 'e': raise Rejected('公式引用值不是独立数值')
            operands.append(decimal(p0.text(c.get('v'))))
        expected = Fraction(operands[0]) * Fraction(operands[1])
        if raw is None or Fraction(decimal(p0.text(raw))) != expected:
            raise Rejected(f'{sheet.name}!{col}{row}：公式缓存缺失或与精确复算不一致')
        proof.update(validation='exact_product_matched', precedents=[evidence(book, sheet, row, columns[f]) for f in factors])
    elif not p0.good(raw):
        if field not in config['optional']: raise Rejected(f'{sheet.name}!{col}{row}：缺少必要字段{field}')
        return '', proof
    return p0.text(raw).strip(), proof


def extract(book, sheet, config, columns, no_code_reasons):
    output = []; previous_carton = None; total_seen = False
    for row in sorted(sheet.rows):
        if row < config['start_row']: continue
        cells = sheet.rows[row]
        for header in config.get('empty_data_headers', []):
            cols = [c for c, record in sheet.rows.get(1, {}).items() if p0.text(record.get('v')) == header]
            if len(cols) > 1: raise Rejected('价格列存在重复表头')
            if cols and (p0.present(sheet.v(row, cols[0])) or sheet.f(row, cols[0])):
                raise Rejected('首期只接第4步申报准备表；已填价格成品请作为历史对照提供')
        marker = config.get('footer_marker')
        if marker and p0.text(sheet.v(row, marker['column'])) == marker['value']:
            for field in ('实际箱数', '箱组总重量KG'):
                col = columns[field]
                if sheet.f(row, col) != f'=SUM({col}{config["start_row"]}:{col}{row-1})':
                    raise Rejected('箱重合计公式范围不明确')
                total = sum((Fraction(decimal(x['values'][field])) for x in output), Fraction(0))
                if Fraction(decimal(p0.text(sheet.v(row, col)))) != total: raise Rejected('箱重合计缓存不一致')
            total_seen = True; continue
        if total_seen and config['role'] == 'weights':
            if p0.expand(sheet.v(row, columns['箱号组'])): raise Rejected('箱重合计后出现实体箱')
            continue
        # Only a trailing isolated amount total is excluded, never an unnamed item.
        total_field = config.get('trailing_total_field')
        if total_field and set(cells) == {columns[total_field]}:
            col = columns[total_field]
            formula = sheet.f(row, col)
            if formula and formula != f'=SUM({col}{config["start_row"]}:{col}{row-1})':
                raise Rejected('业务金额合计公式范围不明确')
            expected = sum((Fraction(decimal(x['values'][total_field])) for x in output), Fraction(0))
            if Fraction(decimal(p0.text(sheet.v(row, col)))) != expected:
                raise Rejected('业务金额尾部合计与逐行金额不一致')
            total_seen = True; continue
        if total_seen: raise Rejected('合计后出现商品或未知内容')
        if config['role'] == 'weights':
            carton = sheet.v(row, columns['箱号组'])
            if not p0.good(carton):
                if any(p0.good(sheet.v(row, columns[f])) for f in ('实际箱数', '箱组总重量KG')):
                    raise Rejected(f'箱重第{row}行缺少箱号')
                continue
            if not p0.expand(carton):
                # Total/volume rows have no carton identity. Reject numeric count/weight.
                if p0.numeric_text(sheet.v(row, columns['实际箱数'])) is not None:
                    raise Rejected(f'箱重第{row}行箱号无法识别')
                continue
        values = {}; proofs = {}
        for field in columns:
            if field == '商品编码' and not p0.good(sheet.v(row, columns[field])):
                values[field] = ''; proofs[field] = evidence(book, sheet, row, columns[field]); continue
            if field == '箱号' and not p0.good(sheet.v(row, columns[field])) and config['inherit_carton']:
                if previous_carton is None: raise Rejected('首行缺少箱号，不能继承')
                values[field], proofs[field] = copy.deepcopy(previous_carton)
                proofs[field]['inherited_to_row'] = row; continue
            values[field], proofs[field] = value(book, sheet, row, field, columns, config)
        if '箱号' in values: previous_carton = (values['箱号'], proofs['箱号'])
        rid = digest([book.sha256, sheet.name, row])[:24]
        if '商品编码' in values and not values['商品编码']:
            reason = no_code_reasons.get(rid, '')
            if not isinstance(reason, str) or not reason.strip(): raise Rejected(f'无码行{rid}需补充真实原因及唯一配对资料')
            values['无码原因'] = reason.strip()
        values['稳定行号'] = rid
        for name, col in config.get('observed_dimensions', {}).items():
            if p0.present(sheet.v(row, col)):
                if sheet.f(row, col): raise Rejected('尺寸只接受原表独立观测数值')
                decimal(p0.text(sheet.v(row, col)))
                values[name] = p0.text(sheet.v(row, col)); proofs[name] = evidence(book, sheet, row, col)
        preserved = {name: evidence(book, sheet, row, col)
                     for name, col in config.get('preserve_columns', {}).items()}
        if config['role'] == 'declaration':
            preserved.update({name: evidence(book, sheet, row, col) for name, col in
                              (('brand', 'H'), ('blind', 'K'), ('label', 'L'), ('size', 'N'), ('spec', 'O'), ('public_origin', 'D'))})
        tax_input = None
        if config.get('tax_columns'):
            from .colored_operations import read_tax_input
            tax_input = read_tax_input(book, sheet, row, config, values)
        output.append({'values': values, 'fields': proofs, 'row': row, 'config': config['id'],
                       'preserved': preserved, 'tax_input': tax_input})
    if not output: raise Rejected(f'{sheet.name}没有明细')
    return output


def carton_key(raw):
    expanded = p0.expand(raw)
    if not expanded: raise Rejected('箱号不能识别：' + raw)
    return tuple(sorted(expanded))


def keyed(rows):
    result = {}
    for row in rows:
        v = row['values']; key = (carton_key(v['箱号']), v['商品编码'])
        if not key[1]: raise Rejected('无码商品无法按箱号＋编码唯一配对，需单独提供关联依据')
        if key in result: raise Rejected('箱号＋编码存在重复候选，禁止按行序配对')
        result[key] = row
    return result


def pair_in_source_order(business, declaration):
    """Unique lookup never controls output order; duplicate proposals require explicit confirmation."""
    def key(row):
        v = row['values']
        if not v['商品编码']: raise Rejected('无码商品需单独提供关联依据，不能猜测配对')
        return carton_key(v['箱号']), v['商品编码']
    groups = {}
    for row in declaration: groups.setdefault(key(row), []).append(row)
    business_groups = {}
    for row in business: business_groups.setdefault(key(row), []).append(row)
    if set(groups) != set(business_groups): raise Rejected('第3/4步商品身份不能双向完整匹配')
    pairs = {}; ambiguous = []
    for identity, originals in business_groups.items():
        candidates = groups[identity]
        if len(originals) != len(candidates): raise Rejected('重复身份两侧行数不同，不能配对')
        if len(originals) == 1:
            pairs[originals[0]['values']['稳定行号']] = candidates[0]
        else:
            if [r['row'] for r in originals] != [r['row'] for r in candidates]:
                raise Rejected('重复身份缺少同原行对应依据，禁止自动配对')
            for a, b in zip(originals, candidates):
                pairs[a['values']['稳定行号']] = b
                ambiguous.append({'row_id': a['values']['稳定行号'], 'business_row': a['row'],
                                  'declaration_row': b['row'], 'code': a['values']['商品编码'],
                                  'specification': a['preserved'].get('specification', {}).get('literal_value'),
                                  'basis': '同原行对应候选，待本批明确确认'})
    return [(a, pairs[a['values']['稳定行号']]) for a in business], ambiguous


def prepare(files, comparison_files=(), no_code_reasons=None):
    config_list = configs(); sources = []; matched = {}; books = []
    preloaded = []
    for name in files:
        path = Path(name).resolve(strict=True)
        if path.suffix.lower() != '.xlsx': raise Rejected('只接受xlsx运营原件')
        preloaded.append(p0.Book(path.read_bytes(), str(path)))
    completed = any(locate(s, config_list[-1]) for b in preloaded for s in b.sheets)
    if completed:
        config_list[1] = copy.deepcopy(config_list[1])
        config_list[1]['empty_data_headers'] = []
        config_list[1]['headers'].update({'执行申报单价': '單價', '执行申报金额': '金額'})
        config_list[1]['columns'].update({'执行申报单价': 'G', '执行申报金额': 'P'})
        config_list[1]['optional'] += ['执行申报单价', '执行申报金额']
        config_list[1]['formula_products']['执行申报金额'] = ['执行申报单价', '申报数量']
    for name, book in zip(files, preloaded):
        path = Path(name).resolve(strict=True)
        if path.suffix.lower() != '.xlsx': raise Rejected('只接受xlsx运营原件')
        books.append(book)
        sources.append({'path': str(path), 'sha256': book.sha256, 'kind': 'operations_original'})
        used = False
        for config in config_list:
            for sheet in book.sheets:
                if config.get('sheet') and sheet.name != config['sheet']: continue
                columns = locate(sheet, config)
                if not columns: continue
                role = config['role']
                if role in matched: raise Rejected('同角色出现多个工作簿/版本：' + role)
                matched[role] = extract(book, sheet, config, columns, no_code_reasons or {})
                used = True
        if not used: raise Rejected('无法识别运营文件结构：' + path.name)
    missing = sorted({'business', 'declaration', 'weights'} - set(matched))
    if missing: raise Rejected('缺少必要来源：' + '、'.join(missing))
    paired, ambiguous = pair_in_source_order(matched['business'], matched['declaration'])
    business = {a['values']['稳定行号']: a for a, _ in paired}
    details = []; special = []
    origins = {'Japan': 'JP', 'China': 'CN', 'Vietnam': 'VN', '日本': 'JP', '中国': 'CN', '越南': 'VN'}
    for a, b in paired:
        av, bv = a['values'], b['values']; fields = copy.deepcopy(a['fields'])
        for f in ('品类', '最终申报材质'):
            if av[f] != bv[f]: raise Rejected(f'{av["稳定行号"]}第3/4步{f}冲突')
        if origins.get(bv['原产国'], bv['原产国']) != av['原产国']: raise Rejected('原产国冲突')
        if '来源盒数' in av and Fraction(decimal(av['来源盒数'])) != Fraction(decimal(bv['内部盒数'])):
            raise Rejected('内部来源盒数与海关盒装数量不一致')
        v = dict(av)
        for f in ('海关申报品名', '内部盒数', '申报数量'):
            v[f] = bv[f]; fields[f] = b['fields'][f]
        for f in ('业务确认数量', '内部盒数', '申报数量'):
            if decimal(v[f]) <= 0: raise Rejected(f'{v["稳定行号"]}{f}必须为正数')
        qty, boxes, declared = [Fraction(decimal(v[f])) for f in ('业务确认数量', '内部盒数', '申报数量')]
        if decimal(v['业务确认金额JPY']) < 0: raise Rejected('金额不能为负数')
        if Fraction(decimal(av['业务单价'])) * qty != Fraction(decimal(v['业务确认金额JPY'])):
            raise Rejected('业务单价×数量与第3步金额不一致')
        v.update(内部数量=v['业务确认数量'], 每盒件数=finite(qty / boxes), 申报换算系数=finite(declared / qty), 无码原因=v.get('无码原因', ''))
        for f, expression, inputs in (('内部数量', '业务确认数量', ['业务确认数量']),
                ('每盒件数', '业务确认数量/内部盒数', ['业务确认数量', '内部盒数']),
                ('申报换算系数', '申报数量/业务确认数量', ['申报数量', '业务确认数量'])):
            fields[f] = {'derived': expression, 'sources': [fields[x] for x in inputs]}
        if declared != qty:
            special.append({'row_id': v['稳定行号'], 'source_row': a['row'], 'business_quantity': v['业务确认数量'],
                            'carton': v['箱号'], 'code': v['商品编码'], 'internal_name': v['内部中文名'],
                            'declaration_quantity': v['申报数量'], 'factor': v['申报换算系数']})
        details.append({'values': v, 'fields': fields,
                        'source_rows': {'business': a['row'], 'declaration': b['row']},
                        'preserved': dict(a['preserved'], **b['preserved']), 'tax_input': a['tax_input'],
                        'execution_input': {'price': bv.get('执行申报单价'), 'amount': bv.get('执行申报金额'),
                                            'fields': {f: b['fields'][f] for f in ('执行申报单价','执行申报金额') if f in b['fields']}}})
    differences = []
    for name in comparison_files:
        path = Path(name).resolve(strict=True); book = p0.Book(path.read_bytes(), str(path))
        if not book.sheets or not locate(book.sheets[0], config_list[1]):
            raise Rejected('历史金额对照结构不符合KMT海关表')
        if p0.text(book.sheets[0].v(1, 'P')) != '金額': raise Rejected('历史对照缺少明确金额列')
        sources.append({'path': str(path), 'sha256': book.sha256, 'kind': 'comparison_only'})
        # Explicit comparison role; never feeds standard values.
        rows = p0.records(book, 5)
        ref = {}
        for r in rows:
            key = (carton_key(r['carton']), r['code'])
            if key in ref: raise Rejected('历史对照存在重复身份')
            ref[key] = r
        keyed_business = keyed(matched['business'])
        if set(ref) != set(keyed_business): raise Rejected('历史对照不能完整匹配')
        for key, a in keyed_business.items():
            r = ref[key]; amount = decimal(a['values']['业务确认金额JPY'])
            if Fraction(r['amount']) != Fraction(amount):
                differences.append({'row_id': a['values']['稳定行号'], 'business': a['fields']['业务确认金额JPY'],
                    'historical': evidence(book, book.sheets[0], r['row'], 'P'),
                    'delta': str(Fraction(r['amount']) - Fraction(amount)), 'applied': False})
    sources += [config['source'] for config in config_list]
    for s in sources:
        if sha(s['path']) != s['sha256']: raise Rejected('读取期间来源发生变化')
    business_profile = next(c for c in config_list if c['id'] == matched['business'][0]['config'])
    source_book = next(b for b in books if b.sha256 == matched['business'][0]['fields']['箱号']['sha256'])
    source_sheet = next(s for s in source_book.sheets if s.name == business_profile['sheet'])
    snapshot = {'file': source_book.label, 'sha256': source_book.sha256, 'sheet': source_sheet.name,
                'headers': [original_literal(source_sheet.v(1, get_column_letter(c))) for c in range(1, 1 + max(column_index_from_string(c) for c in source_sheet.rows[1]))],
                'rows': [], 'cells': {a: {'value': original_literal(c.get('v')), 'formula': c.get('f'), 'type': c.get('type')}
                                     for a, c in source_sheet.cells.items()}}
    # The internal source appendix includes original blanks/footer totals too;
    # only the calculation details exclude the independently verified total row.
    snapshot['rows'] = [[original_literal(source_sheet.v(n, get_column_letter(c)))
                         for c in range(1, len(snapshot['headers']) + 1)]
                        for n in range(2, max(source_sheet.rows) + 1)]
    weight_origin = matched['weights'][0]['fields']['箱号组']
    wb_weights = next(b for b in books if b.sha256 == weight_origin['sha256'])
    ws_weights = next(s for s in wb_weights.sheets if s.name == weight_origin['sheet'])
    weight_snapshot = {'sheet': ws_weights.name, 'source': weight_origin,
        'rows': [[original_literal(ws_weights.v(n, get_column_letter(c))) for c in range(1,9)]
                 for n in range(1, max(ws_weights.rows)+1)],
        'cells': {a: {'value': original_literal(c.get('v')), 'formula': c.get('f'), 'type': c.get('type')}
                  for a,c in ws_weights.cells.items()}}
    return {'schema': 'operations-kmt-adapter-v1', 'sources': sources, 'details': details,
            'completed_input': completed, 'weight_snapshot': weight_snapshot,
            'business_profile': business_profile['id'], 'source_quantity_scope': business_profile.get('source_quantity_scope'),
            'original_snapshot': snapshot, 'pairing_candidates': ambiguous, 'pairing_hash': digest(ambiguous),
            'weights': matched['weights'], 'special_conversions': special,
            'special_conversion_hash': digest(special), 'historical_amount_differences': differences,
            'source_bundle_hash': digest(sources), 'formal_output': False}


def run(output, files, *, units_confirmation=None, comparison_files=(), no_code_reasons=None, reference_evidence=None, tax_allocations=None):
    output = Path(output).resolve()
    for name in [*files, *comparison_files]:
        path = Path(name).resolve()
        if output.is_relative_to(path.parent) or path.is_relative_to(output):
            raise Rejected('输出必须与原件目录隔离')
    output.mkdir(parents=True, exist_ok=True)
    try:
        plan = prepare(files, comparison_files, no_code_reasons)
        from .operations_preservation import load_reference_evidence
        reference_data = load_reference_evidence(reference_evidence, plan)
        from .colored_operations import prepare_calculation_inputs
        allocation_hash = prepare_calculation_inputs(plan, reference_data, tax_allocations)
        preview = {'status': 'needs_confirmation', 'source_bundle_hash': plan['source_bundle_hash'],
                   'special_conversion_hash': plan['special_conversion_hash'], 'rows': len(plan['details']),
                   'sources': [s for s in plan['sources'] if s['kind'] != 'mapping'],
                   'required_units': ['business_unit', 'internal_unit', 'declaration_unit'],
                   'special_conversions': plan['special_conversions'], 'formal_output': False,
                   'business_profile': plan['business_profile'], 'source_quantity_scope': plan['source_quantity_scope'],
                   'pairing_candidates': plan['pairing_candidates'], 'pairing_hash': plan['pairing_hash'],
                   'reference_evidence_hash': digest(reference_data) if reference_data else None}
        if tax_allocations: preview['tax_allocation_hash'] = allocation_hash
        atomic_json(output / (plan['source_bundle_hash'][:24] + '-preview.json'), preview)
        if units_confirmation is None: return preview
        units = units_confirmation
        required = {'business_unit', 'internal_unit', 'declaration_unit', 'reference', 'confirm',
                    'source_bundle_hash', 'special_conversion_hash'}
        if plan['pairing_candidates']: required.add('pairing_hash')
        if reference_data: required.add('reference_evidence_hash')
        if tax_allocations: required.add('tax_allocation_hash')
        if not isinstance(units, dict) or set(units) != required: raise Rejected('批次单位/对应关系/参考试算确认字段不完整或未知')
        if plan['pairing_candidates'] and units['pairing_hash'] != plan['pairing_hash']: raise Rejected('重复身份对应范围已变化，需重新确认')
        if reference_data and units['reference_evidence_hash'] != digest(reference_data): raise Rejected('参考试算依据已变化，需重新确认')
        if tax_allocations and units['tax_allocation_hash'] != allocation_hash: raise Rejected('分摊依据已变化，需重新确认')
        if units['confirm'] is not True or units['source_bundle_hash'] != plan['source_bundle_hash'] or units['special_conversion_hash'] != plan['special_conversion_hash']:
            raise Rejected('先展示当前来源及特殊换算清单，再明确确认当前哈希')
        allowed_units = UNITS | ({'Qty（原表计价口径）', '单个'} if plan['source_quantity_scope'] else set())
        if any(units[k] not in allowed_units for k in ('business_unit', 'internal_unit', 'declaration_unit')):
            raise Rejected('必须提供明确实际单位，不能使用待确认占位')
        if units['business_unit'] != units['internal_unit']: raise Rejected('首期引擎要求业务数量与内部数量同层级')
        if plan['source_quantity_scope'] and (units['business_unit'] != plan['source_quantity_scope'] or units['declaration_unit'] != '单个'):
            raise Rejected('混合Qty运营表需明确保留原计价口径及原表单个申报口径，不能全列猜单位')
        if not isinstance(units['reference'], str) or not units['reference'].strip(): raise Rejected('单位确认依据必填')
        key = digest({'source': plan['source_bundle_hash'], 'units': units})[:24]
        folder = output / key
        if folder.exists() and not (folder / 'batch.json').exists():
            raise Rejected('发现不完整运行目录；保留现场，请选择新的输出目录重试')
        folder.mkdir(exist_ok=True)
        canonical = folder / 'standard' / 'confirmed-arrival-v1.xlsx'
        provenance = folder / 'evidence' / 'provenance.json'
        if not canonical.exists():
            canonical.parent.mkdir(exist_ok=True); provenance.parent.mkdir(exist_ok=True)
            wb = Workbook(); meta = wb.active; meta.title = '批次确认'
            for k, v in zip(('模板版本', '批次号', '币种', '重量单位', '确认依据编号', '数据版本'),
                           ('confirmed-arrival-v1', key, 'JPY', 'KG', units['reference'], plan['source_bundle_hash'])):
                meta.append([k, v])
            sheet = wb.create_sheet('已确认明细'); headers = list(DRAFT_DETAIL_HEADERS) + ['YD', '渠道', '内部备注']
            headers += ['关税率', '增值税率', '税基单价', '税基单位', '税费公式版本', '关务依据']
            sheet.append(headers)
            for index, row in enumerate(plan['details'], 2):
                row['values'].update(计价单位=units['business_unit'], 申报单位=units['declaration_unit'])
                if row['preserved'].get('yd', {}).get('literal_value') == 0:
                    row['values']['YD'] = ''  # Raw numeric placeholder remains in typed provenance/output.
                if reference_data:
                    ref = reference_data['by_source_row'][str(row['source_rows']['business'])]
                    if ref['duty_rate'] is not None:
                        row['values'].update(关税率=ref['duty_rate'], 增值税率=reference_data['vat_rate'],
                            税基单价=row['values']['业务单价'], 税基单位=units['business_unit'],
                            税费公式版本='historical-box-v1', 关务依据='受控参考试算，非关务批准：' + ref['basis'])
                supplied_tax = row.get('tax_input')
                if supplied_tax and supplied_tax.get('rate') is not None and supplied_tax.get('vat_rate') is not None:
                    row['values'].update(关税率=supplied_tax['rate'], 增值税率=supplied_tax['vat_rate'],
                        税基单价=row['values']['业务单价'], 税基单位=units['business_unit'],
                        税费公式版本='historical-box-v1', 关务依据='本批运营输入，非正式批准：' + supplied_tax['basis'])
                for f in ('计价单位', '申报单位'):
                    row['fields'][f] = {'source': 'explicit_batch_confirmation', 'confirmation': units, 'trusted_identity': False}
                for col, field in enumerate(headers, 1): literal(sheet.cell(index, col), row['values'].get(field, ''))
            weights = wb.create_sheet('箱重'); weight_fields = ['箱号组', '实际箱数', '箱组总重量KG']
            weights.append(weight_fields)
            weight_fields += ['单箱重KG', '长度CM', '宽度CM', '高度CM']
            for c, name in enumerate(weight_fields, 1): weights.cell(1, c, name)
            for row in plan['weights']: weights.append([row['values'].get(f, '') for f in weight_fields])
            wb.save(canonical); wb.close()
            atomic_json(provenance, plan)
        batch_file = folder / 'batch.json'
        if batch_file.exists():
            batch = json.loads(batch_file.read_text(encoding='utf-8')); check_sources(batch)
        else:
            batch = import_confirmed_arrival(canonical)
            batch['sources'] += plan['sources'] + [{'path': str(provenance), 'sha256': sha(provenance), 'kind': 'field_provenance'}]
            batch['operations_adapter'] = {'source_bundle_hash': plan['source_bundle_hash'], 'provenance': str(provenance), 'units': units, 'trusted_identity': False}
            batch['operations_delivery_policy'] = {'version': 'rounded-colored-v1', 'preserve_business_amount': True,
                'formal_output': False, 'original_snapshot': plan['original_snapshot'],
                'weight_snapshot': plan['weight_snapshot']}
            if reference_data:
                batch['sources'] += reference_data['sources']
                batch['reference_trial_notice'] = reference_data['notice']
            for row, original in zip(batch['rows'], plan['details']):
                row['operations_provenance'] = original['fields']
                row['operations_preserved'] = original['preserved']
                row['original_pairing'] = original['source_rows']
                row['operations_tax_reference'] = reference_data['by_source_row'][str(original['source_rows']['business'])] if reference_data else None
                row['operational_execution'] = original['execution_input']
                row['operational_tax_input'] = original['tax_input']
            atomic_json(batch_file, batch)
        pair = folder / 'pair'
        from .colored_tables import render, verify, payload
        if pair.exists(): result = {'directory': str(pair), 'comparison': verify(batch, pair, payload(batch)), 'replayed': True}
        else: result = render(batch, pair)
        check_sources(batch)
        delivery = result['comparison']
        gates = gate_status(batch)
        gates['draft']['price_pending_rows'] = 0
        gates['draft']['rounding_rows'] = delivery['rounding_rows']
        gates['draft']['unexplained_amount_difference_jpy'] = delivery['batch_difference_jpy']
        for item in gates['rows']:
            item['exact_price_notice'] = item['price_pending']
            item['formal_price_review_pending'] = bool(item['price_pending'])
            item['price_pending'] = False
        return {'status': 'ok', 'batch': batch['id'], 'standard_input': str(canonical), 'provenance': str(provenance),
                'result': result, 'gates': gates, 'special_conversion_count': len(plan['special_conversions']),
                'historical_amount_difference_count': len(plan['historical_amount_differences']),
                'delivery_report': delivery, 'business_files': [str(pair/name) for name in ('海关报关单_受控测试草稿.xlsx','内部到货单_受控测试草稿.xlsx')],
                'formal_output': False}
    except (Rejected, OSError, ValueError, KeyError) as exc:
        report = {'status': 'needs_input', 'reason': str(exc), 'formal_output': False}
        atomic_json(output / 'input-problems.json', report)
        return report
