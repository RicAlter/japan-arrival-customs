"""P2-B2: registered B1 workbook -> immutable preview -> explicit trusted CAS commit.

Only two business columns are editable. This is a local adapter, not an online UI.
"""
import copy
import hashlib
import io
import json
import os
import re
import shutil
import tempfile
import zipfile
from pathlib import Path
from uuid import uuid4

from openpyxl import load_workbook

from .compact_review import (FILENAME,SHEETS,CHOICES,build_plan,material_catalog,export_compact,_write)
from .core import (Rejected,digest,now,validate,context,changed,review_package,confirm_products,
                   supplement_many,box_tax,exact_price)
from .storage import Store,atomic_json,check_sources
from .return_auth import authorize,Session
from .declaration_decisions import revise,head,FACT_FIELDS,FINAL_FIELDS,validate_values,selected_material
from .excel_writer import literal,save_exact

CONTRACT='japan-arrival/excel-return-1'
ALIASES={'申报材质':'declaration_material','材质':'declaration_material','税号':'classification_code',
 '归类依据':'classification_basis','税制版本':'tariff_schedule_version','关税率':'duty_rate','增值税率':'vat_rate',
 '生效日':'valid_from','失效日':'valid_through','适用日期':'applicability_date','适用辖区':'jurisdiction',
 '原产地条件':'origin_conditions','贸易条件':'trade_conditions','商品类型':'product_type','用途':'purpose',
 '组成':'complete_composition','关键材质特征':'key_material_features','构造':'construction','版本':'version',
 '规格':'specification','适用条件':'inclusion_conditions','排除条件':'exclusion_conditions','原因':'reason',
 '无YD原因':'product.yd_absence_reason','中文名':'product.internal_name','申报品名':'product.customs_name',
 '品类':'product.category','渠道':'product.channel','YD':'product.yd',
 '最终金额':'business_amount.amount','币种':'business_amount.currency',
 '计价数量':'business_amount.quantity','计价单位':'business_amount.quantity_unit'}
KEYS=set(FACT_FIELDS)|set(FINAL_FIELDS)|{'reason'}|{v for v in ALIASES.values() if v.startswith(('product.','business_amount.'))}


def sha(data): return hashlib.sha256(data).hexdigest()


def safe_id(value,length=64):
    if not isinstance(value,str) or not re.fullmatch('[a-f0-9]{'+str(length)+'}',value):raise Rejected('非法登记标识')
    return value


def read_json(path):return json.loads(Path(path).read_text(encoding='utf-8'))


def _archive(path,data):
    path.parent.mkdir(parents=True,exist_ok=True)
    temp=path.with_name('.'+path.name+'.'+uuid4().hex+'.tmp')
    try:
        with temp.open('xb') as f:f.write(data);f.flush();os.fsync(f.fileno())
        try:os.link(temp,path)  # Atomic exclusive publication, never replace a business file.
        except FileExistsError:
            if path.read_bytes()!=data:raise Rejected('只读归档冲突')
    finally:temp.unlink(missing_ok=True)


def _read_workbook(data):
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            if sum(i.file_size for i in z.infolist())>128*1024*1024:raise Rejected('Excel解压大小超限')
            names=z.namelist()
            if len(names)!=len(set(names)):raise Rejected('Excel包包含重复部件')
            if any(any(t in n.lower() for t in ('vbaproject','externallinks/','embeddings/','connections.xml','querytables/')) for n in names):
                raise Rejected('禁止宏、外部连接或嵌入对象')
            # B1 internal navigation uses fragment relationships; no external targets.
            import xml.etree.ElementTree as ET
            for n in names:
                if not n.endswith('.rels'):continue
                for rel in ET.fromstring(z.read(n)):
                    if rel.get('TargetMode')=='External' and not rel.get('Target','').startswith('#'):
                        raise Rejected('禁止外部活动链接')
        wb=load_workbook(io.BytesIO(data),data_only=False,keep_links=False)
        if wb.sheetnames!=SHEETS:raise Rejected('仅接受B1三页工作簿，不能删除/新增页签')
        for ws in wb:
            if ws.max_row>50000 or ws.max_column>192:raise Rejected('Excel结构/范围异常')
            for row in ws:
                for c in row:
                    if c.data_type=='f':raise Rejected('审核回传禁止可执行公式')
        return wb
    except Rejected:raise
    except (ValueError,KeyError,zipfile.BadZipFile) as exc:raise Rejected('无法读取受控Excel') from exc


def _table(ws):
    rows={}
    for cells in ws.iter_rows(min_row=7):
        value=[c.value for c in cells]
        if all(v is None for v in value):continue
        # Some Excel writers leave an internal navigation anchor on an inserted blank
        # row. It carries no business input; never ignore a row containing actual notes.
        if value[0] is None and all(c.value is None or (c.hyperlink and c.value==c.hyperlink.target and str(c.value).startswith('#')) for c in cells):
            continue
        key=str(value[0]).split('｜',1)[0] if ws.title==SHEETS[1] else str(value[0])
        if value[0] is None or key in rows:raise Rejected('空/重复行标识，不能靠显示行号猜测')
        rows[key]=value
    return rows


def compare_workbooks(base,returned,*,recover_missing_identifier=False):
    if base.properties.identifier!=returned.properties.identifier and not (recover_missing_identifier and not returned.properties.identifier):
        raise Rejected('导出登记标识不匹配')
    if [(k,v.attr_text) for k,v in base.defined_names.items()]!=[(k,v.attr_text) for k,v in returned.defined_names.items()]:
        raise Rejected('命名公式/范围被修改')
    edits=[]
    for name in SHEETS:
        old,new=base[name],returned[name]
        if old.sheet_state!=new.sheet_state:raise Rejected('页签显示状态被修改')
        if old.max_column!=new.max_column:raise Rejected('列结构被修改')
        for r in range(1,7):
            if [c.value for c in old[r]]!=[c.value for c in new[r]]:raise Rejected('表头/只读说明被修改')
        a,b=_table(old),_table(new)
        if set(a)!=set(b):raise Rejected('存在缺失/未知问题或商品行')
        for key,values in a.items():
            for col,(before,after) in enumerate(zip(values,b[key]),1):
                if before==after and (not isinstance(before,bool) and not isinstance(after,bool) or type(before)==type(after)):continue
                if name!=SHEETS[1] or col not in (6,7):raise Rejected('只读字段被修改：'+name+' / '+key)
                if after is not None and not isinstance(after,str):raise Rejected('问题填写必须是文本')
            if name==SHEETS[1] and values[5:7]!=b[key][5:7]:
                choice,note=b[key][5:7]
                if choice and choice not in CHOICES.split(','):raise Rejected('未知处理选择')
                edits.append({'issue_id':key,'before':values[5:7],'after':[choice,note]})
    return sorted(edits,key=lambda e:e['issue_id'])


def _parse_note(note):
    if not note:return {},{},''
    values={};control={};unparsed=[]
    for part in re.split('[；;\n]',note):
        part=part.strip()
        if not part:continue
        if '=' not in part and '＝' not in part:
            unparsed.append(part);continue
        key,value=re.split('[=＝]',part,maxsplit=1);key=key.strip();value=value.strip()
        if key in ('范围','依据'):target=control;field=key
        else:target=values;field=ALIASES.get(key,key)
        if target is values and field not in KEYS:unparsed.append(part);continue
        if field in target:raise Rejected('说明字段重复，需明确而不是取最后一条')
        target[field]=value
    validate_values({k:v for k,v in values.items() if not k.startswith(('product.','business_amount.'))})
    return values,control,'；'.join(unparsed)


def _scope(registry,issue,control):
    requested=control.get('范围')
    if not requested:
        if len(issue['row_ids'])!=1:raise Rejected('成组修改须明确范围（本组全部或本表序号），不能默认覆盖整组')
        return issue['row_ids'][:]
    if requested in ('本组全部','组内全部'):return issue['row_ids'][:]
    inverse={str(n):rid for rid,n in registry['ordinals'].items()}
    ids=[]
    for part in re.split('[,，、 ]+',requested):
        if not part:continue
        if not re.fullmatch('[0-9]+',part) or str(int(part)) not in inverse:raise Rejected('范围不是本导出版的有效商品序号')
        ids.append(inverse[str(int(part))])
    if not ids or len(ids)!=len(set(ids)) or set(ids)-set(issue['row_ids']):raise Rejected('问题范围外或重复商品')
    return ids


class ReturnService:
    def __init__(self,root):
        self.store=Store(root);self.root=self.store.root

    def _path(self,kind,identifier):
        return self.root/kind/(safe_id(identifier)+'.json')

    def register_b1(self,batch_id,folder,destination,*,recommendations=None,official_materials=None,reopen_ids=None):
        """Trusted export operation, never register an uploaded client's sidecar."""
        batch=self.store.load(batch_id);check_sources(batch)
        folder=Path(folder);manifest=read_json(folder/'export_manifest.json');plan=read_json(folder/'review_plan.json')
        if manifest['workbook_sha256']!=sha((folder/FILENAME).read_bytes()):raise Rejected('原导出文件哈希失配')
        catalog=material_catalog(plan['source_catalog'].get('path'))
        if digest(batch)!=plan['base_hash'] or plan!=build_plan(batch,catalog):
            raise Rejected('B1导出镜像不是当前可信版本')
        destination=Path(destination).resolve();Store(destination.parent).guard(batch)
        if destination.suffix.lower()!='.xlsx':raise Rejected('审核导出必须为.xlsx')
        if destination.exists():raise Rejected('拒绝覆盖已有业务文件')
        recommendations=recommendations or {};official_materials=official_materials or {}
        ids={r['id'] for r in batch['rows']}
        by_id={r['id']:r for r in batch['rows']}
        if set(recommendations)-ids or set(official_materials)-ids:raise Rejected('推荐/参考包含未知商品')
        layers={}
        for r in batch['rows']:
            rid=r['id'];m=plan['observations'][rid]['material']
            histories=[{'kind':'historical_reference','value':m['business'],'source':r['source'],'review_status':'candidate_only'},
                       {'kind':'historical_reference','value':m['history6'],'source':r.get('historical',{}).get('internal_reference'), 'review_status':'candidate_only'},
                       {'kind':'historical_query','value':m['query'],'source':m['evidence'],'review_status':'candidate_only'}]
            history=[]
            for h in histories:
                if h['value']:history.append(dict(h,id=digest(h)))
            official=[]
            for item in official_materials.get(rid,[]):
                if not isinstance(item,dict) or not item.get('original') or not item.get('source_ref'):
                    raise Rejected('官方参考须保留原文与来源，不可伪造为已审核申报选择')
                official.append(dict(copy.deepcopy(item),id=digest(item),is_final_decision=False))
            rec=recommendations.get(rid)
            if rec is not None:
                if not isinstance(rec,dict) or not all(isinstance(rec.get(k),str) and rec[k].strip() for k in ('declaration_material','reason','algorithm_version')):
                    raise Rejected('AI推荐须包含明确选择、理由和生成版本')
                rec=dict(copy.deepcopy(rec),id=digest(rec),review_status='recommendation_only')
            layers[rid]={'official_materials':official,'historical_reference_materials':history,
                'recommendations':[rec] if rec else [],'current_recommendation':rec}
        reopen_ids=reopen_ids or []
        if len(reopen_ids)!=len(set(reopen_ids)) or set(reopen_ids)-ids:raise Rejected('纠错范围含重复/未知商品')
        covered={rid for i in plan['issues'] if i['domain']=='material' for rid in i['row_ids']}
        # A prior physical-material approval is NOT a declaration selection approval.
        # Reviewed declaration heads need no routine task; explicitly reopened rows do.
        extra=[r for r in batch['rows'] if r['id'] not in covered and (head(r) is None or r['id'] in reopen_ids)]
        view=dict(plan,issues=copy.deepcopy(plan['issues']),memberships=copy.deepcopy(plan['memberships']),stats=copy.deepcopy(plan['stats']))
        for issue in view['issues']:
            if issue['domain'] not in ('material','tax'):continue
            candidates={}
            for rid in issue['row_ids']:
                r=by_id[rid];cur=head(r);rec=layers[rid]['current_recommendation']
                val=cur['human_final'].get('declaration_material') if cur else rec['declaration_material'] if rec else '尚无明确AI推荐'
                candidates[val]=candidates.get(val,0)+1
            issue['candidate']='\n'.join(f'{v}（{n}行）' for v,n in candidates.items())
        extra_groups={}
        for r in extra:
            rec=layers[r['id']]['current_recommendation']
            key=digest([bool(head(r)),rec['declaration_material'] if rec else None,(r['material']['selected'] or {}).get('translation')])
            extra_groups.setdefault(key,[]).append(r)
        for group in extra_groups.values():
            r=group[0];rec=layers[r['id']]['current_recommendation'];qid='Q'+str(len(view['issues'])+1).zfill(3)
            view['issues'].append({'id':qid,'domain':'material','row_ids':[x['id'] for x in group],'subject':'最终申报材质选择/按需纠错',
                'candidate':rec['declaration_material'] if rec else '没有AI推荐，需明确选择',
                'reason':'原物理材质审核不等于最终申报选择' if head(r) is None else '业务明确要求纠错；未操作不改变旧批准',
                'needed':'只填实际改选与原因；审核后记录新版本，不直接覆盖旧确认','priority':'需明确选择' if head(r) is None else '按需纠错',
                'navigation_group_only':True,'group_key':None,'contexts':{x['id']:context(x) for x in group}})
            for x in group:view['memberships'].setdefault(x['id'],[]).append(qid)
        if extra:
            view['stats']['issue_groups']=len(view['issues'])
            view['stats']['issue_groups_by_domain']['material']=view['stats']['issue_groups_by_domain'].get('material',0)+len(extra_groups)
        destination.parent.mkdir(parents=True,exist_ok=True)
        with tempfile.TemporaryDirectory(prefix='.b2-view-',dir=destination.parent) as tmp:
            path=Path(tmp)/'view.xlsx';_write(view,path);wb=_read_workbook(path.read_bytes())
        plan=view;export_id=uuid4().hex+uuid4().hex
        wb.properties.identifier='arrival-return:'+export_id
        label='【合成权限技术测试，不是真实业务审核】' if batch.get('execution_mode')=='synthetic_technical_test' else ''
        for ws in wb:
            literal(ws['A3'],label+'内部审核回传版；填写仅为提议，预览后须可信身份明确确认。正式输出关闭。')
        literal(wb[SHEETS[0]]['E18'],'只改必要内容；回传先看差异再明确确认，操作者不从Excel取值。')
        literal(wb[SHEETS[1]]['A5'],'浅黄色两列可填；范围=本组全部或本表序号；用“申报材质=…；原因=…”记录必要修改。')
        detail=wb[SHEETS[2]];literal(detail.cell(6,11),'当前申报材质推荐/最终选择');literal(detail.cell(6,12),'推荐依据及状态')
        literal(detail.cell(6,18),'归类及有效提示')
        literal(detail.cell(6,14),'当前税率/待核验')
        for r in batch['rows']:
            rid=r['id'];n=6+plan['ordinals'][rid];current=head(r);rec=layers[rid]['current_recommendation']
            literal(detail.cell(n,11),current['human_final'].get('declaration_material') if current else rec['declaration_material'] if rec else None)
            literal(detail.cell(n,12),('人工选择：'+current['state']) if current else ('AI推荐，未审核：'+rec['reason']) if rec else '无AI推荐，历史参考保存在后台')
            literal(detail.cell(n,18),current['human_final'].get('classification_code') if current else '归类未确认')
            literal(detail.cell(n,14),f"关税{current['human_final'].get('duty_rate','待定')} / VAT{current['human_final'].get('vat_rate','待定')}" if current else '未确认，历史参考在后台')
        with tempfile.TemporaryDirectory(prefix='.b2-save-',dir=destination.parent) as tmp:
            path=Path(tmp)/'view.xlsx';save_exact(wb,path);_archive(destination,path.read_bytes())
        wb.close()
        payload=destination.read_bytes();basepath=self.root/'export_files'/(export_id+'.xlsx');_archive(basepath,payload)
        registry={'contract':CONTRACT,'export_id':export_id,'batch_id':batch_id,'base_hash':digest(batch),
            'workbook_sha256':sha(payload),'base_file':str(basepath),'created_at':now(),
            'issues':plan['issues'],'ordinals':plan['ordinals'],'ordinary_candidate_ids':plan['ordinary_candidate_ids'],
            'product_candidates':{r['id']:r['product'] for r in plan['candidate_snapshot']['rows']},'layers':layers,
            'source_catalog':plan['source_catalog'],'execution_mode':batch.get('execution_mode','business')}
        path=self._path('exports',export_id);path.parent.mkdir(parents=True,exist_ok=True);atomic_json(path,registry)
        return {'export_id':export_id,'base_hash':registry['base_hash'],'file':str(destination),
                'workbook_sha256':registry['workbook_sha256'],'formal_output':False}

    def _resolve_export(self, wb, *, export_id=None, batch_id=None):
        """Recover only an absent property against trusted, immutable exports."""
        identifier = wb.properties.identifier
        if identifier:
            if not isinstance(identifier,str) or not identifier.startswith('arrival-return:'):
                raise Rejected('非法登记标识')
            bound = safe_id(identifier.removeprefix('arrival-return:'))
            if export_id and bound != export_id:
                raise Rejected('上传文件与指定导出ID不一致')
            candidates = [self._path('exports',bound)]
        elif export_id:
            candidates = [self._path('exports',export_id)]
        elif batch_id:
            candidates = sorted((self.root/'exports').glob('*.json'))
        else:
            raise Rejected('回传登记属性缺失；请在原批次会话上传，由宿主选择已登记导出版本')
        matches=[]
        for candidate in candidates:
            try:
                registry=read_json(candidate)
                if batch_id and registry['batch_id'] != batch_id:
                    if identifier or export_id: raise Rejected('审核表不属于指定批次')
                    continue
                base_data=Path(registry['base_file']).read_bytes()
                if sha(base_data)!=registry['workbook_sha256']:raise Rejected('可信原表被修改')
                base=_read_workbook(base_data)
                try:
                    if base.properties.identifier != 'arrival-return:'+candidate.stem:
                        raise Rejected('可信登记与原导出属性不一致')
                    edits=compare_workbooks(base,wb,recover_missing_identifier=not identifier)
                finally:base.close()
                matches.append((candidate.stem,registry,edits))
            except (Rejected,OSError,ValueError,KeyError):
                if identifier or export_id:raise
        if len(matches)!=1:
            if not matches:raise Rejected('没有匹配的可信导出；不能猜测恢复或搬运修改')
            raise Rejected('多个已登记导出内容相同；请由宿主展示并选择具体导出版本')
        bound,registry,edits=matches[0]
        binding = None if identifier else {'schema':'japan-arrival/return-binding-1',
            'method':'trusted_readonly_content_match','export_id':bound,'batch_id':registry['batch_id'],
            'readonly_match_sha256':digest({'workbook':registry['workbook_sha256'],'base_hash':registry['base_hash']})}
        return bound,registry,edits,binding

    def preview(self,uploaded,*,export_id=None,batch_id=None,operations=None):
        data=Path(uploaded).read_bytes();wb=_read_workbook(data)
        try:
            export_id,registry,edits,binding=self._resolve_export(wb,export_id=export_id,batch_id=batch_id)
        finally:wb.close()
        batch=self.store.load(registry['batch_id']);check_sources(batch)
        notes=[];proposals=[];issues={i['id']:i for i in registry['issues']}
        for e in edits:
            issue=issues[e['issue_id']];choice,note=e['after'];values,control,unparsed=_parse_note(note)
            if choice is None and note is None:
                notes.append(dict(e,status='cleared_request_remains_pending'));continue
            selected=_scope(registry,issue,control)
            if issue['domain']=='price':
                if choice=='保留待确认':
                    notes.append(dict(e,row_ids=selected,status='business_amount_pending'));continue
                if unparsed or set(values)-{'reason'}-{k for k in values if k.startswith('business_amount.')}:
                    raise Rejected('金额修改须明确结构化金额、单位、数量、原因和依据')
                from .business_amount import preview as amount_preview
                items={}
                for rid in selected:
                    row=next(r for r in batch['rows'] if r['id']==rid)
                    v={k.split('.',1)[1]:value for k,value in values.items() if k.startswith('business_amount.')}
                    if choice=='采用候选':
                        v.setdefault('amount',row['authority']['amount'])
                        v.setdefault('currency',row['authority']['currency'])
                        v.setdefault('quantity',row['authority']['quantity'])
                        if row['authority']['quantity_unit']:v.setdefault('quantity_unit',row['authority']['quantity_unit'])
                    v['reason']=values.get('reason','明确采用本导出版逐行展示金额' if choice=='采用候选' else '')
                    items[rid]=v
                amount=amount_preview(self.root,batch['id'],items,reference=control.get('依据',''))
                notes.append(dict(e,row_ids=selected,status='business_amount_preview_only',amount_preview=amount))
                continue
            if issue['domain'] not in ('material','tax','product'):
                notes.append(dict(e,row_ids=selected,status='pending_external_business_decision'));continue
            if issue['domain']=='product':
                if set(values)-{k for k in values if k.startswith('product.')}-{'reason'}:raise Rejected('商品问题不能混入专业选择')
                proposals.append({'action':'edit_product','row_ids':selected,'values':{k[8:]:v for k,v in values.items() if k.startswith('product.')},
                    'reason':values.get('reason',''),'reference':control.get('依据',''),'unparsed_note':unparsed,'issue_id':issue['id']})
            else:
                if any(k.startswith('business_amount.') for k in values):raise Rejected('专业归类问题不能夹带金额修改')
                if any(k.startswith('product.') for k in values):raise Rejected('专业问题不能静默修改基础资料')
                action='adopt_recommendation' if choice=='采用候选' else 'return_decision' if choice=='保留待确认' else 'edit_decision'
                proposals.append({'action':action,'row_ids':selected,'values':values,'reference':control.get('依据',''),
                                  'unparsed_note':unparsed,'issue_id':issue['id']})
        ops=copy.deepcopy(operations or [])
        self._validate_operations(registry,ops)
        differences=[]
        current_rows={r['id']:r for r in batch['rows']}
        display_proposals=proposals+[dict(op,issue_id='Skill明确补充',unparsed_note='') for op in ops if op['action']=='edit_decision']
        for p in display_proposals:
            for rid in p['row_ids']:
                row=current_rows[rid];old=head(row);new=copy.deepcopy(old['human_final']) if old else {}
                fields=dict(p.get('values',{}),**p.get('per_row_values',{}).get(rid,{}))
                new.update({k:v for k,v in fields.items() if k in FINAL_FIELDS})
                if p['action']=='adopt_recommendation':
                    rec=registry['layers'][rid]['current_recommendation']
                    if not rec:raise Rejected('该范围没有AI推荐，不能把历史材质当作推荐采用')
                    new['declaration_material']=rec['declaration_material']
                differences.append({'row_id':rid,'issue_id':p['issue_id'],'action':p['action'],
                    'before_final':copy.deepcopy(old['human_final']) if old else None,'proposed_final':new,
                    'proposed_facts':{k:v for k,v in fields.items() if k in FACT_FIELDS},
                    'original_recommendation':registry['layers'][rid]['current_recommendation'],
                    'context':context(row),'unresolved_text':p['unparsed_note'],
                    'impact':'只改必要域；最终材质/用途构造变化使关联归类税率和税金待重审，原金额/无关审核保留'})
        upload_hash=sha(data);archive=self.root/'uploads'/(upload_hash+'.xlsx');_archive(archive,data)
        preview={'contract':CONTRACT,'export_id':export_id,'batch_id':batch['id'],'base_hash':registry['base_hash'],
            'upload_sha256':upload_hash,'upload_file':str(archive),'edits':edits,'proposals':proposals,'issue_notes':notes,
            'operations':ops,'decision_differences':differences,'base_is_current':digest(batch)==registry['base_hash'],
            'notice':'只预览，没有保存审核。须明确确认此预览；旧版本新操作将拒绝。'}
        if binding is not None:preview['binding']=binding
        identifier=digest(preview);path=self._path('previews',identifier);path.parent.mkdir(parents=True,exist_ok=True)
        if path.exists() and read_json(path)!=preview:raise Rejected('预览登记冲突')
        if not path.exists():atomic_json(path,preview)
        return dict(preview,preview_id=identifier,confirmation_hash=identifier)

    def _validate_operations(self,registry,operations):
        if not isinstance(operations,list):raise Rejected('明确操作必须为列表')
        all_ids=set(registry['ordinals'])
        for op in operations:
            if not isinstance(op,dict) or set(op)-{'action','row_ids','scope','reference','tax_parameters','reason','mode','resolutions','values','per_row_values'}:
                raise Rejected('操作含未知字段；身份/权限不能来自请求')
            if op.get('action') not in ('confirm_products','approve_material','approve_classification','return_decision','revoke_decision','edit_decision'):
                raise Rejected('不支持的明确操作')
            if op.get('scope'):
                if op['scope']!='ordinary_products' or op['action']!='confirm_products' or op.get('row_ids'):
                    raise Rejected('仅普通商品确认支持已登记的普通范围')
                op['row_ids']=registry['ordinary_candidate_ids'][:];op.pop('scope')
            ids=op.get('row_ids')
            if not isinstance(ids,list) or not ids or len(ids)!=len(set(ids)) or set(ids)-all_ids:raise Rejected('操作范围为空/重复/未知')
            if not isinstance(op.get('reference'),str) or not op['reference'].strip():raise Rejected('明确审核操作须提供凭据引用')
            if op.get('tax_parameters') is not None and (not isinstance(op['tax_parameters'],dict) or set(op['tax_parameters'])-set(ids)):
                raise Rejected('税基参数含范围外商品')
            if 'values' in op or 'per_row_values' in op:
                if op['action']!='edit_decision':raise Rejected('审核批准不能夹带未预览的修改；修改须单独列明')
                validate_values(op.get('values',{}));per=op.get('per_row_values',{})
                if not isinstance(per,dict) or set(per)-set(ids):raise Rejected('逐行事实包含范围外商品')
                for values in per.values():
                    validate_values(values)
                    if set(values)&set(op.get('values',{})):raise Rejected('公共和逐行事实不允许静默覆盖')
            if op['action'] in ('approve_material','approve_classification') and len(ids)>1:
                # Each row must have a context explicitly pinned in the trusted export. Approval
                # reference is an explicit assertion of coverage, not a matching-text exemption.
                if not op.get('reason'):raise Rejected('跨商品成组审核须明确说明凭据覆盖范围')

    def commit(self,preview_id,*,confirmed_hash,explicit_confirmation,reference,session):
        safe_id(preview_id)
        if explicit_confirmation is not True or confirmed_hash!=preview_id or not isinstance(reference,str) or not reference.strip():
            raise Rejected('必须明确确认已展示的预览版本并提供操作依据')
        preview=read_json(self._path('previews',preview_id))
        if digest(preview)!=preview_id:raise Rejected('预览登记被修改')
        registry=read_json(self._path('exports',preview['export_id']))
        data=Path(preview['upload_file']).read_bytes()
        if sha(data)!=preview['upload_sha256']:raise Rejected('归档上传被修改')
        batch=self.store.load(preview['batch_id']);check_sources(batch)
        # Reauthenticate/config/D08 checks precede replay as well as new mutations.
        if not isinstance(session,Session):raise Rejected('审核身份必须来自可信执行上下文')
        permissions=authorize(self.root,session,batch,'business' if 'business' in session.roles else 'material' if 'material' in session.roles else 'customs')
        for op in preview['operations']:
            authorize(self.root,session,batch,{'confirm_products':'business','approve_material':'material',
                'approve_classification':'customs','return_decision':'material','revoke_decision':'material','edit_decision':'material'}[op['action']])
        key=digest({'export_id':preview['export_id'],'edits':preview['edits'],'operations':preview['operations'],
                    'actor':session.principal,'mode':session.mode,'reference':reference})
        if key in batch.get('return_receipts',{}):
            return dict(batch['return_receipts'][key],replayed=True,current_hash=digest(batch))
        if digest(batch)!=preview['base_hash']:raise Rejected('旧版本/并发冲突：不得覆盖新版本，请重新导出')
        # Never trust cached parsed instructions after an archive or registry change.
        rebuilt=self.preview(preview['upload_file'],export_id=preview['export_id'],operations=preview['operations'])
        if rebuilt['preview_id']!=preview_id:raise Rejected('回传内容/范围与已确认预览不一致')
        out=copy.deepcopy(batch);start_events=len(out['events']);decision_start=len(out.get('decision_events',[]))
        out['active_review_lock']=None
        by_id={r['id']:r for r in out['rows']};out['product_review_profile']='exception-first-v1'
        for proposal in preview['proposals']:
            ids=proposal['row_ids'];values=proposal['values'];action=proposal['action']
            if action=='edit_product':
                if values:
                    out=supplement_many(out,{rid:values for rid in ids},actor=session.principal,
                        evidence=proposal['reference'] or reference,permissions=permissions,_declaration_commit=True)
                    by_id={r['id']:r for r in out['rows']}
                    if set(values)&{'category','customs_name'}:
                        for rid in ids:
                            if by_id[rid].get('declaration'):
                                revise(out,by_id[rid],registry['layers'][rid],'invalidate_classification',{},session,proposal['reference'] or reference)
                out.setdefault('return_issue_notes',[]).append(proposal)
                continue
            for rid in ids:
                revise(out,by_id[rid],registry['layers'][rid],action,values,session,
                       proposal['reference'] or reference,note=proposal['unparsed_note'])
        for op in preview['operations']:
            ids=op['row_ids'];action=op['action']
            if action=='confirm_products':
                if any(p['action']=='edit_product' and p['unparsed_note'] and set(p['row_ids'])&set(ids) for p in preview['proposals']):
                    raise Rejected('商品修改说明尚未澄清，不能直接确认')
                fill={rid:{'customs_name':registry['product_candidates'][rid]['customs_name']} for rid in ids
                      if not by_id[rid]['product']['customs_name'] and registry['product_candidates'][rid]['customs_name']}
                if fill:
                    out=supplement_many(out,fill,actor=session.principal,evidence=op['reference'],permissions=permissions,_declaration_commit=True)
                    by_id={r['id']:r for r in out['rows']}
                    for rid in fill:
                        if head(by_id[rid]) and head(by_id[rid])['reviews']['classification']:
                            revise(out,by_id[rid],registry['layers'][rid],'invalidate_classification',{},session,op['reference'])
                out=confirm_products(out,review_package(out),ids,actor=session.principal,mode=op.get('mode','batch'),
                    evidence=op['reference'],permissions=permissions,resolutions=op.get('resolutions'))
                by_id={r['id']:r for r in out['rows']}
            else:
                for rid in ids:
                    values=dict(op.get('values',{}),**op.get('per_row_values',{}).get(rid,{}))
                    if action in ('return_decision','revoke_decision') and op.get('reason'):values['reason']=op['reason']
                    resolved=[]
                    if action in ('approve_material','approve_classification'):
                        domain='material' if action=='approve_material' else 'tax'
                        disputes=[d for d in by_id[rid]['disputes'] if d.get('domain')==domain]
                        supplied=op.get('resolutions')
                        if supplied is not None or disputes:
                            if len(ids)!=1:raise Rejected('专业争议必须单项处理，不得混入成组批准')
                            if not isinstance(supplied,dict) or set(supplied)!={d['id'] for d in disputes}:
                                raise Rejected('须逐项明确本域全部争议及处理依据，不能解除身份/未知争议')
                            if not disputes or any(not isinstance(v,str) or not v.strip() for v in supplied.values()):
                                raise Rejected('争议处理依据不能留空或虚构争议')
                            resolved=[dict(d,resolution=supplied[d['id']],reference=op['reference']) for d in disputes]
                            by_id[rid]['disputes']=[d for d in by_id[rid]['disputes'] if d.get('domain')!=domain]
                    revise(out,by_id[rid],registry['layers'][rid],action,
                        values,
                        session,op['reference'],tax_parameters=op.get('tax_parameters',{}).get(rid),resolved_disputes=resolved)
        for note in preview['issue_notes']:out.setdefault('return_issue_notes',[]).append(note)
        if not preview['proposals'] and not preview['operations'] and not preview['issue_notes']:raise Rejected('没有修改或明确操作，不能默认批准')
        staged_core=out['events'][start_events:];out['events']=out['events'][:start_events];out['revision']=batch['revision']
        new_events=out.get('decision_events',[])[decision_start:]
        for e in new_events:e['transaction_id']=key
        # Transaction ID is envelope metadata; the immutable event content hash excludes it.
        receipt={'transaction_id':key,'preview_id':preview_id,'upload_sha256':preview['upload_sha256'],
            'revision':batch['revision']+1,'decision_event_ids':[e['event_id'] for e in new_events],
            'execution':session.evidence(),'formal_output':False,'knowledge_index':'pending_delivery'}
        out.setdefault('return_receipts',{})[key]=receipt
        ids=sorted({rid for p in preview['proposals']+preview['operations'] for rid in p['row_ids']})
        out=changed(batch,out,session.principal,'excel_return_commit',ids,{'reference':reference,
            'transaction_id':key,'preview_id':preview_id,'execution':session.evidence(),
            'decision_event_ids':receipt['decision_event_ids'],'reused_core_events':staged_core,'issue_notes':preview['issue_notes']})
        self.store.save(out,expected_hash=digest(batch))
        return dict(receipt,replayed=False,current_hash=digest(out))

    def status(self,batch_id):
        batch=self.store.load(batch_id)
        return {'batch_id':batch_id,'revision':batch['revision'],'hash':digest(batch),
            'decisions':[head(r) for r in batch['rows'] if head(r)],
            'calculations':{r['id']:{'box_tax':box_tax(r),'price':exact_price(r)} for r in batch['rows'] if head(r)},
            'formal_output':False,'automatic_rules':False,'active_review_lock':batch.get('active_review_lock')}

    def lock_review(self,batch_id,*,session,reference):
        """Internal audited snapshot; a lock never authorizes formal export."""
        from .batchflow import workflow_status
        batch=self.store.load(batch_id);check_sources(batch);authorize(self.root,session,batch,'release')
        if not isinstance(reference,str) or not reference.strip():raise Rejected('锁版须明确审核依据')
        status=workflow_status(batch)
        if (not status['contract']['all_contract_items_complete'] or not status['logistics']['passed']
            or not status['quantities']['passed'] or any(status['rows'].values())
            or any(not head(r) or not head(r)['technical_approval_complete'] for r in batch['rows'])):
            raise Rejected('必要业务输入、各域审核或数量金额箱重核验未齐，不能锁版；正式输出仍关闭')
        lock={'reviewed_hash':digest(batch),'source_hashes':[s['sha256'] for s in batch['sources']],
            'decision_versions':{r['id']:head(r)['version'] for r in batch['rows']},
            'execution':session.evidence(),'reference':reference,'at':now(),'formal_output_authorized':False}
        lock['lock_id']=digest(lock);out=copy.deepcopy(batch)
        out.setdefault('review_locks',[]).append(lock);out['active_review_lock']=lock['lock_id']
        out=changed(batch,out,session.principal,'review_lock',list(lock['decision_versions']),lock)
        self.store.save(out,expected_hash=digest(batch));return lock


def decision_event_stream(batch):
    """Reliable D1 interface; no knowledge store is implemented in B2."""
    previous=None
    for event in batch.get('decision_events',[]):
        payload={k:v for k,v in event.items() if k not in ('event_id','transaction_id')}
        if digest(payload)!=event['event_id']:raise Rejected('决策事件内容哈希错误')
        if event['previous_event_hash']!=previous:raise Rejected('决策事件链不完整')
        previous=event['event_id']
        yield copy.deepcopy(event)


def deliver_events(store,batch_id,sink):
    """At-least-once delivery. Consumer must upsert by event_id and verify payload_hash."""
    batch=store.load(batch_id);folder=store.root/'event_delivery'/batch_id
    with store.lock(folder):
        path=folder/'receipts.json';acks=read_json(path) if path.exists() else {}
        failures=[]
        for event in decision_event_stream(batch):
            eid=event['event_id'];payload_hash=digest(event)
            if eid in acks:
                if acks[eid]['payload_hash']!=payload_hash:raise Rejected('事件投递回执冲突')
                continue
            try:
                ack=sink(event_id=eid,payload_hash=payload_hash,payload=event)
                if ack!={'event_id':eid,'payload_hash':payload_hash}:raise Rejected('索引未返回匹配的幂等回执')
                acks[eid]=dict(ack,delivered_at=now());atomic_json(path,acks)
            except Exception as exc:
                failures.append({'event_id':eid,'error_type':type(exc).__name__});break
        return {'delivered':len(acks),'pending':len(batch.get('decision_events',[]))-len(acks),
                'failures':failures,'batch_saved':True,'business_learning_not_asserted':True}
