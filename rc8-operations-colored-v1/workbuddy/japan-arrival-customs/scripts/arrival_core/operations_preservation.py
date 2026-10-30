"""Source-bound controlled delivery, never a tax authority or formal release path."""
import hashlib
import json
from fractions import Fraction
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, PatternFill

from .core import Rejected
from .storage import atomic_json, check_sources
from .excel_writer import literal


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def load_reference_evidence(request, plan):
    """Replay an explicitly supplied batch trial dossier, never search/guess rates.

    The dossier is business evidence outside the install package. It is bound to
    both original files, source row identity and original rate cells. Confirmation
    binds its complete digest; no identity or official approval is inferred.
    """
    if request is None: return None
    if not isinstance(request, dict) or set(request) != {'path', 'sha256'}:
        raise Rejected('参考试算证据须提供文件及SHA256')
    path = Path(request['path']).resolve(strict=True)
    if sha(path) != request['sha256']: raise Rejected('参考试算证据已变化')
    doc = json.loads(path.read_text(encoding='utf-8-sig'))
    expected = sorted(s['sha256'] for s in plan['sources'] if s['kind'] == 'operations_original')
    if doc.get('schema') != 'operations-reference-trial-v1' or sorted(doc.get('operations_source_hashes', [])) != expected:
        raise Rejected('参考试算证据不属于当前两份运营原件，禁止跨批沿用')
    if doc.get('formal_output') is not False or not doc.get('notice'):
        raise Rejected('参考试算必须明确非正式范围')
    vat = Fraction(doc['vat_rate'])
    if not 0 <= vat <= 1: raise Rejected('增值税率范围无效')
    by_row = {str(r['source_row']): r for r in doc['rows']}
    if len(by_row) != len(doc['rows']) or len(by_row) != len(plan['details']):
        raise Rejected('参考试算行数不全或重复')
    books = {}; sources = [{'path': str(path), 'sha256': request['sha256'], 'kind': 'batch_reference_trial_dossier'}]
    try:
        for item in plan['details']:
            row = by_row.get(str(item['source_rows']['business']))
            v = item['values']
            if row is None or any(row.get(key) != v[field] for key, field in
                    (('code', '商品编码'), ('category', '品类'), ('material', '最终申报材质'))):
                raise Rejected('参考试算身份/材质/来源行不匹配')
            if row['duty_rate'] is None:
                if not row.get('reason'): raise Rejected('待算税金缺少原因')
                continue
            duty = Fraction(row['duty_rate'])
            if not 0 <= duty <= 1 or not row.get('basis') or not row.get('references'):
                raise Rejected('税率范围或参考依据不完整')
            for ref in row['references']:
                name = str(Path(ref['path']).resolve(strict=True))
                if name not in books:
                    if sha(name) != ref['sha256']: raise Rejected('税率参考原件已变化')
                    books[name] = load_workbook(name, data_only=False, keep_links=False)
                    sources.append({'path': name, 'sha256': ref['sha256'], 'kind': 'batch_reference_rate_source'})
                elif next(s['sha256'] for s in sources if s['path'] == name) != ref['sha256']:
                    raise Rejected('同一税率来源哈希冲突')
                ws = books[name][ref['sheet']]; cell = ws[ref['cell']]
                if cell.data_type == 'f' or cell.value is None or Fraction(str(cell.value)) != duty:
                    raise Rejected('参考税率不是原件独立数值或与证据不一致')
                if ref.get('kind') == 'tax_table':
                    for c in (cell, ws[ref['label_cell']]):
                        color = c.fill.fgColor
                        gray = c.fill.patternType == 'solid' and (
                            (color.type == 'theme' and (color.theme, color.tint) in ((0, -0.5), (1, 0.25)))
                            or (color.type == 'rgb' and str(color.rgb)[-6:].upper() in ('808080', 'C0C0C0', 'BFBFBF', 'D9D9D9', 'A6A6A6')))
                        if gray: raise Rejected('参考税目或税率已标灰过期，禁止使用')
    finally:
        for book in books.values(): book.close()
    # Recheck after reading; no rate lookup writes or implicit defaults.
    for s in sources:
        if sha(s['path']) != s['sha256']: raise Rejected('读取期间参考依据发生变化')
    return {'by_source_row': by_row, 'vat_rate': doc['vat_rate'], 'notice': doc['notice'], 'sources': sources}


def source_value(row, field):
    return row.get('operations_preserved', {}).get(field, {}).get('literal_value')


def privacy_tokens(batch):
    tokens = set()
    for row in batch['rows']:
        for name in ('internal_name', 'yd', 'channel', 'internal_note'):
            value = row['product'].get(name)
            raw_yd = source_value(row, 'yd')
            if name == 'yd' and isinstance(raw_yd, (int, float)) and raw_yd == 0:
                continue  # Typed placeholder, not a global string replacement token.
            if value not in (None, '') and len(str(value).strip()) >= 4: tokens.add(str(value))
        for name in ('specification', 'remarks'):
            value = source_value(row, name)
            if value not in (None, '') and len(str(value).strip()) >= 4: tokens.add(str(value))
    return tokens


def write_delivery_report(batch, pair, folder):
    from .drafts import draft_payload, verify_pair
    from .report_templates import FILES
    from .batchflow import require_formal_release
    check_sources(batch)
    proof = verify_pair(batch, pair, draft_payload(batch))
    price_rows = []; tax_rows = []; total = Fraction()
    i = load_workbook(Path(pair) / FILES['internal'], data_only=True)
    c = load_workbook(Path(pair) / FILES['customs'], data_only=True)
    try:
        for er, row in enumerate(batch['rows'], 2):
            source = row['original_pairing']; n = source['business']
            amount = Fraction(row['authority']['amount']); total += amount
            if Fraction(str(i.worksheets[0].cell(er, 11).value)) != amount or Fraction(str(c.worksheets[0].cell(er, 16).value)) != amount:
                raise Rejected(f'输出第{er}行确认金额未完整保留')
            if str(i.worksheets[0].cell(er, 3).value) != row['identity']['code'] or str(c.worksheets[0].cell(er, 2).value) != row['identity']['code']:
                raise Rejected('两表记录位置不一致')
            for field, col in (('specification', 6), ('yd', 19), ('remarks', 20)):
                preserved = row.get('operations_preserved', {}).get(field)
                if preserved and i.worksheets[0].cell(er, col).value != preserved['literal_value']:
                    raise Rejected('原表内容或备注未完整保留')
            base = {'source_file': row['operations_provenance']['业务确认金额JPY']['path'],
                    'source_sheet': row['operations_provenance']['业务确认金额JPY']['sheet'],
                    'source_row': n, 'output_row': er, 'declaration_source_row': source['declaration'],
                    'code': row['identity']['code'], 'business_amount_jpy': str(amount)}
            if row['confirmed_input'].get('price_blocker'):
                quantity = Fraction(row['authority']['quantity']) * Fraction(row['packaging']['declaration_factor'])
                price_rows.append(dict(base, exact_price=str(amount / quantity),
                    reason='单价不能精确保存；确认金额固定保留，不随G/F修改联算；不自动调平'))
            if row['tax'] is None:
                ref = row.get('operations_tax_reference')
                tax_rows.append(dict(base, category=row['product']['category'],
                    reason=ref['reason'] if ref else '未提供本批已确认的税率/税基单位/公式及参考依据，未套用默认税率'))
    finally:
        i.close(); c.close()
    try: require_formal_release(batch)
    except Rejected: pass
    else: raise Rejected('正式门禁异常，拒绝交付')
    report = {'schema': 'operations-delivery-check-v1', 'rows': len(batch['rows']), 'business_amount_jpy': str(total),
              'source_order': '原运营表逐行，不分组排序', 'price_exceptions': price_rows,
              'tax_pending': tax_rows, 'verification': proof, 'formal_rejected': True,
              'formal_output': False, 'workbuddy_live': '未验证'}
    path = Path(folder) / 'delivery-report.json'
    exception_path = Path(folder) / '单价异常与待算税金.xlsx'
    result = {'json': str(path), 'exceptions_workbook': str(exception_path), 'price_pending_rows': len(price_rows),
              'tax_pending_rows': len(tax_rows), 'confirmed_amount_jpy': str(total), 'formal_output': False}
    if exception_path.exists():
        if not path.is_file(): raise Rejected('异常说明输出不完整，保留现场后换新目录')
        saved = json.loads(path.read_text(encoding='utf-8'))
        expected_hash = saved.pop('exception_workbook_sha256', None)
        if saved != report or expected_hash != sha(exception_path):
            raise Rejected('既有异常说明已变化，禁止覆盖')
        return result
    wb = Workbook(); wb.remove(wb.active)
    for title, rows, fields in (
        ('单价异常', price_rows, ['source_row', 'output_row', 'code', 'business_amount_jpy', 'exact_price', 'reason']),
        ('待算税金', tax_rows, ['source_row', 'output_row', 'code', 'business_amount_jpy', 'category', 'reason'])):
        ws = wb.create_sheet(title)
        headings = {'source_row': '原表行号', 'output_row': '输出行号', 'code': '商品编码', 'business_amount_jpy': '已确认金额JPY',
                    'exact_price': '精确单价分数', 'category': '品类', 'reason': '原因及补充条件'}
        for col, f in enumerate(fields, 1): literal(ws.cell(1, col), headings[f])
        for er, item in enumerate(rows, 2):
            for col, f in enumerate(fields, 1): literal(ws.cell(er, col), item[f])
        ws.freeze_panes = 'A2'
        for col in ('A', 'B', 'C', 'D', 'E', 'F'): ws.column_dimensions[col].width = 24
        ws.column_dimensions[chr(64 + len(fields))].width = 85
        for r in ws:
            for cell in r: cell.alignment = Alignment(wrap_text=True, vertical='top')
        ws.oddHeader.center.text = '受控参考草稿，不得外发或用于申报'
    wb.save(exception_path); wb.close()
    report['exception_workbook_sha256'] = sha(exception_path)
    atomic_json(path, report)
    return result
