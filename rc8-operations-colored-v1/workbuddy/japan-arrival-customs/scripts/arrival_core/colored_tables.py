"""Two operational result tables, exact recalculation and minimal in-cell colors."""
import hashlib
import copy
import os
import re
import tempfile
import unicodedata
import zipfile
from fractions import Fraction
from pathlib import Path
from xml.etree import ElementTree as ET

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, PatternFill
from openpyxl.workbook.properties import CalcProperties

from .core import Rejected, digest, box_tax
from .storage import atomic_json, check_sources
from .excel_writer import literal, save_exact
from .operations_preservation import source_value, privacy_tokens
from .colored_operations import execution, fraction
from .report_templates import layout, view_sheet, cache_formulas, optional_number, FILES

COLORS = {'yellow':'FFFFF2CC', 'orange':'FFF4B183', 'red':'FFFFC7CE'}
WARNING = '受控草稿，不得外发或用于申报'
LEGEND = '浅黄=取整/分摊/参考  橙=人工调整  红=待补'


def mark(colors, kind, address, state):
    priority = {'yellow':1,'orange':2,'red':3}
    previous = colors[kind].get(address)
    if not previous or priority[state] > priority[previous]: colors[kind][address] = state


def payload(batch):
    # Reuse the accepted field whitelist/presentation and unchanged business tax core.
    from .report_templates import payload as accepted_payload
    data = accepted_payload(batch)
    colors = {'customs':{},'internal':{}}; traces=[]
    business_total=Fraction(); public_total=Fraction(); allocations=0; tax_pending=0
    for n, row in enumerate(batch['rows'],2):
        c, i = data['customs']['rows'][n-2], data['internal']['rows'][n-2]
        executed=execution(row); business_total+=executed['business_amount']; public_total+=executed['amount']
        c[6]=executed['price']; c[15]=f'=G{n}*F{n}'
        if source_value(row,'public_origin') is not None: c[3]=source_value(row,'public_origin')
        data['caches']['customs'][f'P{n}']=optional_number(executed['amount'])
        i[10]=f'=I{n}*J{n}'; i[13]=f'=I{n}/H{n}'
        if executed['adjusted']:
            for col in ('G','P'):mark(colors,'customs',f'{col}{n}','orange')
        elif executed['rounded']:
            for col in ('G','P'):mark(colors,'customs',f'{col}{n}','yellow')
        tax_input = row.get('operational_tax_input') or {}
        rate = tax_input.get('rate'); vat_rate=tax_input.get('vat_rate')
        unit=fraction(row['authority']['amount'])/fraction(row['authority']['quantity'])
        pieces=fraction(row['packaging']['pieces_per_box'])
        i[14:18]=[optional_number(fraction(rate)) if rate is not None else None,None,None,None]
        for col in ('P','Q','R'):data['caches']['internal'].pop(f'{col}{n}',None)
        if rate is not None:
            duty=unit*fraction(rate); i[16]=f'=O{n}*J{n}'
            data['caches']['internal'][f'Q{n}']=optional_number(duty)
            if vat_rate is not None:
                tax_values=box_tax(row)
                if tax_values['status']=='pending':raise Rejected('完整税率输入未进入本批税金模型')
                i[15]=f'=J{n}*(1+O{n})*{vat_rate}'; i[17]=f'=(P{n}+Q{n})*N{n}'
                data['caches']['internal'].update({f'P{n}':optional_number(tax_values['vat']),
                    f'Q{n}':optional_number(tax_values['duty']),f'R{n}':optional_number(tax_values['box_tax'])})
        if tax_input.get('allocation'): allocations+=1
        if tax_input.get('allocation') or tax_input.get('reference'):
            for col in ('O','P','Q','R'):mark(colors,'internal',f'{col}{n}','yellow')
        if rate is None:
            for col in ('O','P','Q','R'):mark(colors,'internal',f'{col}{n}','red')
        elif vat_rate is None:
            for col in ('P','R'):mark(colors,'internal',f'{col}{n}','red')
        if i[17] is None:tax_pending+=1
        traces.append({'source_rows':row['original_pairing'],'code':row['identity']['code'],
                       'exact_price':executed['exact_price'],'round0':executed['round0'],'selected_price':executed['price'],
                       'business_amount':str(executed['business_amount']),'customs_amount':str(executed['amount']),
                       'line_difference':str(executed['amount']-executed['business_amount']),
                       'adjusted':executed['adjusted'],'rounded':executed['rounded'],'source':executed['source'],
                       'tax_input':tax_input})
    difference=public_total-business_total
    if difference:
        mark(colors,'customs','P1','red')
        mark(colors,'internal',f'K{len(batch["rows"])+2}','red')
    data.update(colors=colors,operational_trace=traces,business_total=str(business_total),customs_total=str(public_total),
        batch_difference=str(difference),allocation_rows=allocations,tax_pending_rows=tax_pending,
        simple_delivery=True,warning=WARNING,weight_snapshot=batch['operations_delivery_policy'].get('weight_snapshot'),
        rounding_rows=sum(t['rounded'] for t in traces),adjustment_rows=sum(t['adjusted'] for t in traces))
    return data


def text_lines(value, width):
    if value is None:return 1
    return sum(max(1,(sum(2 if unicodedata.east_asian_width(ch) in ('W','F') else 1 for ch in line)+width-1)//width)
               for line in str(value).split('\n'))


def weight_view(data):
    source=data.get('weight_snapshot')
    if source:
        return {'headers':source['rows'][0], 'rows':source['rows'][1:]},source['sheet']
    return data['weights'],layout()['weights']['sheet']


def write_pair(data, directory):
    formats=layout()
    for kind,name in FILES.items():
        wb=Workbook();wb.remove(wb.active)
        wb.properties.title=WARNING;wb.properties.description=data['basis_hash']
        wb.calculation=CalcProperties(calcId=0,fullCalcOnLoad=True,forceFullCalc=True)
        ws=view_sheet(wb,formats[kind],data[kind],trusted_formulas={16} if kind=='customs' else {4,11,14,16,17,18})
        ws.oddHeader.center.text=WARNING;ws.oddFooter.center.text=LEGEND
        caches={1:data['caches'][kind]}
        if kind=='customs':
            ws.column_dimensions['H'].width=33
            for n in range(2,len(data[kind]['rows'])+2):
                ws.cell(n,7).number_format='"¥"#,##0'
                ws.cell(n,8).alignment=Alignment(wrap_text=True,vertical='top')
        else:
            widths={5:62,6:45,7:85,12:26,13:34,19:24,20:45}
            for col,width in widths.items():ws.column_dimensions[ws.cell(1,col).column_letter].width=width
            for col in ('O','P','Q','R'):ws.column_dimensions[col].width=16
            for n in range(2,len(data[kind]['rows'])+2):
                for col in widths:ws.cell(n,col).alignment=Alignment(wrap_text=True,vertical='top')
                ws.cell(n,15).number_format='0.00%'
                ws.row_dimensions[n].height=max(ws.row_dimensions[n].height or 15,
                    15*max(text_lines(ws.cell(n,c).value,width-2) for c,width in widths.items()))
            last=len(data[kind]['rows'])+1
            ws.cell(last+1,11,f'=SUM(K2:K{last})')
            caches[1][f'K{last+1}']=optional_number(fraction(data['business_total']))
            view,title=weight_view(data)
            weights=view_sheet(wb,formats['weights'],view,trusted_formulas=set())
            weights.title=title;weights.oddHeader.center.text=WARNING
        for address,state in data['colors'][kind].items():
            # The accepted template shares column styles; detach before a cell-only color.
            ws[address]._style=copy.copy(ws[address]._style)
            ws[address].fill=PatternFill('solid',fgColor=COLORS[state])
        path=Path(directory)/name;save_exact(wb,path);wb.close();cache_formulas(path,caches)


def verify(batch,directory,data=None):
    check_sources(batch);directory=Path(directory);data=data or payload(batch)
    results={};tokens=privacy_tokens(batch)
    privacy=re.compile('|'.join(re.escape(t) for t in sorted(tokens,key=len,reverse=True))) if tokens else None
    for kind,name in FILES.items():
        path=directory/name;wb=load_workbook(path,data_only=False);cached=load_workbook(path,data_only=True)
        try:
            expected_names=[layout()[kind]['sheet']]
            if kind=='internal':expected_names.append(weight_view(data)[1])
            if wb.sheetnames!=expected_names:raise Rejected('结果只能包含业务明细及原箱重页')
            if wb.properties.title!=WARNING or wb.properties.description!=digest(batch):raise Rejected('受控标识或明细版本不一致')
            if wb.defined_names or wb._external_links:raise Rejected('结果含未允许的名称或外链')
            for sheet in wb:
                if sheet.sheet_state!='visible' or sheet.merged_cells or sheet.tables or sheet.data_validations.count or sheet._charts or sheet._images:
                    raise Rejected('结果含未允许的隐藏内容或对象')
                if any(x.hidden for x in sheet.row_dimensions.values()) or any(x.hidden for x in sheet.column_dimensions.values()):raise Rejected('结果不得隐藏业务字段')
                for rr in sheet:
                    for cell in rr:
                        if cell.comment or cell.hyperlink or cell.data_type=='e':raise Rejected('结果含批注、外链或错误')
            main=wb.worksheets[0];observed=cached.worksheets[0]
            if main.freeze_panes!=layout()[kind]['freeze_panes']:raise Rejected('冻结位置不一致')
            last=len(data[kind]['rows'])+1
            if main.max_row!=last+(kind=='internal') or main.max_column!=layout()[kind]['columns']:raise Rejected('业务明细行列范围不一致')
            for n,values in enumerate([data[kind]['headers'],*data[kind]['rows']],1):
                for col,value in enumerate(values,1):
                    cell=main.cell(n,col)
                    if cell.value!=(None if value=='' else value):raise Rejected(f'保存后内容不符：{kind}/{cell.coordinate}')
                    formula=n>1 and col in ({16} if kind=='customs' else {4,11,14,16,17,18}) and isinstance(value,str) and value.startswith('=')
                    if (cell.data_type=='f')!=formula:raise Rejected('公式或原文本类型发生变化')
                    expected=data['caches'][kind].get(cell.coordinate) if formula else (None if value=='' else value)
                    if observed.cell(n,col).value!=expected:raise Rejected(f'公式缓存/内容复核失败：{kind}/{cell.coordinate}')
            for address,state in data['colors'][kind].items():
                if main[address].fill.fgColor.rgb!=COLORS[state]:raise Rejected(f'结果颜色标记不一致：{kind}/{address}')
            if main.oddHeader.center.text!=WARNING or main.oddFooter.center.text!=LEGEND:raise Rejected('草稿标识或颜色图例丢失')
            if kind=='internal':
                if main.cell(last+1,11).value!=f'=SUM(K2:K{last})' or observed.cell(last+1,11).value!=optional_number(fraction(data['business_total'])):
                    raise Rejected('内部合计不一致')
                view,title=weight_view(data);sheet=wb[title]
                if sheet.max_row!=len(view['rows'])+1 or sheet.max_column!=8:raise Rejected('原箱重内容范围不一致')
                for n,values in enumerate([view['headers'],*view['rows']],1):
                    for col,value in enumerate(values,1):
                        if sheet.cell(n,col).value!=(None if value=='' else value) or sheet.cell(n,col).data_type=='f':raise Rejected('原箱重观测值或文本发生变化')
            with zipfile.ZipFile(path) as z:
                for part in z.namelist():
                    if part.startswith(('customXml/','xl/media/','xl/drawings/','xl/embeddings/')) or 'externalLink' in part or 'comments' in part:
                        raise Rejected('结果包含未允许部件')
                    if part.endswith('.rels') and any(e.get('TargetMode')=='External' for e in ET.fromstring(z.read(part))):raise Rejected('结果包有外部关系')
            if kind=='customs' and privacy:
                for sheet in wb:
                    for rr in sheet:
                        for cell in rr:
                            if cell.data_type!='f' and isinstance(cell.value,str) and privacy.search(cell.value):raise Rejected('海关表发现内部文本')
                if privacy.search(str(wb.properties)):raise Rejected('海关元数据发现内部文本')
            results[kind]={'path':str(path),'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'rows':len(batch['rows']),'reopened':True}
        finally:wb.close();cached.close()
    from .batchflow import require_formal_release
    try:require_formal_release(batch)
    except Rejected:pass
    else:raise Rejected('正式输出门禁异常')
    check_sources(batch)
    return {'files':results,'rows':len(batch['rows']),'confirmed_amount_jpy':data['business_total'],
        'customs_amount_jpy':data['customs_total'],'batch_difference_jpy':data['batch_difference'],
        'rounding_rows':data['rounding_rows'],'adjustment_rows':data['adjustment_rows'],
        'line_amount_difference_rows':sum(t['line_difference']!='0' for t in data['operational_trace']),
        'allocation_rows':data['allocation_rows'],'tax_pending_rows':data['tax_pending_rows'],'price_pending_rows':0,
        'color_counts':{k:{state:sum(v==state for v in values.values()) for state in COLORS} for k,values in data['colors'].items()},
        'formal_rejected':True,'formal_output':False,'privacy_scan':'passed','same_detail_basis':data['basis_hash']}


def render(batch,folder):
    check_sources(batch);folder=Path(folder).resolve()
    for source in batch['sources']:
        path=Path(source['path']).resolve()
        if folder.is_relative_to(path.parent) or path.is_relative_to(folder):raise Rejected('输出须与原件目录隔离')
    if not batch['logistics_check']['passed']:raise Rejected('箱号、箱数或箱组重量核对失败')
    if folder.exists():raise Rejected('不覆盖已有结果版本')
    folder.parent.mkdir(parents=True,exist_ok=True);data=payload(batch)
    with tempfile.TemporaryDirectory(prefix='.colored-arrival-',dir=folder.parent) as temp:
        staging=Path(temp);write_pair(data,staging);comparison=verify(batch,staging,data)
        for kind,name in FILES.items():comparison['files'][kind]['path']=str(folder/name)
        atomic_json(staging/'_verification.json',comparison)
        atomic_json(staging/'_calculation_trace.json',{'basis':data['basis_hash'],'rows':data['operational_trace'],
            'batch_difference':data['batch_difference'],'colors':data['colors'],'formal_output':False})
        check_sources(batch);os.replace(staging,folder)
    return {'directory':str(folder),'comparison':comparison,'files':[str(folder/name) for name in FILES.values()]}
