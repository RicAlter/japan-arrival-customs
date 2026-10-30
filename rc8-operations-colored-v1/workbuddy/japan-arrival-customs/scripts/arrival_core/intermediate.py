"""P2-A single-workbook review export. No Excel return/apply/locking operation exists yet."""
import gc
import json
import os
import tempfile
from decimal import Decimal
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
from openpyxl.worksheet.hyperlink import Hyperlink

from . import _p0_frozen as p0
from .batchflow import conversion
from .core import (DOMAINS, Rejected, box_tax, context, digest, exact_price, now, product_errors,
                   proof_valid, select_historical_material, validate)
from .drafts import excel_number
from .excel_writer import literal, save_exact, sheet
from .intermediate_contract import SCHEMA, LEVELS, DOMAIN_NAMES, build_contract
from .storage import Store, atomic_json, check_sources

FILENAME='统一中间表_待业务审核.xlsx'
WARNING='内部审核中间表，非正式报单。黄色可填；置信度/状态为导出快照，修改后须回传核验，不自动确认或锁定。'

def source_catalog(batch, exported_at):
    result=[]
    for i,src in enumerate(batch['sources']):
        role=src.get('role') or ('authority' if src.get('stage')==3 else 'source'+str(i))
        result.append({'source_id':role+':'+src['sha256'][:16], 'role':role,'file':Path(src['path']).name,
            'sha256':src['sha256'],'sheet':src.get('sheet'),'source_record_id':None,'native_record_id':None,
            'native_version':None,'link':None,'queried_at':None,'file_read_at':exported_at,
            'status':'本地Excel只读快照；未访问飞书，不虚构原生ID/版本/查询时间'})
    return result

def source_ref(row, raw, role='authority'):
    src=row['source'];cell=src.get('cells',{}).get(raw or '',{})
    address=cell.get('cell') or ('第'+str(src['row'])+'行（原字段未提供）')
    return f"{role}:{src['sha256'][:16]} / {src['sheet']}!{address}"

def material_prefill(row):
    selected=row['material'].get('selected')
    if selected:return selected,'current'
    history=select_historical_material(row,row['material'].get('candidates',[]))
    if history['status']=='conflict':return {},'conflict'
    if history['selected']:return history['selected'],'reviewed_history'
    return {'translation':row['material'].get('candidate') or None},'candidate'

def business_value(row,spec,material):
    root,key=spec['key'].split('.',1)
    if root=='proposal':return None
    if root=='material':value=material.get(key)
    else:value=(row.get(root) or {}).get(key)
    if value in ('',None):return None
    if spec['kind']=='boolean':return '是' if value is True else ('否' if value is False else str(value))
    if spec['kind'] in ('decimal','positive_decimal','rate'):
        try:return excel_number(value)
        except (ValueError,TypeError,ZeroDivisionError):return str(value)
    return str(value)

def field_assessment(row,spec,value,material_mode,material):
    key=spec['key'];domain=spec['domain'];proof=row['reviews'][domain]
    reasons=[d.get('reason') or d['id'] for d in row['disputes']
             if d.get('domain')==domain or (domain in ('material','packaging','tax') and d.get('domain')=='identity')]
    if key.startswith('material.') and material_mode=='conflict':reasons.append('已审核同范围历史材质互相冲突')
    if reasons:level='conflict'
    elif value is None:
        level='allowed_absence' if key=='product.yd' and row['product']['yd_absence_reason'].strip() else 'missing'
    elif (proof_valid(proof) and (domain not in ('material','packaging','tax')
          or (proof_valid(row['reviews']['identity']) and (material if domain=='material' else row[domain] or {}).get('context')==context(row)))):
        level='reviewed'
    else:level='candidate' if domain in ('material','packaging','tax') or not spec.get('raw') else 'observed'
    label=LEVELS[level]
    if (domain=='material' and material_mode=='reviewed_history' and not reasons
            and value is not None and proof_valid(row['reviews']['identity'])):
        label=LEVELS['reviewed_history'];level='reviewed_history'
    if key.startswith('material.'):
        evidence=(material.get('source') or source_ref(row,'material'))
        if material_mode=='candidate':evidence+='；仅材质候选，不等于官方完整原文或已审记录'
        if material_mode=='reviewed_history':
            evidence+='；已审历史：'+json.dumps(material.get('approval'),ensure_ascii=False)
    elif domain in ('packaging','tax'):
        evidence=(row.get(domain) or {}).get('source') or '无已采用参数；历史填写仅在候选列展示'
    elif key=='proposal.quantity_unit' or key=='authority.quantity_unit':
        evidence='输入合同D07：计价单位待实际业务签核' if not row['authority']['quantity_unit'] else '当前输入合同的明示计价单位'
    elif key=='authority.currency':evidence='KMT第3步JPY表头及输入合同；不是其他来源价格覆盖'
    else:evidence=source_ref(row,spec.get('raw')) if spec.get('raw') else '业务补录或人工审核；目前无独立字段来源'
    if proof_valid(proof):evidence+=f"；本域确认：{proof['confirmed_by']} / {proof['confirmed_at']} / {proof['reference']}"
    if level=='allowed_absence':evidence+='；原因：'+row['product']['yd_absence_reason']
    if reasons:evidence+='；异常：'+'；'.join(reasons)
    return label,evidence

def row_system(batch,row,hash_value,identity_evidence=None):
    historical=row.get('historical',{});h=(historical.get('internal_reference',{}).get('record') or {})
    c=(historical.get('customs_reference',{}).get('record') or {})
    identity_evidence=identity_evidence or {}
    original=row.get('original_purchase',{});old=([original['record']] if original.get('record') else
        identity_evidence.get('original_candidates',original.get('candidates',[])))
    errors=[];calculations={}
    for name,fn in [('declaration_quantity',conversion),('exact_price',exact_price),('box_tax',box_tax)]:
        try:
            result=fn(row)
            if name=='declaration_quantity':calculations[name]=result.get('declaration_quantity')
            elif name=='exact_price':calculations[name]=(result['unit_price_numerator']+'/'+result['unit_price_denominator'] if result['status']=='exact_ratio_only' else None)
            else:calculations[name]=result.get('box_tax')
        except Rejected as exc:calculations[name]=None;errors.append(str(exc))
    entities=set(p0.expand(row['logistics']['carton_raw']) or [])
    weights=[str(i) for i,g in enumerate(batch.get('weight_data',{}).get('groups',[]),1) if entities&set(g['entities'])]
    missing_labels={'internal_name':'内部商品名','customs_name':'申报品名','category':'品类'}
    missing=[('缺少'+missing_labels.get(e[8:],e[8:])) if e.startswith('missing_') else
             {'yd_absence_reason_required':'无YD原因待说明','frontend_status_pending':'前台记录状态待确认',
              'frontend_absence_reason_required':'无前台原因待说明'}.get(e,e) for e in product_errors(row)]
    value={'row_id':row['id'],'batch_id':batch['id'],'data_revision':batch['revision'],'batch_hash':hash_value,'row_hash':digest(row),
        'pending_domains':'、'.join(DOMAIN_NAMES[d] for d in DOMAINS if not proof_valid(row['reviews'][d])),
        'exceptions':'；'.join([d.get('reason') or d['id'] for d in row['disputes']]+missing+errors),
        'original_status':{'identity_pending':'身份连接待单独核实','unique_source_alignment_not_approval':'同码/箱号/日文名唯一对齐，非身份审核'}.get(original.get('status'),'无日方记录'),
        'original_code':'；'.join(str(r.get('code') or '') for r in old) or None,
        'original_source_rows':'、'.join(str(r.get('row') or '') for r in old) or None,
        'related_source3_rows':'、'.join(str(r) for r in identity_evidence.get('related_authority_rows',[])) or None,
        'identity_candidate_method':identity_evidence.get('candidate_method'),
        'identity_candidate_snapshot':json.dumps(identity_evidence,ensure_ascii=False,separators=(',',':')) if identity_evidence else None,
        'original_quantity':'；'.join(str(r.get('qty') or '') for r in old) or None,
        'original_amount':'；'.join(str(r.get('amount') or '') for r in old) or None,
        'historical_material':h.get('material'),'historical_pack':c.get('pack'),'historical_pieces':h.get('pieces'),
        'historical_rate':h.get('rate'),'historical_vat_amount':h.get('vat'),
        'carton':row['logistics']['carton_raw'],'origin':row['logistics']['origin'],'weight_groups':','.join(weights) or None,
        'source_row':row['source']['row'],'source_record_id':row['source']['sheet']+'!'+str(row['source']['row']),
        'source_version':row['source']['sha256'],'native_record_id':None,'source_link':None,'source_queried_at':None,
        'row_snapshot':json.dumps(row,ensure_ascii=False,separators=(',',':')),**calculations}
    for domain in DOMAINS:
        review=row['reviews'][domain];disputed=any(d.get('domain')==domain for d in row['disputes'])
        value['review_status.'+domain]=('已确认；仍有争议' if disputed else '已确认') if proof_valid(review) else '待确认'
        value['review_proof.'+domain]=json.dumps(review,ensure_ascii=False) if review else None
    return value

def make_view(batch):
    validate(batch);contract=build_contract();state_hash=digest(batch);records=[];quality_counts={}
    from .group_review import group_package
    group_refs={row['id']:{} for row in batch['rows']}
    for domain in ('material','packaging','tax'):
        for group in group_package(batch,domain)['groups']:
            for member in group['rows']:group_refs[member['id']]['group.'+domain]=group['id']
    identity_map={}
    if 'input_contract' in batch and 'original_input_summary' in batch:
        from .business_review import identity_checklist
        identity_map={r['id']:r for r in identity_checklist(batch)['rows']}
    for row in batch['rows']:
        system=row_system(batch,row,state_hash,identity_map.get(row['id']));material,mode=material_prefill(row)
        system.update({'group.'+domain:group_refs[row['id']].get('group.'+domain) for domain in ('material','packaging','tax')})
        values={};assessments={}
        for field in contract['fields']:
            if field['role']=='value':
                value=business_value(row,field,material);values[field['key']]=value
                assessments[field['key']]=field_assessment(row,field,value,mode,material)
        for field in contract['fields']:
            key=field['key']
            if field['role'] in ('confidence','evidence'):
                values[key]=assessments[field['field_key']][0 if field['role']=='confidence' else 1]
                if field['role']=='confidence':quality_counts[values[key]]=quality_counts.get(values[key],0)+1
            elif field['role']=='system':values[key]=system[key]
            elif field['role']=='request':values[key]=None
        record=[values[f['key']] for f in contract['fields']]
        if any(isinstance(v,str) and len(v)>32767 for v in record):raise Rejected('单元格内容超过Excel限制；不能静默截断来源/原文')
        records.append(record)
    return contract,records,quality_counts

def index_records(headers,records,expected_ids,fields):
    if headers!=[f['label'] for f in fields]:
        raise Rejected('主明细字段头不匹配字段合同')
    result={}
    for row in records:
        if all(v is None for v in row):continue
        row_id=row[0]
        if not isinstance(row_id,str) or row_id not in expected_ids:raise Rejected('存在空/未知行ID，新增商品须先从源导入建立身份')
        if row_id in result:raise Rejected('稳定行ID重复，禁止按显示行号/商品码猜测覆盖')
        result[row_id]=list(row)
    if set(result)!=set(expected_ids):raise Rejected('本批行缺失，筛选或删除不能使待确认商品静默消失')
    return result

def index_rows(ws,expected_ids,fields):
    """Structural indexing only, not a return importer. Sorting never changes the identity key."""
    if ws.max_column!=len(fields):raise Rejected('主明细列数不匹配字段合同')
    return index_records([ws.cell(6,i).value for i in range(1,len(fields)+1)],
        ws.iter_rows(min_row=7,max_col=len(fields),values_only=True),expected_ids,fields)

def verify_intermediate(batch,path,contract,records,export_id):
    # Independent parser checks all sheet parts, strings/numbers, metadata and formula injection.
    book=p0.Book(Path(path).read_bytes(),FILENAME)
    if [s.name for s in book.sheets]!=['统一主明细','字段合同','来源目录','箱重资料']:
        raise Rejected('中间表工作表范围异常')
    if book.metadata['external_links'] or book.metadata['comment_parts'] or book.metadata['media_count']:
        raise Rejected('中间表包含未允许的链接/对象')
    if any(s.state!='visible' or any(c.get('f') or c.get('type')=='e' for c in s.cells.values()) for s in book.sheets):
        raise Rejected('中间表含隐藏工作表/公式或错误')
    main=book.sheets[0]
    if main.v(3,'A')!=WARNING or main.v(4,'A')!='导出ID：'+export_id:raise Rejected('中间表版本标记失配')
    from openpyxl.utils import get_column_letter,column_index_from_string
    fields=contract['fields'];columns=[get_column_letter(i) for i in range(1,len(fields)+1)]
    if any(column_index_from_string(col)>len(fields) for row in main.rows.values() for col in row):
        raise Rejected('主明细含字段合同之外的单元格')
    index=index_records([main.v(6,col) for col in columns],
        ([main.v(ri,col) for col in columns] for ri in sorted(main.rows) if ri>=7),
        {r['id'] for r in batch['rows']},fields)
    expected={r[0]:r for r in records}
    for row_id,record in index.items():
        for i,(actual,value) in enumerate(zip(record,expected[row_id])):
            if value in (None,''):value=None
            if isinstance(value,(float,int)) and not isinstance(value,bool):
                if actual is None or Decimal(str(actual))!=Decimal(str(value)):raise Rejected('中间表保存后数字变化：'+contract['fields'][i]['key'])
            elif actual!=value:raise Rejected('中间表保存后字段变化：'+contract['fields'][i]['key'])
    return {'rows':len(index),'columns':len(contract['fields']),'stable_ids_unique':True,'all_rows_preserved':True,
        'all_cells_reopened_equal':True,'formulas':0,'business_state_changed':False,'return_import_implemented':False,
        'review_lock_created':False,'formal_output_enabled':False,'sha256':book.sha256}

def export_intermediate(batch,folder):
    check_sources(batch);Store(folder).guard(batch);before=digest(batch)
    folder=Path(folder).resolve()
    if folder.exists():raise Rejected('中间表导出目录已存在，拒绝覆盖业务填写或旧版本')
    contract,records,counts=make_view(batch)
    export_id=digest({'schema':SCHEMA,'batch_hash':before,'contract_hash':contract['contract_hash']})[:24]
    exported_at=now();catalog=source_catalog(batch,exported_at)
    folder.parent.mkdir(parents=True,exist_ok=True)
    with tempfile.TemporaryDirectory(prefix='.arrival-middle-',dir=folder.parent) as tmp:
        staging=Path(tmp);wb=Workbook();wb.remove(wb.active)
        fields=contract['fields'];editable=[i for i,f in enumerate(fields,1) if f['editable']]
        ws=sheet(wb,'统一主明细','日本到货统一中间表',[f['label'] for f in fields],records,[f['width'] for f in fields],
            warning=WARNING,basis='导出ID：'+export_id,summary=f"{len(records)}行；数据版本{batch['revision']}；未锁定；置信度和审核状态相互独立",inputs=editable)
        ws.freeze_panes='D7';ws.sheet_view.zoomScale=85
        literal(ws['A5'],f'{len(records)}行；未锁定')
        by_key={f['key']:f for f in fields}
        for ci,(label,key) in enumerate([('商品资料','product.internal_name'),('身份版本','identity.version'),
            ('材质','material.original'),('盒规','packaging.source_unit'),('税率','tax.classification'),
            ('审核申请','request.product.request'),('异常原因','exceptions')],2):
            cell=ws.cell(5,ci);literal(cell,label)
            cell.hyperlink=Hyperlink(ref=cell.coordinate,location=f"'统一主明细'!{by_key[key]['column']}6",display=label)
            cell.font=Font(name='Arial',size=10,color='0563C1',underline='single')
        for ri in range(7,7+len(records)):ws.row_dimensions[ri].height=64
        # Column grouping provides navigation without hiding any business values.
        for ci,f in enumerate(fields,1):
            if f['role'] in ('confidence','evidence'):
                ws.column_dimensions[f['column']].outlineLevel=1
            if f.get('choices'):
                validation=DataValidation(type='list',formula1='"'+','.join(f['choices'])+'"',allow_blank=True)
                validation.showErrorMessage=True;validation.error='请选择允许值；如需新业务取值请先确认字段合同'
                ws.add_data_validation(validation);validation.add(f"{f['column']}7:{f['column']}{6+len(records)}")
            if f['role']=='confidence':
                for label,color in [(LEVELS['conflict'],'FCE4D6'),(LEVELS['missing'],'FFF1DB'),(LEVELS['reviewed'],'E2EFDA')]:
                    ws.conditional_formatting.add(f"{f['column']}7:{f['column']}{6+len(records)}",
                        CellIsRule(operator='equal',formula=['"'+label+'"'],fill=PatternFill('solid',fgColor=color)))
            if f['role']=='system':
                for ri in range(7,7+len(records)):ws.cell(ri,ci).font=Font(name='Arial',size=10,color='59636E')
        rows=[[f['column'],f['key'],f['label'],DOMAIN_NAMES.get(f['domain'],'系统'),
               '可填提议/申请' if f['editable'] else '只读快照',f['kind'],'、'.join(f.get('choices',[])),f['rule'],f.get('invalidate','不适用')] for f in fields]
        sheet(wb,'字段合同','字段合同与可修改范围',
            ['列','字段ID','显示字段','审核域','可修改性','数据类型','允许选项','填写/校验规则','P2-B拟定失效范围'],
            rows,[10,48,44,18,24,20,48,76,80],warning='黄色字段是业务提议；Excel颜色/只读标记不等于安全授权，P2-B按可信基准核验',
            basis='字段合同SHA256：'+contract['contract_hash'],summary='所有业务修改集中在统一主明细，不另发独立确认文件')
        sheet(wb,'来源目录','输入来源及版本',
            ['来源ID','用途','源文件','文件SHA256','原生记录ID','原生版本','链接','来源查询时间','本次读取时间','状态'],
            [[s['source_id'],s['role'],s['file'],s['sha256'],s['native_record_id'],s['native_version'],s['link'],s['queried_at'],s['file_read_at'],s['status']] for s in catalog],
            [38,28,80,68,28,28,42,32,32,68],warning='本版仅展示现有本地输入；飞书导出/CLI适配后续接入，不虚构来源',basis='导出ID：'+export_id)
        sheet(wb,'箱重资料','原始箱号组与总重（只读）',['箱重组ID','箱号组','物理箱数','组总重KG','来源工作表/行','人工物流签核'],
            [[str(i),g['carton_raw'],excel_number(g['count']),excel_number(g['total_kg']),
              str(g.get('source')), '待业务签核'] for i,g in enumerate(batch.get('weight_data',{}).get('groups',[]),1)],
            [16,24,18,22,90,28],warning='按物理箱号组核对，不按商品行重复累计；不生成第三张物流业务单据',basis='导出ID：'+export_id)
        save_exact(wb,staging/FILENAME);wb.close();del wb;gc.collect()
        verified=verify_intermediate(batch,staging/FILENAME,contract,records,export_id)
        atomic_json(staging/'field_contract.json',contract)
        # Trusted local export mirror, not a second business confirmation file or a self-signed receipt.
        baseline={'schema':SCHEMA,'export_id':export_id,'batch_id':batch['id'],'revision':batch['revision'],
            'batch_hash':before,'contract_hash':contract['contract_hash'],'exported_at':exported_at,
            'columns':[f['key'] for f in fields],'rows':{r[0]:r for r in records},'source_catalog':catalog}
        atomic_json(staging/'baseline.json',baseline)
        manifest={'schema':SCHEMA,'export_id':export_id,'batch_id':batch['id'],'revision':batch['revision'],'batch_hash':before,
            'contract_hash':contract['contract_hash'],'baseline_hash':digest(baseline),'exported_at':exported_at,
            'confidence_counts':counts,'verification':verified,'business_review_applied':False,'locked':False,
            'notice':'P2-A只导出中间表；回传差异识别/审核生效/重算/版本锁定属于P2-B，当前未实现'}
        atomic_json(staging/'export_manifest.json',manifest)
        check_sources(batch)
        if digest(batch)!=before:raise Rejected('导出不得改变业务明细或审核状态')
        os.replace(staging,folder)
    return {'directory':str(folder),'workbook':str(folder/FILENAME),**manifest}
