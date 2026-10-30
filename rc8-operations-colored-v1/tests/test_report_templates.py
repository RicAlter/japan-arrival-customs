"""Repository-safe template integration tests, with no real product data."""
import copy
import hashlib
import json
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'workbuddy/japan-arrival-customs/scripts'))
from arrival_core.batchflow import import_batch
from arrival_core.contracts import create_raw_contract
from arrival_core.core import Rejected, digest
from arrival_core.drafts import draft_payload, render_pair, verify_pair
from arrival_core.report_templates import layout, load_preview
from arrival_core.workflow import execute
from arrival_core.storage import Store


class TemplateTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        inputs = self.root / 'inputs'; inputs.mkdir()
        self.source = inputs / 'source.xlsx'
        w = Workbook(); s = w.active; s.title = 'invoice (計算)'
        for c, v in {'C17':'ProductCode(JANCode)', 'D17':'DescriptionofGoods', 'G17':'Origin',
                     'H17':'UnitPrice(日本JPY)', 'I17':'Qty', 'K17':'Subtotal',
                     'A18':'HA021-022', 'C18':'0012345678901', 'D18':'TEST raw description', 'G18':'JP', 'H18':50, 'I18':2, 'K18':100,
                     'C19':'TEST-2', 'D19':'TEST second', 'G19':'CN', 'H19':25, 'I19':4, 'K19':100}.items(): s[c] = v
        ws = w.create_sheet('重さとサイズ')
        ws.append(['箱番','重さ','カートン個数','合計重量（KGS)','容積横(CM)','容積幅(CM)','容積高(CM)','容積重量（Ｍ3）'])
        ws.append(['HA21-22',5,2,10,20,30,40,4.8])
        w.save(self.source); w.close()
        self.hash = hashlib.sha256(self.source.read_bytes()).hexdigest()
        self.batch = import_batch(create_raw_contract(self.source, '重さとサイズ'))

    def tearDown(self): self.tmp.cleanup()

    def test_accepted_columns_blank_pending_and_current_source(self):
        from arrival_core.core import product_errors_for_batch
        self.assertNotIn('yd_absence_reason_required', product_errors_for_batch(self.batch,self.batch['rows'][0]))
        from arrival_core.compact_review import build_plan,material_catalog
        plan=build_plan(self.batch,material_catalog(None))
        self.assertFalse(any('YD空缺' in issue['subject'] for issue in plan['issues']))
        before = digest(self.batch)
        out = self.root / 'output' / 'pair'
        result = render_pair(self.batch, out)
        self.assertEqual(result['comparison']['layout'], 'accepted-stage5-stage6-v1')
        w = load_workbook(out / '海关报关单_受控测试草稿.xlsx', data_only=True)
        s = w['图片处理']
        self.assertEqual(s.max_column, 33)
        self.assertEqual(s['B2'].value, '0012345678901')
        self.assertIsNone(s['F2'].value); self.assertIsNone(s['G2'].value); self.assertIsNone(s['P2'].value)
        self.assertEqual(s['Q2'].value, 2); self.assertIsNone(s['Q3'].value)
        self.assertEqual(s['S2'].value, 10); self.assertEqual(s['A3'].value, 'HA021-022')
        self.assertEqual(s.freeze_panes, 'A2'); w.close()
        w = load_workbook(out / '内部到货单_受控测试草稿.xlsx', data_only=True)
        s = w['Sheet1']
        self.assertEqual(s.max_column, 20)
        self.assertEqual([s['I2'].value,s['I3'].value], [2,4])
        self.assertIsNone(s['K4'].value)
        for cell in ('H2','J2','K2','N2','O2','P2','Q2','R2'): self.assertIsNone(s[cell].value, cell)
        weights = w['重さとサイズ (2)']
        self.assertEqual(weights['D2'].value, 10); self.assertIsNone(weights['H2'].value)
        self.assertEqual([weights.cell(2,c).value for c in (5,6,7)], [20,30,40])
        w.close()
        self.assertEqual(before, digest(self.batch))
        self.assertEqual(self.hash, hashlib.sha256(self.source.read_bytes()).hexdigest())

    def test_same_entrypoint_replay_and_formal_rejection(self):
        store = Store(self.root / 'state'); store.save(self.batch)
        request = {'message':'生成海关和内部到货表','batch':self.batch['id']}
        a = execute(store.root, request); b = execute(store.root, request)
        self.assertEqual(a['result']['directory'], b['result']['directory'])
        self.assertFalse(b['formal_output'])
        self.assertEqual(store.load(self.batch['id'])['revision'], 0)

    def test_raw_upload_produces_initial_review_and_accepted_pair(self):
        state=self.root/'natural-entry'
        first=execute(state,{'message':'处理这批日本到货','files':[str(self.source)]})
        result=execute(state,{'message':'生成海关和内部到货表','batch':first['batch']['batch_id']})
        self.assertEqual(result['result']['comparison']['layout'],'accepted-stage5-stage6-v1')
        self.assertFalse(first['formal_output'])

    def test_literal_names_and_internal_isolation(self):
        row = self.batch['rows'][0]
        row['product'].update(internal_name='PRIVATE-FRONTEND', yd='PRIVATE-YD', channel='PRIVATE-CHANNEL',
                              customs_name='=MALICIOUS()', category='TEST category')
        data = draft_payload(self.batch)
        out = self.root/'output'/'literal'
        render_pair(self.batch, out)
        w = load_workbook(out/'海关报关单_受控测试草稿.xlsx')
        self.assertEqual(w['图片处理']['C2'].data_type, 's')
        self.assertEqual(w['图片处理']['C2'].value, '=MALICIOUS()')
        w['图片处理']['C2'].data_type='f'; w.save(out/'海关报关单_受控测试草稿.xlsx'); w.close()
        with self.assertRaisesRegex(Rejected, '单元格类型'): verify_pair(self.batch,out,data)

    def test_unregistered_or_stale_preview_cannot_provide_candidates(self):
        with self.assertRaises((OSError, Rejected)): load_preview(self.root, self.batch, 'a'*64)
        with self.assertRaisesRegex(Rejected, '版本不一致'):
            from arrival_core.report_templates import payload
            payload(self.batch, {'batch_id':self.batch['id'],'base_hash':'0'*64,'proposals':[]})

    def test_approved_tax_parameters_reuse_core_not_historical_constant(self):
        # In-memory arithmetic fixture only; no business commit or D1 import.
        row = self.batch['rows'][0]
        row['authority'].update(stage=3, status='business_confirmed_stage3', quantity_unit='盒')
        row['packaging'] = {'source_unit':'盒','declaration_unit':'件','declaration_factor':'2',
                            'internal_boxes':'1','internal_quantity':'2','pieces_per_box':'2','internal_unit':'盒'}
        for domain in ('price','packaging','tax'): row['reviews'][domain]={'fixture':'technical_arithmetic_only'}
        row['tax']={'tax_unit_price':'50','tax_unit':'盒','duty_rate':'0.1','vat_rate':'0.09','formula_version':'historical-box-v1'}
        data=draft_payload(self.batch)
        self.assertIn('*0.09', data['internal']['rows'][0][15])
        self.assertEqual(data['caches']['internal']['R2'], 19.9)
        self.assertEqual(data['caches']['customs']['P2'], 100)
        row['reviews']['price']=None
        self.assertIsNone(draft_payload(self.batch)['caches']['customs']['P2'])


if __name__ == '__main__': unittest.main()
