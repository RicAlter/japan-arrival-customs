"""Synthetic coverage for the direct confirmed-arrival-v1 P0 path."""
import hashlib
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "workbuddy/japan-arrival-customs/scripts"))

from arrival_core.confirmed_excel import import_confirmed_arrival, REQUIRED_DETAIL_HEADERS, OPTIONAL_DETAIL_HEADERS
from arrival_core.core import Rejected
from arrival_core.batchflow import require_formal_release
from arrival_core.storage import Store
from arrival_core.workflow import execute
from arrival_core.workbuddy_router import inspect_workbooks


def workbook(path, *, no_code_reason="无码采购", amount=100, declaration_qty=2, declared_price=""):
    book = Workbook()
    confirmation = book.active; confirmation.title = "批次确认"
    for key, value in (("模板版本", "confirmed-arrival-v1"), ("批次号", "SYNTHETIC-BATCH-1"),
                       ("币种", "JPY"), ("重量单位", "KG"), ("确认依据编号", "SYNTHETIC-REF"),
                       ("数据版本", "v1")):
        confirmation.append([key, value])
    details = book.create_sheet("已确认明细")
    headers = list(REQUIRED_DETAIL_HEADERS) + list(OPTIONAL_DETAIL_HEADERS)
    details.append(headers)
    values = {"稳定行号":"ROW-1", "箱号":"SYN001", "原产国":"JP", "商品编码":"", "无码原因":no_code_reason,
              "日文品名":"SYNTHETIC JAPANESE", "内部中文名":"PRIVATE-INTERNAL-NAME", "海关申报品名":"Synthetic toy",
              "品类":"toy", "最终申报材质":"PVC", "材质证据引用":"MAT-SYN", "业务确认数量":2,
              "计价单位":"盒", "业务确认金额JPY":amount, "内部盒数":1, "内部数量":2, "每盒件数":2,
              "申报换算系数":declaration_qty / 2, "申报数量":declaration_qty, "申报单位":"盒", "盒规依据":"PACK-SYN",
              "HS/归类编码":"950300", "关税率":0.1, "增值税率":0.13, "税基单价":amount / 2,
              "税基单位":"盒", "税费公式版本":"historical-box-v1", "适用日期":"2026-10-27", "关务依据":"CUSTOMS-SYN",
              "已批准申报单价":declared_price, "单价调整依据":"DECLARED-SYN" if declared_price else "",
              "YD":"", "渠道":"", "内部备注":"INTERNAL-NOTE", "品牌":"", "是否盲盒":"", "是否需要打标签":""}
    details.append([values[key] for key in headers])
    weights = book.create_sheet("箱重")
    weights.append(["箱号组", "实际箱数", "箱组总重量KG", "单箱重KG", "长度CM", "宽度CM", "高度CM"])
    weights.append(["SYN001", 1, 2, 2, 20, 20, 20])
    book.save(path); book.close()


class ConfirmedArrivalTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(); self.output_tmp = tempfile.TemporaryDirectory(); self.root = Path(self.tmp.name)
        self.input = self.root / "confirmed.xlsx"; workbook(self.input)
        self.before = hashlib.sha256(self.input.read_bytes()).hexdigest()

    def tearDown(self): self.output_tmp.cleanup(); self.tmp.cleanup()

    def test_direct_import_route_and_draft_pair(self):
        inspection = inspect_workbooks([self.input])
        self.assertTrue(inspection["ready"]); self.assertEqual(inspection["input_mode"], "confirmed_arrival_v1")
        state = Path(self.output_tmp.name) / "state"
        created = execute(state, {"message":"处理这批日本到货", "files":[str(self.input)]})
        batch_id = created["batch"]["batch_id"]
        batch = Store(state).load(batch_id)
        self.assertEqual(batch["execution_mode"], "confirmed_excel_untrusted")
        self.assertFalse(batch["confirmed_input"]["trusted_identity_verified"])
        result = execute(state, {"message":"生成海关和内部到货表", "batch":batch_id})["result"]
        self.assertFalse(result["comparison"]["formal_output"])
        customs = Path(result["directory"]) / "海关报关单_受控测试草稿.xlsx"
        book = load_workbook(customs, data_only=True)
        content = "\n".join(str(cell.value or "") for sheet in book.worksheets for row in sheet.iter_rows() for cell in row)
        book.close()
        self.assertNotIn("PRIVATE-INTERNAL-NAME", content); self.assertNotIn("INTERNAL-NOTE", content)
        with self.assertRaises(Rejected): require_formal_release(batch)
        self.assertEqual(self.before, hashlib.sha256(self.input.read_bytes()).hexdigest())

    def test_no_code_requires_reason(self):
        bad = self.root / "bad-code.xlsx"; workbook(bad, no_code_reason="")
        with self.assertRaisesRegex(Rejected, "商品编码或无码原因"):
            import_confirmed_arrival(bad)

    def test_nonterminating_declaration_price_is_pending_not_adjusted(self):
        bad = self.root / "nonterminating.xlsx"; workbook(bad, amount=100, declaration_qty=3)
        batch = import_confirmed_arrival(bad)
        from arrival_core.drafts import draft_payload
        data = draft_payload(batch)
        self.assertIsNone(data['customs']['rows'][0][6])
        self.assertEqual(data['caches']['internal']['K2'], 100)
        self.assertIn("尾差规则", batch["rows"][0]["confirmed_input"]["price_blocker"])

    def test_missing_formal_date_is_reported_but_does_not_block_draft_input(self):
        bad = self.root / "missing-date.xlsx"; workbook(bad)
        book = load_workbook(bad); sheet = book["已确认明细"]
        date_column = next(cell.column for cell in sheet[1] if cell.value == "适用日期")
        sheet.cell(2, date_column).value = None; book.save(bad); book.close()
        batch = import_confirmed_arrival(bad)
        self.assertIn("适用日期", batch["rows"][0]["confirmed_input"]["formal_missing"])
        self.assertIsNotNone(batch["rows"][0]["tax"])

    def test_absent_formal_columns_allow_pair_with_gate_report(self):
        book = load_workbook(self.input); sheet = book['已确认明细']
        for col in reversed(range(1, sheet.max_column + 1)):
            if sheet.cell(1, col).value in ('HS/归类编码', '适用日期', '关务依据'):
                sheet.delete_cols(col)
        book.save(self.input); book.close()
        from arrival_core.drafts import render_pair
        result = render_pair(import_confirmed_arrival(self.input), Path(self.output_tmp.name) / 'pair')
        self.assertTrue((Path(result['directory']) / 'gate_report.json').is_file())

    def test_unsafe_quantity_or_weight_and_price_override_are_blocked(self):
        from arrival_core.drafts import render_pair
        for field, value in [('业务确认数量', '0'), ('申报数量', '99'), ('已批准申报单价', '49')]:
            path = self.root / (field + '.xlsx'); workbook(path)
            book = load_workbook(path); sheet = book['已确认明细']
            col = next(c.column for c in sheet[1] if c.value == field)
            sheet.cell(2, col, value); book.save(path); book.close()
            with self.assertRaises(Rejected): import_confirmed_arrival(path)
        batch = import_confirmed_arrival(self.input)
        batch['weight_data']['groups'][0]['entities'] = ['OTHER1']
        with self.assertRaisesRegex(Rejected, '箱号'):
            render_pair(batch, Path(self.output_tmp.name) / 'bad-weight')


if __name__ == "__main__": unittest.main()
