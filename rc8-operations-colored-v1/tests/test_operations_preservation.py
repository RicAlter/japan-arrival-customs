"""Behavior checks for real operations layout, native row identity and partial drafts."""
import copy
import hashlib
import json
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill, Color

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'workbuddy/japan-arrival-customs/scripts'))
from arrival_core.operations_adapter import run, prepare
from arrival_core.core import Rejected
from arrival_core.report_templates import cache_formulas
from arrival_core.workflow import execute
from arrival_core.colored_tables import payload as draft_payload, verify as verify_pair
from arrival_core.workbuddy_router import route


def fixtures(root):
    business = root / 'business.xlsx'; declaration = root / 'declaration.xlsx'
    wb = Workbook(); ws = wb.active; ws.title = 'Sheet1'
    ws.append(['CartonNo.', 'Origin', 'ProductCode(JANCode)', None, None, None, 'DescriptionofGoods', 'Qty',
               'UnitPrice(日本JPY)', 'Subtotal', None, '材质', None, '关税率', '增值税', '关税', '盒装合计税金（日元', 'YD编码', '备注'])
    # Non-contiguous duplicate identity must remain red, black, blue.
    for row, code, spec in ((2, '001001', 'RED-PRIVATE'), (3, '001002', 'BLACK-PRIVATE'), (4, '001001', 'BLUE-PRIVATE')):
        ws.append(['HA001', 'JP', code, code, 'INTERNAL-PRIVATE-' + spec, spec, 'Japanese-' + spec, 2, 50,
                   f'=H{row}*I{row}', 'toy', 'PVC', None, None, None, None, None, 0 if row == 2 else None,
                   '=HYPERLINK("https://invalid.example")' if row == 3 else '原备注-' + spec])
    ws['S3'].data_type = 's'  # Literal user text starting '=' must remain text.
    weights = wb.create_sheet('weights'); weights.append(['箱番', '重さ', 'カートン個数', '合計重量（KGS)']); weights.append(['HA001', 2, 1, 2])
    wb.save(business); wb.close(); cache_formulas(business, {1: {'J2': 100, 'J3': 100, 'J4': 100}, 2: {}})
    wb = Workbook(); ws = wb.active; ws.title = '图片处理'
    ws.append(['箱号', '品番', 'Items & Descriptions', '原產地', '数量（盒装', 'Qty（数量按单个）',
               '單價', '品牌', '品名', '品类', '是否\n盲盒', '是否需要打标签', '材质'])
    for code in ('001001', '001002', '001001'):
        ws.append(['HA001', code, 'Public toy', '日本', 1, 3, None, 'Public brand', 'Public toy', 'toy', '是', None, 'PVC'])
    wb.save(declaration); wb.close()
    return [str(business), str(declaration)]


def confirm(preview):
    result = {'confirm': True, 'business_unit': 'Qty（原表计价口径）', 'internal_unit': 'Qty（原表计价口径）',
              'declaration_unit': '单个', 'reference': 'TEST explicit source scope, conversions and duplicate pairing',
              'source_bundle_hash': preview['source_bundle_hash'], 'special_conversion_hash': preview['special_conversion_hash'],
              'pairing_hash': preview['pairing_hash']}
    if preview.get('reference_evidence_hash'): result['reference_evidence_hash'] = preview['reference_evidence_hash']
    return result


class PreservationTests(unittest.TestCase):
    def setUp(self):
        self.src = tempfile.TemporaryDirectory(); self.out = tempfile.TemporaryDirectory()
        self.root = Path(self.src.name); self.dest = Path(self.out.name) / 'state'; self.files = fixtures(self.root)

    def tearDown(self): self.src.cleanup(); self.out.cleanup()

    def generate(self, **kwargs):
        preview = run(self.dest, self.files, **kwargs)
        self.assertEqual(preview['status'], 'needs_confirmation', preview)
        result = run(self.dest, self.files, units_confirmation=confirm(preview), **kwargs)
        self.assertEqual(result['status'], 'ok', result)
        return result

    def test_order_fields_raw_zero_formula_text_and_known_amount(self):
        result = self.generate(); pair = Path(result['result']['directory'])
        wi = load_workbook(pair / '内部到货单_受控测试草稿.xlsx', data_only=False)
        wc = load_workbook(pair / '海关报关单_受控测试草稿.xlsx', data_only=True)
        try:
            self.assertEqual([wi.worksheets[0].cell(n, 6).value for n in (2, 3, 4)], ['RED-PRIVATE', 'BLACK-PRIVATE', 'BLUE-PRIVATE'])
            self.assertEqual(wi.worksheets[0]['S2'].value, 0)
            self.assertEqual(wi.worksheets[0]['T3'].data_type, 's')
            self.assertEqual(wi.worksheets[0]['T3'].value, '=HYPERLINK("https://invalid.example")')
            self.assertNotIn('②运营原表', wi.sheetnames)
            self.assertNotIn('①处理说明', wi.sheetnames)
            self.assertEqual([wc.worksheets[0].cell(n, 16).value for n in (2, 3, 4)], [99, 99, 99])
            self.assertTrue(all(wc.worksheets[0].cell(n, 7).value == 33 for n in (2, 3, 4)))
            self.assertEqual(wc.worksheets[0]['H2'].value, 'Public brand')
            self.assertEqual(wc.worksheets[0]['K2'].value, '是')
            self.assertEqual(result['delivery_report']['tax_pending_rows'], 3)
        finally: wi.close(); wc.close()

    def test_duplicate_pairing_requires_current_confirmation(self):
        preview = run(self.dest, self.files); units = confirm(preview); units.pop('pairing_hash')
        self.assertEqual(run(self.dest, self.files, units_confirmation=units)['status'], 'needs_input')
        units = confirm(preview); units['pairing_hash'] = 'stale'
        self.assertEqual(run(self.dest, self.files, units_confirmation=units)['status'], 'needs_input')

    def test_mixed_scope_cannot_be_globally_invented(self):
        preview = run(self.dest, self.files); units = confirm(preview)
        units.update(business_unit='件', internal_unit='件')
        self.assertEqual(run(self.dest, self.files, units_confirmation=units)['status'], 'needs_input')

    def test_source_content_and_warning_tampering_rejected(self):
        result = self.generate(); folder = Path(result['standard_input']).parent.parent
        batch = json.loads((folder / 'batch.json').read_text(encoding='utf-8')); pair = folder / 'pair'
        path = pair / '内部到货单_受控测试草稿.xlsx'
        wb = load_workbook(path); wb.worksheets[0]['F3'] = 'MOVED'; wb.save(path); wb.close()
        with self.assertRaises(Rejected): verify_pair(batch, pair, draft_payload(batch))

    def test_reference_evidence_cannot_cross_batch_or_change_hash(self):
        f = self.root / 'reference.json'; f.write_text(json.dumps({'schema': 'operations-reference-trial-v1',
            'operations_source_hashes': ['wrong'], 'formal_output': False}), encoding='utf-8')
        ref = {'path': str(f), 'sha256': hashlib.sha256(f.read_bytes()).hexdigest()}
        self.assertIn('跨批', run(self.dest, self.files, reference_evidence=ref)['reason'])
        ref['sha256'] = 'stale'
        self.assertIn('变化', run(self.dest, self.files, reference_evidence=ref)['reason'])

    def test_route_and_main_workflow_use_same_new_intent(self):
        self.assertEqual(route('处理已确认运营表')['command'], 'operations-intake')
        self.assertEqual(execute(self.dest, {'message': '处理已确认运营表', 'files': self.files})['rows'], 3)

    def test_short_variant_and_numeric_note_do_not_erase_arithmetic(self):
        wb = load_workbook(self.files[0]); wb['Sheet1']['F2'] = 'F2'; wb['Sheet1']['S2'] = 6
        wb.save(self.files[0]); wb.close()
        cache_formulas(self.files[0], {1: {'J2': 100, 'J3': 100, 'J4': 100}, 2: {}})
        result = self.generate()
        self.assertEqual(result['delivery_report']['confirmed_amount_jpy'], '300')

    def reference_fixture(self, gray=False):
        rate = self.root / 'rate.xlsx'; wb = Workbook(); ws = wb.active; ws.title = 'Rates'
        ws['A2'] = '测试税目'; ws['B2'] = 0
        if gray: ws['A2'].fill = PatternFill('solid', fgColor=Color(theme=1, tint=0.25))
        wb.save(rate); wb.close()
        plan = prepare(self.files)
        rows = [{'source_row': item['source_rows']['business'], 'code': item['values']['商品编码'],
                 'category': item['values']['品类'], 'material': item['values']['最终申报材质'],
                 'duty_rate': '0', 'basis': 'TEST零税率参考', 'reason': None,
                 'references': [{'path': str(rate), 'sha256': hashlib.sha256(rate.read_bytes()).hexdigest(),
                                 'sheet': 'Rates', 'cell': 'B2', 'label_cell': 'A2', 'kind': 'tax_table'}]}
                for item in plan['details']]
        doc = {'schema': 'operations-reference-trial-v1', 'operations_source_hashes': sorted(hashlib.sha256(Path(p).read_bytes()).hexdigest() for p in self.files),
               'vat_rate': '0.13', 'notice': 'TEST本批参考，非关务批准', 'formal_output': False, 'rows': rows}
        path = self.root / 'confirmed-reference.json'; path.write_text(json.dumps(doc, ensure_ascii=False), encoding='utf-8')
        return {'path': str(path), 'sha256': hashlib.sha256(path.read_bytes()).hexdigest()}

    def test_zero_duty_is_computed_and_reference_requires_confirmation(self):
        ref = self.reference_fixture(); preview = run(self.dest, self.files, reference_evidence=ref)
        units = confirm(preview); units.pop('reference_evidence_hash')
        self.assertEqual(run(self.dest, self.files, reference_evidence=ref, units_confirmation=units)['status'], 'needs_input')
        result = self.generate(reference_evidence=ref)
        self.assertEqual(result['delivery_report']['tax_pending_rows'], 0)

    def test_gray_expired_reference_is_rejected(self):
        ref = self.reference_fixture(gray=True)
        self.assertIn('标灰过期', run(self.dest, self.files, reference_evidence=ref)['reason'])

    def test_internal_text_in_public_field_blocks_instead_of_silent_deletion(self):
        wb = load_workbook(self.files[1]); wb['图片处理']['H2'] = 'RED-PRIVATE'; wb.save(self.files[1]); wb.close()
        preview = run(self.dest, self.files)
        result = run(self.dest, self.files, units_confirmation=confirm(preview))
        self.assertEqual(result['status'], 'needs_input')
        self.assertIn('内部敏感文本冲突', result['reason'])


if __name__ == '__main__': unittest.main()
