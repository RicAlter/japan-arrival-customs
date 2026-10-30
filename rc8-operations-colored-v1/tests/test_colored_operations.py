"""Operational calculation, new input profile and observable cell colors."""
import json
import sys
import tempfile
import unittest
import zipfile
from xml.etree import ElementTree as ET
from fractions import Fraction
from pathlib import Path

from openpyxl import Workbook, load_workbook

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'workbuddy/japan-arrival-customs/scripts'))
from arrival_core.colored_operations import allocated_rate,round0,prepare_calculation_inputs
from arrival_core.colored_tables import COLORS,payload,verify
from arrival_core.operations_adapter import run,prepare
from arrival_core.report_templates import cache_formulas
from arrival_core.core import Rejected


def fixtures(root,rate='0.05',selected=34):
    business=root/'confirmed.xlsx';customs=root/'executed.xlsx'
    wb=Workbook();s=wb.active;s.title='Sheet1'
    s.append(['CartonNo.','Origin','ProductCode(JANCode)',None,None,None,'DescriptionofGoods','盒装数','Qty','UnitPrice(日本JPY)','Subtotal',None,'材质',None,'关税率','增值税','关税','盒装合计税金（日元','YD编码','备注'])
    s.append(['HA001','JP','0010001','0010001','PRIVATE INTERNAL','PRIVATE SPEC','Japanese',1,2,50,'=I2*J2','card,case','PVC,纸','=I2/H2',float(rate) if rate is not None else None,'=J2*(1+O2)*0.13','=O2*J2','=(P2+Q2)*N2',0,'原备注'])
    w=wb.create_sheet('weights');w.append(['箱番','重さ','カートン個数','合計重量（KGS)']);w.append(['HA001',2,1,2])
    wb.save(business);wb.close()
    cache={'K2':100,'N2':2}
    if rate is not None:
        r=Fraction(rate);cache.update(P2=float(50*(1+r)*Fraction('0.13')),Q2=float(50*r),R2=float((50*(1+r)*Fraction('0.13')+50*r)*2))
    cache_formulas(business,{1:cache,2:{}})
    wb=Workbook();s=wb.active;s.title='图片处理'
    s.append(['箱号','品番','Items & Descriptions','原產地','数量（盒装','Qty（数量按单个）','單價','品牌','品名','品类','是否\n盲盒','是否需要打标签','材质',None,None,'金額'])
    s.append(['HA001','0010001','Public item','日本',1,3,selected,'Public brand','Public item','card,case',None,None,'PVC,纸',None,None,'=G2*F2'])
    wb.save(customs);wb.close();cache_formulas(customs,{1:{'P2':selected*3}})
    return [str(business),str(customs)]


def confirmation(preview):
    result={'confirm':True,'business_unit':'Qty（原表计价口径）','internal_unit':'Qty（原表计价口径）',
            'declaration_unit':'单个','reference':'TEST explicit operational result, units, source and allocation scope',
            **{k:preview[k] for k in ('source_bundle_hash','special_conversion_hash')}}
    if preview.get('pairing_candidates'):result['pairing_hash']=preview['pairing_hash']
    if preview.get('tax_allocation_hash'):result['tax_allocation_hash']=preview['tax_allocation_hash']
    return result


class ColoredTests(unittest.TestCase):
    def setUp(self):
        self.src=tempfile.TemporaryDirectory();self.out=tempfile.TemporaryDirectory()
        self.root=Path(self.src.name);self.dest=Path(self.out.name)/'state';self.files=fixtures(self.root)
    def tearDown(self):self.src.cleanup();self.out.cleanup()
    def execute(self,**kwargs):
        p=run(self.dest,self.files,**kwargs);self.assertEqual(p['status'],'needs_confirmation',p)
        r=run(self.dest,self.files,units_confirmation=confirmation(p),**kwargs);self.assertEqual(r['status'],'ok',r);return r

    def test_round0_and_explicit_allocation_weights(self):
        self.assertEqual(round0(Fraction(928,3)),309)
        self.assertEqual(round0(Fraction(693,2)),347)
        self.assertEqual(allocated_rate([{'rate':'0','weight':'9'},{'rate':'0.1','weight':'9'}],'explicit basis'),Fraction('0.05'))
        self.assertEqual(allocated_rate([{'rate':'0.1','weight':'1'},{'rate':'0.065','weight':'24'}],'explicit basis'),Fraction('0.0664'))

    def test_missing_weights_or_basis_never_guessed(self):
        with self.assertRaises(Rejected):allocated_rate([{'rate':'0.1'}],'basis')
        with self.assertRaises(Rejected):allocated_rate([{'rate':'0.1','weight':1}],'')
        with self.assertRaises(Rejected):allocated_rate([{'rate':'0.1','weight':0}],'basis')

    def test_new_20_column_input_and_only_two_workbooks_with_cell_colors(self):
        r=self.execute();pair=Path(r['result']['directory'])
        self.assertEqual(len(list(pair.glob('*.xlsx'))),2)
        c=load_workbook(pair/'海关报关单_受控测试草稿.xlsx');i=load_workbook(pair/'内部到货单_受控测试草稿.xlsx')
        try:
            self.assertEqual(len(c.sheetnames),1);self.assertEqual(len(i.sheetnames),2)
            self.assertEqual(c.worksheets[0]['G2'].value,34)
            self.assertEqual(c.worksheets[0]['P2'].value,'=G2*F2')
            self.assertEqual(c.worksheets[0]['G2'].fill.fgColor.rgb,COLORS['orange'])
            self.assertEqual(c.worksheets[0]['P1'].fill.fgColor.rgb,COLORS['red'])
            self.assertEqual(i.worksheets[0]['O2'].fill.fgColor.rgb,COLORS['yellow'])
            self.assertEqual(i.worksheets[0]['F2'].value,'PRIVATE SPEC')
            self.assertEqual(i.worksheets[0]['S2'].value,0);self.assertEqual(i.worksheets[0]['T2'].value,'原备注')
        finally:c.close();i.close()

    def test_rounding_cell_does_not_recolor_normal_rows_or_other_columns(self):
        self.files=fixtures(self.root,selected=33);r=self.execute();p=Path(r['business_files'][0]);w=load_workbook(p)
        try:
            self.assertEqual(w.worksheets[0]['G2'].fill.fgColor.rgb,COLORS['yellow'])
            self.assertNotEqual(w.worksheets[0]['B2'].fill.fgColor.rgb,COLORS['yellow'])
        finally:w.close()

    def test_missing_rate_is_red_and_does_not_become_zero(self):
        self.files=fixtures(self.root,rate=None);r=self.execute();p=Path(r['business_files'][1]);w=load_workbook(p,data_only=True)
        try:
            self.assertIsNone(w.worksheets[0]['O2'].value);self.assertIsNone(w.worksheets[0]['R2'].value)
            self.assertEqual(w.worksheets[0]['O2'].fill.fgColor.rgb,COLORS['red'])
        finally:w.close()

    def test_allocation_input_recomputes_rate_and_requires_bound_confirmation(self):
        plan=prepare(self.files);rid=plan['details'][0]['values']['稳定行号']
        allocations={rid:{'components':[{'rate':'0','weight':9},{'rate':'0.1','weight':9}],
                          'reference':'TEST explicit allocation','vat_rate':'0.13'}}
        p=run(self.dest,self.files,tax_allocations=allocations);units=confirmation(p);units['tax_allocation_hash']='stale'
        self.assertEqual(run(self.dest,self.files,tax_allocations=allocations,units_confirmation=units)['status'],'needs_input')
        r=self.execute(tax_allocations=allocations);self.assertEqual(r['delivery_report']['tax_pending_rows'],0)

    def test_zero_rate_is_computed_as_zero_not_missing(self):
        self.files=fixtures(self.root,rate='0');r=self.execute();w=load_workbook(r['business_files'][1],data_only=True)
        try:
            self.assertEqual(w.worksheets[0]['O2'].value,0);self.assertEqual(w.worksheets[0]['Q2'].value,0)
            self.assertEqual(w.worksheets[0]['P2'].value,6.5)
            self.assertNotEqual(w.worksheets[0]['O2'].fill.fgColor.rgb,COLORS['red'])
        finally:w.close()

    def test_explicit_reference_can_supply_only_the_missing_vat_parameter(self):
        plan=prepare(self.files);plan['details'][0]['tax_input']['vat_rate']=None
        reference={'by_source_row':{'2':{'duty_rate':'0.05','basis':'TEST explicit same-batch reference'}},'vat_rate':'0.13'}
        prepare_calculation_inputs(plan,reference)
        tax=plan['details'][0]['tax_input']
        self.assertEqual(tax['rate'],'0.05');self.assertEqual(tax['vat_rate'],'0.13')
        self.assertTrue(tax['reference'])

    def test_bad_tax_cache_blocks_instead_of_silent_recalculation(self):
        cache_formulas(self.files[0],{1:{'K2':100,'N2':2,'P2':999,'Q2':2.5,'R2':18.65},2:{}})
        self.assertIn('缓存',run(self.dest,self.files)['reason'])

    def test_new_profile_boxes_must_match_declaration_source(self):
        w=load_workbook(self.files[1]);w.worksheets[0]['E2']=2;w.save(self.files[1]);w.close()
        cache_formulas(self.files[1],{1:{'P2':102}})
        self.assertIn('盒数',run(self.dest,self.files)['reason'])

    def test_color_tamper_rejected_after_save(self):
        r=self.execute();folder=Path(r['standard_input']).parent.parent
        batch=json.loads((folder/'batch.json').read_text(encoding='utf-8'));pair=folder/'pair'
        path=pair/'海关报关单_受控测试草稿.xlsx'
        with zipfile.ZipFile(path) as z:parts={n:z.read(n) for n in z.namelist()}
        ns={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
        sheet=ET.fromstring(parts['xl/worksheets/sheet1.xml']);styles=ET.fromstring(parts['xl/styles.xml'])
        cell=sheet.find(".//s:c[@r='G2']",ns);xf=styles.find('s:cellXfs',ns)[int(cell.get('s'))]
        fill=styles.find('s:fills',ns)[int(xf.get('fillId'))]
        fill.find('s:patternFill/s:fgColor',ns).set('rgb','FFFFFFFF')
        parts['xl/styles.xml']=ET.tostring(styles,encoding='utf-8',xml_declaration=True)
        with zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as z:
            for name,data in parts.items():z.writestr(name,data)
        with self.assertRaisesRegex(Rejected,'颜色'):verify(batch,pair,payload(batch))


if __name__=='__main__':unittest.main()
