"""Deterministic WorkBuddy language routing and safe workbook preflight."""
from pathlib import Path
from openpyxl import load_workbook
from .core import Rejected
from .storage import Store

ROUTES=(('生成海关和内部到货表','render-drafts'),('生成海关和内部到货','render-drafts'),
 ('查看这批还有哪些问题','status'),('这是我修改后的审核表，继续处理','return-preview'),
 ('生成这批到货的审核表','return-export'),('生成审核表','return-export'),('处理这批日本到货','import-stage3'),
 ('确认提交这批审核','return-commit'),('应用已授权归类规则','rule-apply'),
 ('准备这批业务合同','prepare-contract'),('填写这批业务合同','fill-contract'),
 ('签核这批业务合同','sign-contract'),('检查这批能否开始','preflight'),
 ('预览这批单位换算','preview-conversion'),('确认这批单位换算','confirm-conversion'))
def route(message):
    text=''.join(str(message).split())
    matches=[command for phrase,command in ROUTES if phrase in text]
    if len(set(matches))!=1: raise Rejected('无法唯一识别操作；请使用“处理这批日本到货”等受控表达')
    return {'command':matches[0],'automatic_execution':False,'notice':'仅映射到受控命令；仍须完成预检、差异预览和可信审批'}

def inspect_workbooks(paths):
    found=[]
    for value in paths:
        path=Path(value)
        if path.suffix.lower()!='.xlsx' or not path.is_file(): raise Rejected('只接受存在的 .xlsx 文件')
        wb=load_workbook(path,read_only=True,data_only=False,keep_links=False)
        try:
            sheets=set(wb.sheetnames)
            first=wb.worksheets[0]
            cells={address:str(first[address].value or '').strip() for address in ('A1','C1','D1','J1','K1','L1','D17','M17')}
            weight_sheets=[ws.title for ws in wb.worksheets if str(ws['A1'].value or '').strip()=='箱番' and 'KGS' in str(ws['D1'].value or '')]
        finally: wb.close()
        roles=[]
        if sheets=={'处理总览','待处理问题','完整商品明细'}: kind='review_return'
        # These are the same structural signatures consumed by controlled importers.
        # A familiar filename or a generic “invoice” label is deliberately insufficient.
        elif ('JAN' in cells['C1'].upper() and cells['J1']=='Qty' and
              cells['K1']=='UnitPrice(日本JPY)' and cells['L1']=='Subtotal'): kind='stage3_authority'
        elif 'JAN' in cells['D17'].upper() and cells['M17']=='Qty': kind='invoice'
        elif cells['A1']=='箱番' and 'KGS' in cells['D1']: kind='weight'
        else: kind='unknown'
        if kind!='unknown': roles.append(kind)
        if weight_sheets and 'weight' not in roles: roles.append('weight')
        if kind=='unknown' and roles: kind=roles[0]
        found.append({'file':path.name,'path':str(path),'kind':kind,'sheets':sorted(sheets),
                      'roles':roles,'weight_sheets':weight_sheets,
                      'recognized_by':'controlled_template_structure' if kind!='unknown' else None})
    by={k:[x for x in found if k in x['roles'] or (k=='unknown' and x['kind']=='unknown')] for k in ('stage3_authority','invoice','weight','review_return','unknown')}
    required=['stage3_authority']; conflicts={k:v for k,v in by.items() if k!='unknown' and len(v)>1}
    return {'files':found,'missing':[k for k in required if not by[k]],'conflicts':{k:[x['file'] for x in v] for k,v in conflicts.items()},
      'unknown':[x['file'] for x in by['unknown']],'ready':not any([*[not by[k] for k in required],bool(conflicts),bool(by['unknown'])])}

def batch_brief(root,batch_id):
    from .workflow import describe
    return describe(root,batch_id)

def sync_d1(root,library,batch_id):
    from .decision_library import DecisionLibrary
    saved=False
    try:
        Store(root).load(batch_id);saved=True
        source=DecisionLibrary(library)
        if Path(source._metadata()['source_root']).resolve()!=Path(root).resolve(): raise Rejected('D1库绑定的来源不是本批次存储')
        result=source.import_batch(batch_id)
        return {'batch_saved':True,'knowledge':result,'retry_required':result.get('index_status')!='ready'}
    except Exception as exc:return {'batch_saved':saved,'knowledge':{'status':'pending_retry','error_type':type(exc).__name__},'retry_required':True}


def historical_references(library,criteria,as_of):
    """Render only eligible formal D1 cases; results remain manual candidates."""
    from .decision_library import DecisionLibrary
    try:
        source=DecisionLibrary(library)
        stats=source.statistics(as_of=as_of)
        result=source.search(criteria,as_of=as_of,realm='formal')
        cases=[]
        for match in result['matches']:
            final=match['decision']['human_final']
            cases.append({'declaration_material':final.get('declaration_material'),
                          'classification_code':final.get('classification_code'),
                          'duty_rate':final.get('duty_rate'),'basis':match['match_basis'],
                          'applicable_scope':match['applicable_scope'],
                          'effective_version':match['effective_version'],'eligibility':match['eligibility'],
                          'source':{k:match[k] for k in ('batch_id','decision_id','event_id','source_bundle')},
                          'classification_basis':final.get('classification_basis'),
                          'approval_evidence':match['decision']['reviews']})
        return {'status':'available','formal_learning_samples':stats['formal_learning_samples'],
                'cases':cases,'count':len(cases),'automatic_adoption':False,
                'notice':'历史案例仅供人工审核参考；必须按本票证据重新确认，不能自动采用。'}
    except Exception as exc:
        return {'status':'unavailable','formal_learning_samples':None,'cases':[],
                'automatic_adoption':False,'error_type':type(exc).__name__,
                'notice':'历史案例暂不可查询；不影响已保存审核决定，不能据此跳过人工审核。'}
