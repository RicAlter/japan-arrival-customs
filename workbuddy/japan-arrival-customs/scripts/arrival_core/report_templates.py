"""Accepted stage-5/6 presentation for the existing controlled draft pipeline.

Only styles and headers are installed. Every detail comes from the current batch
or a registered, revalidated B2 preview. No historical business row is copied.
"""
from __future__ import annotations

import copy
import hashlib
import json
import re
import zipfile
from functools import lru_cache
from decimal import Decimal
from fractions import Fraction
from pathlib import Path
from xml.etree import ElementTree as ET

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection
from openpyxl.workbook.properties import CalcProperties
from openpyxl.utils import get_column_letter

from .core import Rejected, digest, validate, pricing_basis
from .excel_writer import literal, save_exact

SCHEMA = "japan-arrival/accepted-draft-pair-1"
LAYOUT = "accepted-stage5-stage6-v1"
FILES = {"customs": "海关报关单_受控测试草稿.xlsx", "internal": "内部到货单_受控测试草稿.xlsx"}
WARNING = "受控草稿，不得外发或用于申报；候选待确认，空白计算项表示条件未齐"


def optional_number(value):
    from .drafts import optional_number as convert
    return convert(str(value)) if value is not None else None


@lru_cache(maxsize=1)
def layout():
    path = Path(__file__).resolve().parents[2] / "templates" / "accepted_layout.json"
    return json.loads(path.read_text(encoding="utf-8"))


def load_preview(root, batch, preview_id):
    """An attachment or JSON request cannot self-register candidate evidence."""
    from .excel_return import ReturnService, read_json, sha
    service = ReturnService(root)
    preview = read_json(service._path("previews", preview_id))
    if digest(preview) != preview_id or preview["batch_id"] != batch["id"] or preview["base_hash"] != digest(batch):
        raise Rejected("候选预览不属于当前批次版本；请重新预览回传表")
    if sha(Path(preview["upload_file"]).read_bytes()) != preview["upload_sha256"]:
        raise Rejected("回传归档已变更")
    rebuilt = service.preview(preview["upload_file"], export_id=preview["export_id"], operations=preview["operations"])
    if rebuilt["preview_id"] != preview_id:
        raise Rejected("候选预览与可信回传登记不一致")
    return preview


def target_name(batch, preview_id=None):
    # New layout cannot reuse an old 12/21-column directory under the same revision.
    suffix = "-candidate-" + preview_id[:16] if preview_id else ""
    return "r" + str(batch["revision"]) + "-" + digest(batch)[:12] + "-" + LAYOUT + suffix


def guarded(expression, *inputs, positive=()):
    checks = [f'{cell}=""' for cell in inputs] + [f'{cell}<=0' for cell in positive]
    return '=IF(OR(' + ','.join(checks) + '),"",' + expression + ')'


def payload(batch, preview=None):
    from . import _p0_frozen as p0
    from .batchflow import conversion, workflow_status
    from .declaration_decisions import selected_material
    from .drafts import private_tokens
    validate(batch)
    formats = layout()
    products, materials = {}, {}
    if preview:
        if preview.get("batch_id") != batch["id"] or preview.get("base_hash") != digest(batch):
            raise Rejected("候选与批次版本不一致")
        for change in preview["proposals"]:
            if change.get("unparsed_note"):
                continue
            for rid in change["row_ids"]:
                if change["action"] == "edit_product":
                    products.setdefault(rid, {}).update(change["values"])
                elif change["action"] == "edit_decision" and change["values"].get("declaration_material"):
                    materials[rid] = change["values"]["declaration_material"]
    groups = batch.get("weight_data", {}).get("groups", [])
    weights, by_entities = [], {}
    for group in groups:
        key = tuple(sorted(group["entities"]))
        if key in by_entities:
            raise Rejected("箱重组重复，不能确定主表重量位置")
        by_entities[key] = group
        dimensions = [None, None, None]
        # Dimensions remain observations. Read only the original source cells.
        source = group.get("source", {})
        if source.get("path") and source.get("sheet") and source.get("row"):
            dimensions = [source.get("dimensions", {}).get(c) for c in ("E", "F", "G")]
        values = [group["carton_raw"], optional_number(group.get("unit_weight_observed")),
                  optional_number(group.get("count")), optional_number(group.get("total_kg")),
                  *[optional_number(v) for v in dimensions], None]
        weights.append(values)
    # Older snapshots do not contain dimension evidence. Recover it from the
    # source workbook by source hash + sheet + row, never historical output rows.
    books = {}
    try:
        for values, group in zip(weights, groups):
            source = group.get("source", {})
            if source.get("path") and source.get("sheet") and source.get("row"):
                path = source["path"]
                if path not in books:
                    data = Path(path).read_bytes()
                    if hashlib.sha256(data).hexdigest() != source["sha256"]:
                        raise Rejected("箱重来源已变更")
                    books[path] = p0.Book(data, path)
                sheet = next(s for s in books[path].sheets if s.name == source["sheet"])
                values[4:7] = [optional_number(p0.text(sheet.v(source["row"], c))) for c in ("E", "F", "G")]
    finally:
        books.clear()
    tokens = private_tokens(batch)
    tokens.update(v for item in products.values() for k, v in item.items()
                  if k in {"internal_name", "yd", "channel", "internal_note"} and v)
    privacy = re.compile('|'.join(re.escape(token) for token in sorted(tokens, key=len, reverse=True))) if tokens else None
    customs, internal, notes, suppressed, caches = [], [], [], [], {"customs": {}, "internal": {}, "weights": {}}
    used = set()
    for er, row in enumerate(batch["rows"], 2):
        product = dict(row["product"], **products.get(row["id"], {}))
        pack = conversion(row)
        basis = pricing_basis(row)
        price_ready = bool(row["reviews"]["price"]) and basis.get("status") in ("business_confirmed_stage3", "business_confirmed_in_batch")
        quantity = optional_number(pack.get("declaration_quantity"))
        amount = Fraction(basis["amount"]) if price_ready else None
        raw_qty = optional_number(pack.get('internal_quantity') if pack.get('status')=='confirmed_exact_conversion' else row["arrival"]["quantity_raw"])
        piece_price, arrival_price = None, None
        if amount is not None and quantity:
            piece_price = optional_number(amount / Fraction(pack["declaration_quantity"]))
            # Excel's finite numeric price must exactly reproduce the approved line.
            if not isinstance(piece_price, (int, float)) or Fraction(str(piece_price)) * Fraction(pack["declaration_quantity"]) != amount:
                piece_price = None
        if amount is not None and raw_qty:
            arrival_price = optional_number(amount / Fraction(str(raw_qty)))
            if not isinstance(arrival_price, (int, float)) or Fraction(str(arrival_price)) * Fraction(str(raw_qty)) != amount:
                arrival_price = None
        material = selected_material(row, approved_only=True)
        candidate_material = materials.get(row["id"]) or (row["material"].get("selected") or {}).get("translation") or row["material"].get("candidate")
        c = [None] * 33
        c[0:4] = [row["logistics"]["carton_raw"], row["identity"]["code"], product["customs_name"] or None,
                   {"JP": "Japan", "CN": "China", "VN": "Vietnam"}.get(row["logistics"]["origin"], row["logistics"]["origin"])]
        c[4:7] = [optional_number(pack.get("internal_boxes")), quantity, piece_price]
        c[8], c[9], c[12] = c[2], product["category"] or None, material or (candidate_material if preview else None)
        c[15] = guarded(f"G{er}*F{er}", f"G{er}", f"F{er}")
        entity_set = set(p0.expand(c[0]) or [])
        matches = [g for key, g in by_entities.items() if entity_set and entity_set.issubset(set(key))]
        if len(matches) == 1:
            group = matches[0]
            key = tuple(sorted(group["entities"]))
            if key not in used:
                c[16:19] = [optional_number(group["count"]), optional_number(group["unit_weight_observed"]), optional_number(group["total_kg"])]
                used.add(key)
        row_notes = []
        if quantity is None: row_notes.append("申报单位及盒规待确认")
        if not price_ready or piece_price is None: row_notes.append("申报金额或单价待确认；原始Subtotal仅作初步到货金额")
        if not material: row_notes.append("申报材质待确认")
        i = [None] * 20
        i[0:3] = [c[0], row["logistics"]["origin"], row["identity"]["code"]]
        i[3] = guarded(f'C{er}&""', f"C{er}")
        i[4], i[6] = product["internal_name"] or None, row["identity"]["japanese_name"] or None
        i[7:13] = [optional_number(pack.get("internal_boxes")), raw_qty, arrival_price,
                     guarded(f"I{er}*J{er}", f"I{er}", f"J{er}"), product["category"] or None, material or candidate_material or None]
        i[13] = guarded(f"I{er}/H{er}", f"I{er}", f"H{er}", positive=(f"H{er}",))
        tax = row.get("tax") if row["reviews"]["tax"] and row["reviews"]["packaging"] else None
        tax_ready = bool(tax and price_ready and arrival_price is not None
                         and Fraction(tax["tax_unit_price"]) == Fraction(str(arrival_price))
                         and Fraction(row["packaging"]["internal_quantity"]) == Fraction(str(raw_qty)))
        if tax_ready:
            i[14] = optional_number(tax["duty_rate"])
            vat_rate = tax["vat_rate"]
            i[15] = guarded(f"J{er}*(1+O{er})*{vat_rate}", f"J{er}", f"O{er}")
            i[16] = guarded(f"O{er}*J{er}", f"O{er}", f"J{er}")
            i[17] = guarded(f"(P{er}+Q{er})*N{er}", f"P{er}", f"Q{er}", f"N{er}")
        else:
            # Do not activate historical 0.13, tax variants or mismatched units.
            i[15:18] = ['=""', '=""', '=""']
            row_notes.append("税率、计税单位及盒装税金待确认")
        i[18] = product["yd"] or None
        i[19] = row["product"].get("internal_note") or None
        if row_notes:
            notes.append([er - 1, row["identity"]["code"], "；".join(row_notes)])
        for col, value in enumerate(c):
            if isinstance(value, str) and privacy and privacy.search(value):
                c[col] = None
                suppressed.append({"ordinal": er - 1, "column": get_column_letter(col + 1)})
        customs.append(c); internal.append(i)
        caches["customs"][f"P{er}"] = optional_number(amount) if piece_price is not None else None
        caches["internal"].update({f"D{er}": row["identity"]["code"] or None,
                                    f"K{er}": optional_number(amount) if arrival_price is not None else None,
                                    f"N{er}": optional_number(Fraction(str(raw_qty)) / Fraction(str(i[7]))) if raw_qty is not None and i[7] else None})
        if tax_ready:
            from .core import box_tax
            tax_values = box_tax(row)
            caches["internal"].update({f"P{er}": optional_number(tax_values["vat"]), f"Q{er}": optional_number(tax_values["duty"]), f"R{er}": optional_number(tax_values["box_tax"])})
    for er, group in enumerate(weights, 2):
        # Weight arithmetic is already checked. Volume uses no unapproved /10000 rule.
        observed_total = group[3]
        if (isinstance(group[1], (int, float)) and isinstance(group[2], (int, float))
            and isinstance(observed_total, (int, float))
            and Fraction(str(group[1])) * Fraction(str(group[2])) == Fraction(str(observed_total))):
            group[3] = guarded(f"B{er}*C{er}", f"B{er}", f"C{er}")
            caches["weights"][f"D{er}"] = observed_total
        group[7] = None
    status = workflow_status(batch)
    return {"schema": SCHEMA, "layout": LAYOUT, "basis_hash": digest(batch), "revision": batch["revision"],
            "preview_id": digest(preview) if preview else None, "warning": WARNING, "formal_output": False,
            "customs": {"headers": formats["customs"]["headers"], "rows": customs},
            "internal": {"headers": formats["internal"]["headers"], "rows": internal},
            "weights": {"headers": formats["weights"]["headers"], "rows": weights}, "notes": notes, "caches": caches,
            "privacy_tokens": sorted(tokens),
            "summary": {"rows": len(customs), "cartons": status["logistics"].get("physical_cartons"),
                        "weight_kg": optional_number(status["logistics"].get("total_weight_kg")),
                        "product_confirmed": status["counts"]["product_confirmed"],
                        "suppressed_private_collisions": suppressed,
                        "candidate_only": bool(preview), "volume_formula": "pending_approval"}}


STYLE_CLASSES = {"font": Font, "fill": PatternFill, "border": Border, "alignment": Alignment, "protection": Protection}


def apply_style(cell, spec):
    for key, cls in STYLE_CLASSES.items():
        setattr(cell, key, cls.from_tree(ET.fromstring(spec[key])))
    cell.number_format = spec["number_format"]


def set_cell(cell, value):
    if isinstance(value, dict) and "formula" in value:
        cell.value = value["formula"]
    else:
        literal(cell, value)


def view_sheet(wb, spec, view, *, trusted_formulas=()):
    ws = wb.create_sheet(spec["sheet"])
    ws.freeze_panes = spec["freeze_panes"]
    ws.sheet_view.showGridLines = spec["show_grid_lines"]
    for column, props in spec["column_dimensions"].items():
        for key, value in props.items():
            if value is not None: setattr(ws.column_dimensions[column], key, value)
    styles = {}
    for er, key in ((1, 'header_styles'), (2, 'data_styles')):
        for col in range(1, spec['columns'] + 1):
            cell = ws.cell(er, col)
            apply_style(cell, spec[key][col - 1])
            styles[(er, col)] = copy.copy(cell._style)
    for er, values in enumerate([view["headers"], *view["rows"]], 1):
        ws.row_dimensions[er].height = spec["header_height"] if er == 1 else spec["data_height"]
        for col, value in enumerate(values, 1):
            cell = ws.cell(er, col)
            # Only internally generated formula positions can execute. User text
            # in names/remarks always passes through literal(), even if it starts '='.
            if er > 1 and col in trusted_formulas and isinstance(value, str) and value.startswith("="):
                cell.value = value
            else:
                literal(cell, value)
            cell._style = styles[(1 if er == 1 else 2, col)]
    ws.oddHeader.center.text = WARNING
    ws.oddHeader.center.size = 9
    ws.oddHeader.center.font = "Arial,Bold"
    ws.oddHeader.center.color = "C00000"
    return ws


def cache_formulas(path, cache):
    """Cache only values calculated independently by the existing exact core."""
    ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    with zipfile.ZipFile(path) as src:
        parts = {name: src.read(name) for name in src.namelist()}
    for index, values in cache.items():
        name = f"xl/worksheets/sheet{index}.xml"
        root = ET.fromstring(parts[name])
        for cell in root.findall(".//s:c", ns):
            if cell.find("s:f", ns) is None: continue
            value = values.get(cell.get("r"))
            old = cell.find("s:v", ns)
            if old is not None: cell.remove(old)
            cell.set("t", "str" if value is None or isinstance(value, str) else "n")
            element = ET.SubElement(cell, "{" + ns["s"] + "}v")
            element.text = "" if value is None else str(value)
        parts[name] = ET.tostring(root, encoding="utf-8", xml_declaration=True)
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as out:
        for name, data in parts.items(): out.writestr(name, data)


def write_pair(data, directory):
    if data["formal_output"] is not False: raise Rejected("模板仅可生成受控草稿")
    formats = layout()
    for kind, filename in FILES.items():
        wb = Workbook(); wb.remove(wb.active)
        wb.properties.title = WARNING
        wb.properties.description = data["basis_hash"]
        wb.calculation = CalcProperties(calcId=0, fullCalcOnLoad=True, forceFullCalc=True)
        formula_columns = {16} if kind == "customs" else {4, 11, 14, 16, 17, 18}
        ws = view_sheet(wb, formats[kind], data[kind], trusted_formulas=formula_columns)
        cache = {1: dict(data["caches"][kind])}
        if kind == "internal":
            last = 1 + len(data[kind]["rows"])
            ws.cell(last + 1, 11, f'=IF(COUNT(K2:K{last})={last - 1},SUM(K2:K{last}),"")')
            amounts = [data["caches"]["internal"].get(f"K{er}") for er in range(2, last + 1)]
            cache[1][f"K{last+1}"] = optional_number(sum((Fraction(str(v)) for v in amounts), Fraction(0))) if all(v is not None for v in amounts) else None
            view_sheet(wb, formats["weights"], data["weights"], trusted_formulas={4})
            cache[2] = data["caches"]["weights"]
        explain = wb.create_sheet("①处理说明")
        info = [["状态", WARNING], ["明细版本SHA256", data["basis_hash"]],
                ["候选预览", data.get("preview_id") or "无独立候选预览"],
                ["数据行数", len(data[kind]["rows"])],
                ["金额口径", "原始Subtotal为初步到货金额；最终申报金额仍待确认"],
                ["缺值处理", "空白计算项表示输入未确认，不能视为金额为零"],
                ["公式口径", "仅使用本批已批准参数；历史0.13、税金变体及体积除数不自动采用"],
                ["商品序号", "商品编码", "本行待确认事项"], *data["notes"]]
        for er, row in enumerate(info, 1):
            for col, value in enumerate(row, 1): literal(explain.cell(er, col), value)
            explain.row_dimensions[er].height = 44 if er == 1 or er > 8 else 28
        explain.column_dimensions["A"].width = 24
        explain.column_dimensions["B"].width = 72
        explain.column_dimensions["C"].width = 90
        for row in explain:
            for cell in row:
                if cell.data_type == "f": cell.data_type = "s"
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        path = Path(directory) / filename
        save_exact(wb, path); wb.close()
        cache_formulas(path, cache)


def verify(batch, directory, data):
    from .drafts import private_tokens
    expected = payload(batch) if not data.get("preview_id") else data
    if data["basis_hash"] != digest(batch) or data["schema"] != SCHEMA:
        raise Rejected("模板结果不属于当前批次版本")
    tokens = private_tokens(batch)
    tokens.update(data.get('privacy_tokens', []))
    tokens.update(str(row[4]) for row in data["internal"]["rows"] if row[4])
    tokens.update(str(row[18]) for row in data["internal"]["rows"] if row[18])
    privacy = re.compile('|'.join(re.escape(token) for token in sorted(tokens, key=len, reverse=True))) if tokens else None
    results = {}
    for kind, filename in FILES.items():
        path = Path(directory) / filename
        book = load_workbook(path, data_only=False)
        cached = load_workbook(path, data_only=True)
        try:
            names = [layout()[kind]["sheet"], *([layout()["weights"]["sheet"]] if kind == "internal" else []), "①处理说明"]
            if book.sheetnames != names or book.properties.title != WARNING or book.properties.description != digest(batch):
                raise Rejected("工作表范围、草稿标识或版本不一致")
            for ws in book:
                if ws.sheet_state != "visible" or ws.merged_cells or ws.tables or ws.data_validations.count or ws._charts or ws._images:
                    raise Rejected("模板含未允许隐藏内容或对象")
                if any(d.hidden for d in ws.row_dimensions.values()) or any(d.hidden for d in ws.column_dimensions.values()):
                    raise Rejected("模板不得隐藏业务字段")
                for row in ws:
                    for cell in row:
                        if cell.comment or cell.hyperlink or cell.data_type == "e": raise Rejected("模板含批注、外链或错误")
            if book.defined_names or book._external_links: raise Rejected("模板含非预期名称或外部引用")
            ws = book.worksheets[0]
            spec = layout()[kind]
            if ws.freeze_panes != spec["freeze_panes"]: raise Rejected("冻结位置不符合验收版式")
            for er, row in enumerate([expected[kind]["headers"], *expected[kind]["rows"]], 1):
                for col, value in enumerate(row, 1):
                    actual = ws.cell(er, col).value
                    if (None if value == "" else value) != actual:
                        raise Rejected(f"模板保存值不一致：{kind}/{er}/{col}")
                    formula_columns = {16} if kind == "customs" else {4, 11, 14, 16, 17, 18}
                    formula_expected = er > 1 and col in formula_columns and isinstance(value, str) and value.startswith("=")
                    if (ws.cell(er, col).data_type == "f") != formula_expected:
                        raise Rejected("公式或用户文本单元格类型被篡改")
            last = len(data[kind]["rows"]) + 1
            expected_end = last + (kind == "internal")
            if ws.max_row != expected_end or ws.max_column != spec["columns"]:
                raise Rejected("模板存在历史残留行或额外字段")
            if kind == "internal":
                w = book.worksheets[1]
                if w.max_row != len(data["weights"]["rows"]) + 1 or w.max_column != 8: raise Rejected("箱重范围不一致")
                for er, row in enumerate([data["weights"]["headers"], *data["weights"]["rows"]], 1):
                    if [w.cell(er, col).value for col in range(1, 9)] != row: raise Rejected("箱重映射不一致")
            for ws in cached:
                if any(c.data_type == "e" for row in ws for c in row): raise Rejected("公式缓存出现错误")
            for address, value in data['caches'][kind].items():
                actual = cached.worksheets[0][address].value
                if actual != value:
                    raise Rejected(f'公式缓存与确定性计算不一致：{kind}/{address}')
            if kind == 'internal':
                end = len(data['internal']['rows']) + 1
                total_formula = f'=IF(COUNT(K2:K{end})={end - 1},SUM(K2:K{end}),"")'
                if book.worksheets[0].cell(end + 1, 11).value != total_formula:
                    raise Rejected('内部合计公式被修改')
            with zipfile.ZipFile(path) as archive:
                for name in archive.namelist():
                    if name.startswith(("customXml/", "xl/media/", "xl/drawings/", "xl/embeddings/")) or "externalLink" in name:
                        raise Rejected("模板包含未允许的包部件")
                    if name.endswith(".rels") and any(e.get("TargetMode") == "External" for e in ET.fromstring(archive.read(name))):
                        raise Rejected("模板出现外部关系")
                    if kind == "customs" and name.endswith(".xml"):
                        root = ET.fromstring(archive.read(name))
                        if privacy and any(privacy.search(el.text or '') for el in root.iter()):
                            raise Rejected("海关表发现内部字段，拒绝交付")
            results[kind] = {"sha256": hashlib.sha256(path.read_bytes()).hexdigest(), "rows": len(data[kind]["rows"]), "reopened": True}
        finally:
            book.close(); cached.close()
    compared = []
    for er in range(2, 2 + len(data['customs']['rows'])):
        left, right = data['caches']['customs'].get(f'P{er}'), data['caches']['internal'].get(f'K{er}')
        if left is not None and right is not None:
            if Fraction(str(left)) != Fraction(str(right)):
                raise Rejected('两份表逐商品金额不一致')
            compared.append(er)
    results.update({"same_detail_basis": data["basis_hash"], "layout": LAYOUT, "privacy_scan": "passed", "formal_output": False,
                    "candidate_only": bool(data.get("preview_id")), "row_amounts_equal": True,
                    "amount_comparison_denominator": len(compared),
                    "declaration_quantity_pending_rows": sum(row[5] is None for row in data["customs"]["rows"]),
                    "authority_amount_total_observed": str(sum(Decimal(r["authority"]["amount"]) for r in batch["rows"])),
                    "quantity_conversion_traceable": True})
    return results
