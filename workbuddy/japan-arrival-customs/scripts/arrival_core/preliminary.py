"""Read-only preliminary business organization generated from a raw Japanese workbook."""
import hashlib
import os
import tempfile
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook, load_workbook

from .core import Rejected, digest, validate
from .drafts import excel_number
from .excel_writer import save_exact, sheet
from .storage import check_sources

FILENAME='KMT_初步业务整理.xlsx'
HEADERS=['序号','原始工作表','原始行','箱号原值及位置','商品编码原值及位置','日文描述原值及位置',
         '产地原值及位置','原始单价JPY及位置','原始Qty及位置','原始到货单位','包装内容及换算依据',
         '原始Subtotal JPY及位置','内部中文名候选','品类候选','YD候选','材质候选','商品资料匹配状态','当前处理状态']
WIDTHS=[9,19,11,25,30,55,22,28,24,20,28,30,42,18,22,30,25,50]


def _cell(row,key):
    cell=row['source'].get('cells',{}).get(key,{})
    value=cell.get('value') or ''
    return (value+'\n'+row['source']['sheet']+'!'+cell.get('cell','')).strip()


def payload(batch):
    validate(batch)
    if batch.get('input_mode')!='raw_invoice_preliminary':
        raise Rejected('初步整理只适用于日方原表直接导入批次')
    rows=[]
    for n,row in enumerate(batch['rows'],1):
        lookup=row.get('product_lookup',{})
        packaging_candidate=(row.get('arrival',{}).get('product_reference_packaging',{}).get('values',{}).get('packaging_label'))
        packaging_display=(row['arrival'].get('packaging_content') or
                           (('资料库候选（未批准，不可换算）：'+packaging_candidate) if packaging_candidate else '待查，不按商品名猜测'))
        missing=[name for name,key in (('中文名','internal_name'),('品类','category'),('YD或无YD原因','yd'),('材质','material'))
                 if not (row['product'].get(key) if key!='material' else row['material'].get('candidate'))]
        rows.append([n,row['source']['sheet'],row['source']['row'],_cell(row,'carton'),_cell(row,'code'),_cell(row,'ja'),
            _cell(row,'origin'),_cell(row,'price'),_cell(row,'qty'),row['arrival'].get('unit_raw') or '待查开定/采购资料',
            packaging_display,_cell(row,'amount'),
            row['product']['internal_name'] or None,row['product']['category'] or None,row['product']['yd'] or None,
            row['material']['candidate'] or None,lookup.get('status','not_configured'),
            '待补：'+'、'.join(missing) if missing else '候选已预填，仍待人工审核'])
    summary=batch['arrival_summary']
    return {'rows':rows,'row_count':len(rows),'subtotal_jpy':summary['subtotal_jpy'],
            'qty_numeric_sum':summary['qty_numeric_sum'],'source_hash':batch['sources'][0]['sha256'],
            'batch_hash':digest(batch),'matches':summary.get('product_reference_matches',0),
            'pending':summary.get('product_reference_pending',len(rows))}


def verify(path,batch,view):
    wb=load_workbook(path,data_only=False,read_only=True)
    try:
        if wb.sheetnames!=['初步整理']:raise Rejected('初步整理页签不一致')
        ws=wb['初步整理']
        if ws.max_row-6!=view['row_count'] or ws.max_column!=len(HEADERS):raise Rejected('初步整理范围不一致')
        if any(c.data_type=='f' for row in ws.iter_rows() for c in row):raise Rejected('初步整理不得包含可执行公式')
        total=Decimal(0)
        saved_rows=ws.iter_rows(min_row=7,values_only=True)
        for values,row in zip(saved_rows,batch['rows']):
            if values[2]!=row['source']['row']:raise Rejected('初步整理原始行号不一致')
            if (values[8] or '').split('\n',1)[0]!=(row['source']['cells']['qty']['value'] or ''):
                raise Rejected('初步整理原始Qty不一致')
            value=(values[11] or '').split('\n',1)[0]
            total+=Decimal(value)
        if total!=Decimal(view['subtotal_jpy']):raise Rejected('初步整理Subtotal合计不一致')
        return {'reopened':True,'rows':view['row_count'],'subtotal_jpy':str(total),
                'no_formulas':True,'raw_qty_units_inferred':False,'formal_output':False}
    finally:wb.close()


def export_preliminary(batch,target):
    check_sources(batch);before=digest(batch);view=payload(batch);target=Path(target).resolve()
    if target.exists():raise Rejected('拒绝覆盖已有初步整理文件')
    target.parent.mkdir(parents=True,exist_ok=True)
    source=Path(batch['sources'][0]['path']);source_before=hashlib.sha256(source.read_bytes()).hexdigest()
    with tempfile.TemporaryDirectory(prefix='.preliminary-',dir=target.parent) as tmp:
        file=Path(tmp)/FILENAME;wb=Workbook();wb.remove(wb.active)
        summary=(f"商品 {view['row_count']} 行；Subtotal合计 {view['subtotal_jpy']} JPY；"
                 f"Qty数值合计 {view['qty_numeric_sum']}（单位混合/未知，不称为件）；商品资料匹配 {view['matches']} 行")
        ws=sheet(wb,'初步整理','KMT初步业务整理',HEADERS,view['rows'],WIDTHS,
            warning='系统从日方原表生成，尚未人工确认；不得视为最终申报金额、单位或归类批准。',
            basis='来源哈希：'+view['source_hash']+'；逐行保留原工作表、行号、单元格位置与原值。',summary=summary)
        ws.freeze_panes='F7';ws.sheet_view.zoomScale=70
        save_exact(wb,file);wb.close();verification=verify(file,batch,view)
        os.replace(file,target)
    if hashlib.sha256(source.read_bytes()).hexdigest()!=source_before or digest(batch)!=before:
        target.unlink(missing_ok=True);raise Rejected('生成期间原件或批次状态发生变化')
    return {'file':str(target),'sha256':hashlib.sha256(target.read_bytes()).hexdigest(),
            'verification':verification,'product_reference_matches':view['matches'],
            'product_reference_pending':view['pending'],'formal_output':False}
