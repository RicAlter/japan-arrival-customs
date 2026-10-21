"""Python-only Excel presentation adapter. No business inference or formula evaluation."""
import math
import io
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter

WIDTHS={'customs':[7,17,20,32,14,28,16,14,25,20,12,57],
        'internal':[7,17,20,65,25,20,20,17,15,15,15,17,15,20,12,38,20,22,25,57,55],
        'weights':[25,18,22,54]}

def save_exact(wb, path):
    """Keep projected decimal literals exact despite openpyxl's %.16g float serializer.

    Only new workbooks are processed. String identifiers/formulas/styles are untouched.
    Independent OOXML verification still checks the final disk bytes against the projection.
    """
    stream=io.BytesIO();wb.save(stream)
    ns={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    with zipfile.ZipFile(stream) as src, zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as dst:
        numeric={f'xl/worksheets/sheet{i}.xml':{cell.coordinate:str(cell.value)
                 for row in ws for cell in row if cell.data_type=='n' and isinstance(cell.value,(int,float))}
                 for i,ws in enumerate(wb.worksheets,1)}
        for name in src.namelist():
            data=src.read(name)
            if name in numeric:
                root=ET.fromstring(data)
                for cell in root.findall('.//s:sheetData/s:row/s:c',ns):
                    value=numeric[name].get(cell.get('r'))
                    if value is not None:
                        element=cell.find('s:v',ns)
                        if element is None:raise ValueError('数值序列化缺少v节点')
                        element.text=value
                data=ET.tostring(root,encoding='utf-8',xml_declaration=True)
            dst.writestr(name,data)

def literal(cell, value):
    if value=='':value=None
    cell.value=value
    if isinstance(value,str):
        cell.data_type='s'  # Never execute strings starting with =, +, - or @.
        # Some preview engines ignore @ on digit strings; explicit zero patterns retain
        # full identifiers there as well. Underlying OOXML type remains string.
        cell.number_format='0'*len(value) if value.isascii() and value.isdigit() else '@'
    elif isinstance(value,(int,float)):
        cell.number_format='#,##0' if isinstance(value,int) else '#,##0.0##############'

def sheet(wb, name, title, headers, rows, widths, *, warning, basis, summary='', inputs=()):
    ws=wb.create_sheet(name);ws.sheet_view.showGridLines=False
    for ri,value in ((2,title),(3,warning),(4,basis),(5,summary)):literal(ws.cell(ri,1),value)
    ws['A2'].font=Font(name='Arial',size=14,bold=True)
    ws['A3'].font=Font(name='Arial',size=10,color='9C2F20')
    # Title, warning and trace stay inspectable, with adequate height instead of shrinking text.
    for ri in (2,3,4,5):
        ws.row_dimensions[ri].height=22
    for col,width in enumerate(widths,1):ws.column_dimensions[get_column_letter(col)].width=width
    for ri,record in enumerate([headers,*rows],6):
        lines=2
        for ci,value in enumerate(record,1):
            cell=ws.cell(ri,ci);literal(cell,value)
            cell.font=Font(name='Arial',size=10,color='17212B')
            cell.alignment=Alignment(wrap_text=True,vertical='center',horizontal='right' if isinstance(value,(int,float)) else 'left')
            if ri==6:
                cell.fill=PatternFill('solid',fgColor='233C56');cell.font=Font(name='Arial',size=10,color='FFFFFF',bold=True)
                cell.alignment=Alignment(wrap_text=True,vertical='center',horizontal='center')
            elif ci in inputs:
                cell.fill=PatternFill('solid',fgColor='FFF1DB');cell.protection=Protection(locked=False)
            if isinstance(value,str):
                lines=max(lines,sum(max(1,math.ceil(sum(2 if ord(c)>255 else 1 for c in line)/max(8,widths[ci-1]-2))) for line in value.split('\n')))
        ws.row_dimensions[ri].height=min(409,max(36,lines*14+8))
    last=6+len(rows);ws.freeze_panes='D7' if len(headers)>4 else 'A7'
    ws.auto_filter.ref=f'A6:{get_column_letter(len(headers))}{max(6,last)}'
    ws.print_title_rows='1:6'  # Creates a standard built-in print name, not business metadata.
    ws.sheet_properties.pageSetUpPr.fitToPage=True
    ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
    # Explicit print area is unnecessary; avoid creating defined names for draft projection checks.
    return ws

def write_pair(payload, directory):
    if payload['formal_output'] is not False:raise ValueError('仅受控草稿')
    summary=f"商品 {payload['summary']['rows']} 行；物理箱数 {payload['summary']['cartons']}；总重 {payload['summary']['weight_kg']} KG"
    for kind,filename,title in [('customs','海关报关单_受控测试草稿.xlsx','海关报关单（受控测试草稿）'),
                                ('internal','内部到货单_受控测试草稿.xlsx','内部到货单（受控测试草稿）')]:
        wb=Workbook();wb.remove(wb.active)
        for name,view,key,heading in [('商品明细',payload[kind],kind,title),('箱重核对',payload['weights'],'weights','箱数与总重量核对（受控测试草稿）')]:
            ws=sheet(wb,name,heading,view['headers'],view['rows'],WIDTHS[key],warning=payload['warning'],
                     basis='明细版本SHA256：'+payload['basis_hash'],summary=summary)
            if kind=='internal' and name=='商品明细':
                for row in ws.iter_rows(min_row=7):
                    for ci in (17,18):row[ci-1].number_format='0.00%'
        save_exact(wb,Path(directory)/filename);wb.close()
