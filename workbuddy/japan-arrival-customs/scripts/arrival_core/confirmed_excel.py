"""Strict ``confirmed-arrival-v1`` workbook adapter.

This adapter deliberately treats the workbook as the batch's supplied facts,
not as an Excel review form and not as a product-library import.  It records
every accepted value against the input file, sheet and row, then reuses the
existing exact arithmetic, logistics and rendering pipeline.
"""
from __future__ import annotations

import hashlib
from datetime import date
from decimal import Decimal
from fractions import Fraction
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from . import _p0_frozen as p0
from .core import Rejected, decimal, new_batch, new_row
from .logistics import reconcile_logistics

SCHEMA = "confirmed-arrival-v1"
CONFIRMATION_SHEET = "批次确认"
DETAIL_SHEET = "已确认明细"
WEIGHT_SHEET = "箱重"

META_FIELDS = ("模板版本", "批次号", "币种", "重量单位", "确认依据编号", "数据版本")
# A controlled draft needs traceable rows and deterministic business arithmetic.
# Classification/tax evidence remains a formal-release gate, rather than making a
# usable draft impossible.  Both layers are recorded separately below.
DRAFT_DETAIL_HEADERS = (
    "稳定行号", "箱号", "原产国", "商品编码", "无码原因", "日文品名", "内部中文名", "海关申报品名", "品类",
    "最终申报材质", "业务确认数量", "计价单位", "业务确认金额JPY",
    "内部盒数", "内部数量", "每盒件数", "申报换算系数", "申报数量", "申报单位",
)
FORMAL_DETAIL_HEADERS = (
    "材质证据引用", "盒规依据", "HS/归类编码", "关税率", "增值税率", "税基单价", "税基单位",
    "税费公式版本", "适用日期", "关务依据",
)
# Backward-compatible export for users of the documented field constants.
REQUIRED_DETAIL_HEADERS = DRAFT_DETAIL_HEADERS
OPTIONAL_DETAIL_HEADERS = FORMAL_DETAIL_HEADERS + (
    "已批准申报单价", "单价调整依据", "YD", "渠道", "内部备注", "品牌", "是否盲盒", "是否需要打标签",
)
TAX_FORMAL_FIELDS = ("关税率", "增值税率", "税基单价", "税基单位", "税费公式版本")
REQUIRED_WEIGHT_HEADERS = ("箱号组", "实际箱数", "箱组总重量KG")
OPTIONAL_WEIGHT_HEADERS = ("单箱重KG", "长度CM", "宽度CM", "高度CM")


def _text(value):
    return "" if value is None else str(value).strip()


def _cell_text(cell):
    if cell.data_type == "f":
        raise Rejected(f"已确认Excel不接受公式单元格：{cell.parent.title}!{cell.coordinate}")
    return _text(cell.value)


def _finite_decimal(value, label, *, positive=False, nonnegative=False):
    try:
        result = decimal(value)
    except Rejected as exc:
        raise Rejected(f"{label}必须是有限十进制数") from exc
    if positive and result <= 0:
        raise Rejected(f"{label}必须大于零")
    if nonnegative and result < 0:
        raise Rejected(f"{label}不得为负数")
    return result


def _terminating(value):
    denominator = Fraction(value).denominator
    while denominator % 2 == 0:
        denominator //= 2
    while denominator % 5 == 0:
        denominator //= 5
    return denominator == 1


def _header_map(ws, required, optional, title):
    headers = {}
    for cell in ws[1]:
        value = _cell_text(cell)
        if not value:
            continue
        if value in headers:
            raise Rejected(f"{title}存在重复表头：{value}")
        headers[value] = cell.column
    missing = [value for value in required if value not in headers]
    unknown = sorted(set(headers) - set(required) - set(optional))
    if missing or unknown:
        detail = []
        if missing:
            detail.append("缺少：" + "、".join(missing))
        if unknown:
            detail.append("未知：" + "、".join(unknown))
        raise Rejected(f"{title}表头不符合{SCHEMA}（{'；'.join(detail)}）")
    return headers


def is_confirmed_arrival_workbook(path):
    """Cheap structural recognition used by the WorkBuddy router."""
    path = Path(path)
    if path.suffix.lower() != ".xlsx" or not path.is_file():
        return False
    book = load_workbook(path, read_only=False, data_only=False, keep_links=False)
    try:
        if set(book.sheetnames) != {CONFIRMATION_SHEET, DETAIL_SHEET, WEIGHT_SHEET}:
            return False
        detail = book[DETAIL_SHEET]
        weights = book[WEIGHT_SHEET]
        _header_map(detail, REQUIRED_DETAIL_HEADERS, OPTIONAL_DETAIL_HEADERS, DETAIL_SHEET)
        _header_map(weights, REQUIRED_WEIGHT_HEADERS, OPTIONAL_WEIGHT_HEADERS, WEIGHT_SHEET)
        meta = {_cell_text(row[0]): _cell_text(row[1]) for row in book[CONFIRMATION_SHEET].iter_rows(max_col=2) if _cell_text(row[0])}
        return meta.get("模板版本") == SCHEMA
    except Rejected:
        return False
    finally:
        book.close()


def _metadata(ws):
    values = {}
    for row in ws.iter_rows(max_col=2):
        key = _cell_text(row[0])
        if not key:
            continue
        if key in values:
            raise Rejected(f"{CONFIRMATION_SHEET}存在重复字段：{key}")
        values[key] = _cell_text(row[1])
    missing = [key for key in META_FIELDS if not values.get(key)]
    if missing:
        raise Rejected(f"{CONFIRMATION_SHEET}缺少必填项：{'、'.join(missing)}")
    if values["模板版本"] != SCHEMA or values["币种"] != "JPY" or values["重量单位"] != "KG":
        raise Rejected("本期仅接受confirmed-arrival-v1、JPY和KG的已确认输入")
    return {key: values[key] for key in META_FIELDS}


def _row_values(ws, headers, row_number):
    return {name: _cell_text(ws.cell(row_number, column)) for name, column in headers.items()}


def _proof(metadata, source_hash, domain):
    return {"kind": SCHEMA, "domain": domain, "reference": metadata["确认依据编号"],
            "source_sha256": source_hash, "trusted_identity": False}


def _formal_missing(values):
    return [field for field in FORMAL_DETAIL_HEADERS if not values.get(field)]


def _date(value, label):
    if not value:
        return ""
    try:
        return date.fromisoformat(value).isoformat()
    except ValueError as exc:
        raise Rejected(f"{label}必须为YYYY-MM-DD") from exc


def _detail_row(path, source_hash, ws, headers, row_number, metadata):
    values = _row_values(ws, headers, row_number)
    if not any(values.values()):
        return None
    # 商品编码与无码原因是一组互斥的替代字段，不能把合法无码行先当作缺少编码。
    missing = [name for name in DRAFT_DETAIL_HEADERS if not values[name] and name not in ("商品编码", "无码原因")]
    if not values["商品编码"] and not values["无码原因"]:
        missing.append("商品编码或无码原因")
    if missing:
        raise Rejected(f"{DETAIL_SHEET}第{row_number}行缺少：{'、'.join(missing)}")
    source = {"path": str(path), "sha256": source_hash, "sheet": DETAIL_SHEET, "row": row_number,
              "stage": "confirmed-arrival-v1", "template": SCHEMA,
              "cells": {name: {"cell": f"{get_column_letter(column)}{row_number}", "column": column,
                                "value": values[name], "formula": None, "type": "s"}
                        for name, column in headers.items()}}
    business_qty = _finite_decimal(values["业务确认数量"], "业务确认数量", positive=True)
    internal_boxes = _finite_decimal(values["内部盒数"], "内部盒数", positive=True)
    internal_qty = _finite_decimal(values["内部数量"], "内部数量", positive=True)
    pieces_per_box = _finite_decimal(values["每盒件数"], "每盒件数", positive=True)
    factor = _finite_decimal(values["申报换算系数"], "申报换算系数", positive=True)
    declaration_qty = _finite_decimal(values["申报数量"], "申报数量", positive=True)
    amount = _finite_decimal(values["业务确认金额JPY"], "业务确认金额JPY", nonnegative=True)
    if internal_qty != business_qty:
        raise Rejected(f"{DETAIL_SHEET}第{row_number}行业务确认数量必须等于内部数量；首期不自动换层")
    if Fraction(internal_qty) / Fraction(internal_boxes) != Fraction(pieces_per_box):
        raise Rejected(f"{DETAIL_SHEET}第{row_number}行内部数量/盒数与每盒件数不一致")
    if Fraction(business_qty) * Fraction(factor) != Fraction(declaration_qty):
        raise Rejected(f"{DETAIL_SHEET}第{row_number}行业务数量×申报换算系数不等于申报数量")
    exact_declaration_price = Fraction(amount) / Fraction(declaration_qty)
    approved_price = values.get("已批准申报单价", "")
    adjustment = values.get("单价调整依据", "")
    if approved_price:
        approved = _finite_decimal(approved_price, "已批准申报单价", nonnegative=True)
        if Fraction(approved) != exact_declaration_price:
            raise Rejected(f"{DETAIL_SHEET}第{row_number}行申报单价与金额不一致；尾差规则未批准")
        if not adjustment:
            raise Rejected(f"{DETAIL_SHEET}第{row_number}行已填申报单价时必须提供调整依据")
    from .drafts import optional_number
    saved_price = optional_number(str(exact_declaration_price))
    price_ready = (isinstance(saved_price, (int, float))
                   and Fraction(str(saved_price)) == exact_declaration_price)
    price_blocker = None if price_ready else "金额/申报数量不能精确写入Excel；须先取得已批准尾差规则"
    formal_missing = _formal_missing(values)
    tax_complete = not any(field in formal_missing for field in TAX_FORMAL_FIELDS)
    tax = None
    if tax_complete:
        tax_price = _finite_decimal(values["税基单价"], "税基单价", nonnegative=True)
        duty = _finite_decimal(values["关税率"], "关税率", nonnegative=True)
        vat = _finite_decimal(values["增值税率"], "增值税率", nonnegative=True)
        if duty > 1 or vat > 1:
            raise Rejected(f"{DETAIL_SHEET}第{row_number}行税率必须在0到1之间")
        if values["税基单位"] != values["计价单位"]:
            raise Rejected(f"{DETAIL_SHEET}第{row_number}行税基单位必须等于计价单位")
        if values["税费公式版本"] != "historical-box-v1":
            raise Rejected(f"{DETAIL_SHEET}第{row_number}行不支持的税费公式版本")
        expected_tax_price = Fraction(amount) / Fraction(business_qty)
        if Fraction(tax_price) != expected_tax_price:
            raise Rejected(f"{DETAIL_SHEET}第{row_number}行税基单价必须精确等于业务金额/业务数量")
        tax = {"classification": values.get("HS/归类编码", ""), "tax_unit_price": str(tax_price),
               "tax_unit": values["税基单位"], "duty_rate": str(duty), "vat_rate": str(vat),
               "formula_version": values["税费公式版本"], "evidence": values.get("关务依据", "")}
    raw = {"carton": values["箱号"], "code": values["商品编码"], "ja": values["日文品名"],
           "origin": values["原产国"], "qty": str(business_qty), "price": tax["tax_unit_price"] if tax else "",
           "amount": str(amount), "internal": values["内部中文名"], "category": values["品类"],
           "material": values["最终申报材质"], "channel": values.get("渠道", ""), "yd": values.get("YD", "")}
    row = new_row(source, raw, currency="JPY", quantity_unit=values["计价单位"], authority_stage=3,
                  authority_status="confirmed_excel_v1", amount_basis="confirmed_excel_line_amount")
    row["confirmed_row_id"] = values["稳定行号"]
    row["identity"]["code_type"] = "absent_not_fabricated" if not values["商品编码"] else "confirmed_source_identifier"
    row["identity"]["specification"] = ""
    row["product"].update({"internal_name": values["内部中文名"], "customs_name": values["海关申报品名"],
                           "category": values["品类"], "yd": values.get("YD", ""),
                           "channel": values.get("渠道", ""), "internal_note": values.get("内部备注", ""),
                           "yd_absence_reason": values["无码原因"] if not values.get("YD", "") else "",
                           "frontend_status": "unknown", "brand": values.get("品牌", ""),
                           "blind_flag": values.get("是否盲盒", ""), "label_flag": values.get("是否需要打标签", "")})
    row["material"] = {"candidate": values["最终申报材质"],
                       "selected": {"translation": values["最终申报材质"], "source": values.get("材质证据引用", ""),
                                    "review_status": "confirmed_input"}, "candidates": []}
    row["packaging"] = {"source_unit": values["计价单位"], "declaration_unit": values["申报单位"],
                         "declaration_factor": str(factor), "internal_boxes": str(internal_boxes),
                         "internal_quantity": str(internal_qty), "pieces_per_box": str(pieces_per_box),
                         "internal_unit": values["计价单位"], "evidence": values.get("盒规依据", "")}
    row["tax"] = tax
    row["declaration_classification"] = {"code": values.get("HS/归类编码", ""),
        "evidence": values.get("关务依据", ""), "applicability_date": values.get("适用日期", "")}
    row["reviews"] = {domain: _proof(metadata, source_hash, domain)
                      for domain in ("product", "identity", "material", "packaging")}
    row["reviews"].update({"price": _proof(metadata, source_hash, "price"),
                           "tax": _proof(metadata, source_hash, "tax") if tax else None})
    row["confirmed_input"] = {"row_id": values["稳定行号"], "no_code_reason": values["无码原因"],
                              "material_evidence": values.get("材质证据引用", ""), "packaging_evidence": values.get("盒规依据", ""),
                              "customs_evidence": values.get("关务依据", ""), "formal_missing": formal_missing,
                              "price_blocker": price_blocker}
    return row


def _weights(path, source_hash, ws):
    headers = _header_map(ws, REQUIRED_WEIGHT_HEADERS, OPTIONAL_WEIGHT_HEADERS, WEIGHT_SHEET)
    groups, issues = [], []
    for row_number in range(2, ws.max_row + 1):
        values = _row_values(ws, headers, row_number)
        if not any(values.values()):
            continue
        missing = [name for name in REQUIRED_WEIGHT_HEADERS if not values[name]]
        if missing:
            raise Rejected(f"{WEIGHT_SHEET}第{row_number}行缺少：{'、'.join(missing)}")
        entities = p0.expand(values["箱号组"])
        if not entities:
            issues.append({"row": row_number, "reason": "unparsed_carton", "carton": values["箱号组"]})
            continue
        count = _finite_decimal(values["实际箱数"], "实际箱数", positive=True)
        total = _finite_decimal(values["箱组总重量KG"], "箱组总重量KG", nonnegative=True)
        if count != count.to_integral_value() or len(entities) != int(count):
            issues.append({"row": row_number, "reason": "range_count_mismatch"})
        unit = values.get("单箱重KG", "")
        if unit:
            _finite_decimal(unit, "单箱重KG", nonnegative=True)
        groups.append({"carton_raw": values["箱号组"], "entities": entities, "count": str(count), "total_kg": str(total),
                       "unit_weight_observed": unit or None,
                       "source": {"path": str(path), "sha256": source_hash, "sheet": WEIGHT_SHEET, "row": row_number,
                                  "dimensions": {"E": values.get("长度CM") or None, "F": values.get("宽度CM") or None,
                                                 "G": values.get("高度CM") or None}}})
    if not groups:
        raise Rejected("箱重表没有可识别的箱组")
    return {"groups": groups, "issues": issues, "unit": "KG", "source_sha256": source_hash}


def import_confirmed_arrival(path):
    path = Path(path).resolve(strict=True)
    data = path.read_bytes()
    source_hash = hashlib.sha256(data).hexdigest()
    book = load_workbook(path, read_only=False, data_only=False, keep_links=False)
    try:
        if set(book.sheetnames) != {CONFIRMATION_SHEET, DETAIL_SHEET, WEIGHT_SHEET}:
            raise Rejected("已确认Excel必须且只能包含批次确认、已确认明细、箱重三页")
        metadata = _metadata(book[CONFIRMATION_SHEET])
        detail = book[DETAIL_SHEET]
        headers = _header_map(detail, REQUIRED_DETAIL_HEADERS, OPTIONAL_DETAIL_HEADERS, DETAIL_SHEET)
        rows, external_ids = [], set()
        for row_number in range(2, detail.max_row + 1):
            row = _detail_row(path, source_hash, detail, headers, row_number, metadata)
            if row is None:
                continue
            if row["confirmed_row_id"] in external_ids:
                raise Rejected(f"{DETAIL_SHEET}稳定行号重复：{row['confirmed_row_id']}")
            external_ids.add(row["confirmed_row_id"])
            rows.append(row)
        if not rows:
            raise Rejected("已确认明细没有商品行")
        weight_data = _weights(path, source_hash, book[WEIGHT_SHEET])
    finally:
        book.close()
    if hashlib.sha256(path.read_bytes()).hexdigest() != source_hash:
        raise Rejected("读取期间已确认Excel发生变化")
    batch = new_batch(rows, [{"path": str(path), "sha256": source_hash, "stage": "confirmed-arrival-v1",
                              "kind": SCHEMA, "role": "confirmed_excel"}])
    batch.update({"input_mode": "confirmed_arrival_v1", "phase": "P0-confirmed-excel-controlled-draft",
                  "product_data_policy": "confirmed-excel-v1", "business_batch_id": metadata["批次号"],
                  "confirmed_input": {"schema": SCHEMA, "metadata": metadata, "source_sha256": source_hash,
                                      "trusted_identity_verified": False, "formal_output_authorized": False},
                  "weight_data": weight_data})
    formal_missing = {row["id"]: row["confirmed_input"]["formal_missing"] for row in rows
                      if row["confirmed_input"]["formal_missing"]}
    price_blocked = {row["id"]: row["confirmed_input"]["price_blocker"] for row in rows
                     if row["confirmed_input"]["price_blocker"]}
    batch["confirmed_input"].update({"formal_missing_rows": formal_missing,
                                      "formal_missing_count": len(formal_missing),
                                      "price_blocked_rows": price_blocked,
                                      "price_blocked_count": len(price_blocked)})
    batch["logistics_check"] = reconcile_logistics(batch["rows"], weight_data)
    batch["confirmed_input"]["gate_version"] = "three-level-v1"
    return batch


def gate_status(batch):
    """Input integrity, draft arithmetic and formal authorization are independent."""
    from .storage import check_sources
    check_sources(batch)
    logistics = reconcile_logistics(batch["rows"], batch["weight_data"])
    rows = [{"row_id": row["confirmed_row_id"], "source_row": row["source"]["row"],
             "price_pending": row["confirmed_input"].get("price_blocker"),
             "formal_missing": row["confirmed_input"].get("formal_missing", []),
             "tax_pending": row["tax"] is None} for row in batch["rows"]]
    return {"input": {"passed": True},
            "draft": {"allowed": logistics["passed"], "logistics": logistics,
                      "price_pending_rows": sum(bool(r["price_pending"]) for r in rows),
                      "tax_pending_rows": sum(r["tax_pending"] for r in rows)},
            "formal": {"allowed": False, "release_implemented": False,
                       "missing_rows": sum(bool(r["formal_missing"]) for r in rows),
                       "host_verified": False}, "rows": rows}
