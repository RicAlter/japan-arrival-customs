"""Explicit, draft-only field projection and independent saved-file verification."""
import json
import os
import subprocess
import tempfile
import zipfile
import xml.etree.ElementTree as ET
from decimal import Decimal, localcontext
from fractions import Fraction
from pathlib import Path
from openpyxl.utils import get_column_letter

from . import _p0_frozen as p0
from .batchflow import conversion, workflow_status
from .core import Rejected, box_tax, decimal, digest, exact_price, validate
from .storage import check_sources

CUSTOMS_HEADERS = ["序号", "箱号", "商品编码", "申报品名（候选/已确认）", "产地", "材质（已确认）",
                   "申报数量", "申报单位", "单价精确比（未量化）", "第3步金额", "币种", "审核状态"]
INTERNAL_HEADERS = ["序号", "箱号", "商品编码", "内部商品名称", "YD", "渠道", "品类", "第3步数量", "计价单位",
                    "内部盒数", "每盒件数", "申报数量", "申报单位", "第3步金额", "币种", "材质（已确认/候选）",
                    "关税率（已确认）", "增值税率（已确认）", "盒装税金（精确值）", "审核状态", "内部备注（含历史候选）"]
WEIGHT_HEADERS = ["箱号组", "实际箱数", "总重量（KG）", "核对状态"]
WARNING = "受控测试草稿，不得外发；未确认规则不参与计算"


def excel_number(value):
    """Preserve non-terminating/over-precision values as explicit exact text, never round prices."""
    if value is None or value == "":
        return None
    fraction = Fraction(value)
    denominator = fraction.denominator
    while denominator % 2 == 0: denominator //= 2
    while denominator % 5 == 0: denominator //= 5
    if denominator != 1:
        return "精确比 " + str(fraction)
    with localcontext() as ctx:
        ctx.prec = max(40, len(str(abs(fraction.numerator))) + len(str(fraction.denominator)) + 10)
        number = Decimal(fraction.numerator) / Decimal(fraction.denominator)
    if len(number.normalize().as_tuple().digits) > 15:
        return "精确值 " + str(number)
    return int(number) if number == number.to_integral_value() else float(number)


def optional_number(value):
    try:
        return excel_number(Fraction(value) if isinstance(value,str) and "/" in value else Fraction(decimal(value)))
    except (Rejected, ValueError, TypeError, ZeroDivisionError):
        return None


def private_tokens(batch):
    tokens = {row["product"][key] for row in batch["rows"] for key in ("internal_name", "yd", "channel", "internal_note")
            if row["product"].get(key)}
    tokens.update(str(note) for row in batch["rows"]
                  if p0.good(note:=(row.get("historical",{}).get("internal_reference",{}).get("record") or {}).get("note")))
    return tokens


def draft_payload(batch, *, candidate_preview=None):
    if batch.get('input_mode') in ('raw_invoice_preliminary', 'confirmed_arrival_v1'):
        from .report_templates import payload
        return payload(batch, candidate_preview)
    if candidate_preview:
        raise Rejected('历史兼容草稿不支持候选预览叠加；请使用原始日方工作簿路径')
    validate(batch)
    status = workflow_status(batch)
    raw=batch.get('input_mode')=='raw_invoice_preliminary'
    customs_headers=list(CUSTOMS_HEADERS);internal_headers=list(INTERNAL_HEADERS)
    warning=("受控测试草稿，不得外发；日方原表Subtotal仅为初步到货金额，未知Qty单位不作换算；未确认规则不参与计算"
             if raw else WARNING)
    if raw:
        customs_headers[9]='原始Subtotal（初步到货金额）'
        internal_headers[7]='原始到货Qty';internal_headers[8]='原始到货单位（待查）'
        internal_headers[13]='原始Subtotal（初步到货金额）'
    tokens = private_tokens(batch)
    customs, internal, suppressed = [], [], []
    for ordinal, row in enumerate(batch["rows"], 1):
        pack = conversion(row)
        price = exact_price(row)
        tax = box_tax(row)
        pending = [d for d in row["reviews"] if not row["reviews"][d]]
        labels={"product":"商品资料","identity":"商品身份","price":"金额","material":"材质","packaging":"盒规","tax":"税率"}
        label = "待确认：" + "、".join(labels[d] for d in pending) if pending else "各域已有确认，仍为受控草稿"
        if row["disputes"]: label += "；有未决争议"
        historical = row.get("historical", {}).get("customs_reference", {}).get("record") or {}
        name = row["product"]["customs_name"]
        if not name:
            name = historical.get("customs_name") or None
            label += "\n申报名候选未审"
        from .declaration_decisions import selected_material
        material = selected_material(row,approved_only=True)
        exact = (price["unit_price_numerator"]+"/"+price["unit_price_denominator"]
                 if price["status"] == "exact_ratio_only" else None)
        amount = optional_number(row["authority"]["amount"])
        c = [ordinal, row["logistics"]["carton_raw"], row["identity"]["code"], name,
             row["logistics"]["origin"], material, optional_number(pack["declaration_quantity"]),
             pack["declaration_unit"], exact, amount, row["authority"]["currency"], label]
        # Conservative test projection: collisions with any known internal field are suppressed.
        for col, value in enumerate(c):
            if isinstance(value, str) and any(token in value for token in tokens):
                c[col] = None
                suppressed.append({"ordinal": ordinal, "column": customs_headers[col]})
        customs.append(c)
        selected = selected_material(row,approved_only=False) if row.get('declaration') else (row["material"]["selected"] or {}).get("translation") or row["material"]["candidate"]
        t = row["tax"] if row["reviews"]["tax"] else None
        historical_note=(row.get("historical",{}).get("internal_reference",{}).get("record") or {}).get("note")
        note=row["product"].get("internal_note") or ("历史候选："+str(historical_note) if p0.good(historical_note) else None)
        internal.append([ordinal, row["logistics"]["carton_raw"], row["identity"]["code"], row["product"]["internal_name"],
            row["product"]["yd"], row["product"]["channel"], row["product"]["category"], optional_number(row["authority"]["quantity"]),
            row["authority"]["quantity_unit"] or None, optional_number(pack.get("internal_boxes")), optional_number(pack.get("pieces_per_box")),
            optional_number(pack["declaration_quantity"]), pack["declaration_unit"], amount, row["authority"]["currency"],
            selected or None, optional_number(t["duty_rate"]) if t else None, optional_number(t["vat_rate"]) if t else None,
            optional_number(tax.get("box_tax")) if tax["status"] != "pending" else None, label, note])
    weight_rows = [[g["carton_raw"], optional_number(g["count"]), optional_number(g["total_kg"]),
                    "算术核对一致，未作人工物流签核" if status["logistics"]["passed"] else "待核对"]
                   for g in batch.get("weight_data", {}).get("groups", [])]
    # Logistics cells are also part of the customs workbook, not a third document.
    for record in weight_rows:
        for i, value in enumerate(record):
            if isinstance(value, str) and any(token in value for token in tokens):
                record[i] = None
                suppressed.append({"scope": "weights", "column": i})
    basis = digest(batch)
    return {"schema": "japan-arrival/draft-pair-1", "basis_hash": basis, "revision": batch["revision"],
            "warning": warning, "formal_output": False,
            "customs": {"headers": customs_headers, "rows": customs},
            "internal": {"headers": internal_headers, "rows": internal},
            "weights": {"headers": WEIGHT_HEADERS, "rows": weight_rows},
            "summary": {"rows": len(batch["rows"]), "cartons": status["logistics"].get("physical_cartons"),
                        "weight_kg": optional_number(status["logistics"].get("total_weight_kg")),
                        "product_confirmed": status["counts"]["product_confirmed"], "suppressed_private_collisions": suppressed}}


def verify_pair(batch, directory, payload):
    if payload.get('schema') == 'japan-arrival/accepted-draft-pair-1':
        from .report_templates import verify
        return verify(batch, directory, payload)
    directory = Path(directory)
    results = {}
    tokens = private_tokens(batch)
    for kind, filename in (("customs", "海关报关单_受控测试草稿.xlsx"), ("internal", "内部到货单_受控测试草稿.xlsx")):
        book = p0.Book((directory/filename).read_bytes(), filename)
        with zipfile.ZipFile(directory/filename) as archive:
            for name in archive.namelist():
                if name.startswith(("customXml/","xl/embeddings/","xl/media/","xl/drawings/")) or name.endswith("vbaProject.bin"):
                    raise Rejected("草稿包含未允许的隐藏对象或自定义内容")
                if name.endswith(".rels"):
                    if any(e.get("TargetMode")=="External" for e in ET.fromstring(archive.read(name))):
                        raise Rejected("草稿不能包含外部关系或超链接")
        if [s.name for s in book.sheets] != ["商品明细", "箱重核对"]:
            raise Rejected("草稿工作表范围不符合双表合同")
        # openpyxl represents the visible autofilter using a built-in range name.
        # Only the exact, bounded filter ranges are allowed; no arbitrary/hidden business names.
        allowed_filters={str(i):f"'{name}'!$A$6:${get_column_letter(len(view['headers']))}${max(6,6+len(view['rows']))}"
                         for i,(name,view) in enumerate((("商品明细",payload[kind]),("箱重核对",payload['weights'])))}
        allowed_titles={str(i):f"'{name}'!$1:$6" for i,name in enumerate(('商品明细','箱重核对'))}
        unexpected_names=[n for n in book.metadata['defined_names'] if not (
            (n.get('name')=='_xlnm._FilterDatabase' and n.get('hidden')=='1'
             and n.get('text')==allowed_filters.get(n.get('localSheetId')))
            or (n.get('name')=='_xlnm.Print_Titles' and n.get('hidden') is None
                and n.get('text')==allowed_titles.get(n.get('localSheetId'))))]
        if book.metadata["external_links"] or unexpected_names or book.metadata["comment_parts"]:
            raise Rejected("草稿包含非预期链接、名称或批注元数据")
        for sheet in book.sheets:
            if sheet.state != "visible" or sheet.hidden_rows or sheet.hidden_columns:
                raise Rejected("草稿不应隐藏数据")
            if any(c.get("type") == "e" or c.get("f") for c in sheet.cells.values()):
                raise Rejected("草稿出现公式/错误；确定性结果快照不应执行用户文本")
        main = book.sheets[0]
        if main.v(3,"A") != payload["warning"] or main.v(4,"A") != "明细版本SHA256："+payload["basis_hash"]:
            raise Rejected("草稿标记或版本失配")
        for sheet, view in ((main,payload[kind]),(book.sheets[1],payload["weights"])):
            allowed = {"A2","A3","A4","A5"}
            if sheet.v(3,"A") != payload["warning"] or sheet.v(4,"A") != "明细版本SHA256："+payload["basis_hash"]:
                raise Rejected("草稿附表标记或版本失配")
            for ri, record in enumerate([view["headers"], *view["rows"]], 6):
                for ci, expected in enumerate(record):
                    addr=get_column_letter(ci+1)+str(ri)
                    allowed.add(addr)
                    actual = sheet.v(ri, get_column_letter(ci+1))
                    if isinstance(expected, (int,float)):
                        if actual != Decimal(str(expected)):
                            raise Rejected(f"草稿保存后数值不一致：{kind}/{sheet.name}!{addr}，保存值{actual}，预期{expected}")
                    elif (p0.text(actual) if actual is not None else None) != (str(expected) if expected not in (None,"") else None):
                        if actual != "'"+str(expected):
                            raise Rejected("草稿保存后文本不一致")
            if set(sheet.cells)-allowed:
                raise Rejected("草稿出现字段投影范围以外的单元格")
        if kind == "customs":
            content = [str(c.get("v") or "") for s in book.sheets for c in s.cells.values()]
            content += [str(book.metadata)]
            if any(token in value for token in tokens for value in content):
                raise Rejected("海关草稿发现内部字段内容，拒绝交付")
        results[kind] = {"sha256": book.sha256, "rows": len(payload[kind]["rows"]), "reopened": True}
    results["same_detail_basis"] = payload["basis_hash"]
    results["row_amounts_equal"] = all(c[9] == i[13] for c,i in zip(payload["customs"]["rows"], payload["internal"]["rows"]))
    results["quantity_conversion_traceable"] = all(c[6:8] == i[11:13] for c,i in zip(payload["customs"]["rows"],payload["internal"]["rows"]))
    results["privacy_scan"] = "passed"
    results["formal_output"] = False
    numeric_amounts=[optional_number(r["authority"]["amount"]) for r in batch["rows"]]
    results["missing_authority_amount_rows"]=sum(v is None for v in numeric_amounts)
    results["authority_amount_total_observed"]=str(sum((decimal(r["authority"]["amount"]) for r in batch["rows"]
                                                       if optional_number(r["authority"]["amount"]) is not None),Decimal(0)))
    results["declaration_quantity_pending_rows"]=sum(row[6] is None for row in payload["customs"]["rows"])
    if not results["row_amounts_equal"] or not results["quantity_conversion_traceable"]:
        raise Rejected("两份草稿逐行金额/换算关系不一致")
    return results


def render_pair(batch, folder, runtime=None, *, previews=False, candidate_preview=None):
    from .storage import atomic_json
    check_sources(batch)
    if batch.get('confirmed_input', {}).get('gate_version') == 'three-level-v1':
        from .confirmed_excel import gate_status
        if not gate_status(batch)['draft']['allowed']:
            raise Rejected('受控草稿阻断：箱号覆盖、箱数或箱组重量核对失败')
    folder = Path(folder).resolve()
    for source in batch["sources"]:
        path = Path(source["path"]).resolve()
        if folder.is_relative_to(path.parent) or path.is_relative_to(folder):
            raise Rejected("草稿目录须与输入原件隔离")
    runtime = runtime or {}
    engine = runtime.get('engine', 'artifact-tool' if runtime.get('node_executable') else 'python')
    if engine not in ('python','artifact-tool'): raise Rejected('未知Excel生成引擎')
    node, modules = runtime.get("node_executable"), runtime.get("module_root")
    if engine=='artifact-tool' and (not node or not modules or not Path(node).is_file() or not Path(modules).is_dir()):
        raise Rejected("可选Node渲染运行时未配置")
    if engine=='python' and previews:
        raise Rejected('Python生成Excel不依赖预览引擎；PNG预览请用开发端独立核验工具')
    folder.parent.mkdir(parents=True, exist_ok=True)
    if folder.exists():
        raise Rejected("该版本草稿目录已存在，拒绝覆盖；重新读取既有版本或审核变更后另存")
    payload = draft_payload(batch, candidate_preview=candidate_preview)
    accepted = payload.get('schema') == 'japan-arrival/accepted-draft-pair-1'
    if accepted and engine != 'python':
        raise Rejected('已验收历史版式使用Python/openpyxl渲染，请移除旧Node引擎配置')
    with tempfile.TemporaryDirectory(prefix=".arrival-draft-", dir=folder.parent) as tmp:
        staging = Path(tmp)
        atomic_json(staging/"draft_payload.json", payload)
        if batch.get('confirmed_input', {}).get('gate_version') == 'three-level-v1':
            atomic_json(staging/"gate_report.json", gate_status(batch))
        if engine=='python':
            if accepted:
                from .report_templates import write_pair
            else:
                from .excel_writer import write_pair
            write_pair(payload,staging)
        else:
            command = [node, str(Path(__file__).resolve().parents[1]/"render_drafts.mjs"), str(staging/"draft_payload.json"), str(staging), modules]
            if previews: command.append("--previews")
            result = subprocess.run(command, capture_output=True, text=True, encoding="utf-8", timeout=180)
            if result.returncode:
                raise Rejected("草稿渲染失败: "+result.stderr[-1200:])
        comparison = verify_pair(batch, staging, payload)
        comparison['excel_engine']=engine
        check_sources(batch)
        atomic_json(staging/"draft_comparison.json", comparison)
        os.replace(staging, folder)
    return {"directory": str(folder), "comparison": comparison, "summary": payload["summary"]}
